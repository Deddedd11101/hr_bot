from collections import defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import List, Optional

from fastapi import APIRouter, Body, Depends, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import FileResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..database import get_session
from ..flow_templates import (
    EMPLOYEE_SCOPE_LABELS,
    NOTIFICATION_RECIPIENT_SCOPE_LABELS,
    RESPONSE_TYPE_LABELS,
    ROLE_SCOPE_LABELS,
    SEND_MODE_LABELS,
    TARGET_FIELD_LABELS,
    TRIGGER_MODE_LABELS,
)
from ..models import Employee, FlowStepTemplate, ScenarioTemplate, StepButtonNotification, SurveyAnswer
from ..time_utils import utc_now
from .scenarios import (
    _apply_workspace_step_update,
    _build_scenario_workspace_payload,
    _copy_template_entity,
    _delete_step_attachment_file,
    _delete_step_subtree,
    _delete_template_entity,
    _generate_workspace_scenario_key,
    _get_workspace_scenario_by_flow_key,
    _load_scenario_editor_data,
    _normalize_workspace_kind,
    _normalize_notification_scope,
    _save_step_attachment,
    _sync_button_notification,
    _sync_workspace_button_notifications,
    _sync_workspace_step_send_notifications,
    _delete_step_tree,
    _workspace_item_label,
    _workspace_node_kind,
)
from .support import render_template, require_api_auth, require_auth

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")


def get_db():
    with get_session() as db:
        yield db


@router.post("/flows/reorder")
@router.post("/surveys/reorder")
async def reorder_templates(
    request: Request,
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    form = await request.form()
    scenario_ids = [int(value) for value in form.getlist("scenario_id") if str(value).isdigit()]
    if not scenario_ids:
        return RedirectResponse(url=request.url.path.rsplit("/", 1)[0], status_code=status.HTTP_303_SEE_OTHER)
    scenarios = db.query(ScenarioTemplate).filter(ScenarioTemplate.id.in_(scenario_ids)).all()
    scenario_map = {scenario.id: scenario for scenario in scenarios}
    for index, scenario_id in enumerate(scenario_ids):
        scenario = scenario_map.get(scenario_id)
        if scenario:
            scenario.sort_order = (index + 1) * 10
    db.commit()
    return RedirectResponse(url=request.url.path.rsplit("/", 1)[0], status_code=status.HTTP_303_SEE_OTHER)


@router.get("/flows")
def scenarios_page(request: Request):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    return RedirectResponse(url="/app/flows/workspace-v2", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/api/flows/workspace")
def scenario_workspace_api(
    request: Request,
    scenario_id: Optional[int] = None,
    kind: str = "scenario",
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    return _build_scenario_workspace_payload(db, scenario_id, kind=kind)


@router.post("/api/flows/workspace/scenarios")
def create_workspace_scenario_api(
    request: Request,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    try:
        kind = _normalize_workspace_kind(str(payload.get("kind") or "scenario"))
        item_label = _workspace_item_label(kind)
        title = str(payload.get("title") or "").strip() or f"Новый {item_label}"
        description = str(payload.get("description") or "").strip() or None

        last_scenario = (
            db.query(ScenarioTemplate)
            .filter(ScenarioTemplate.scenario_kind == kind)
            .order_by(ScenarioTemplate.sort_order.desc(), ScenarioTemplate.id.desc())
            .first()
        )
        next_order = ((last_scenario.sort_order or 0) + 10) if last_scenario else 10
        now = utc_now()
        scenario_key = _generate_workspace_scenario_key(kind)

        table_info = db.execute(text("PRAGMA table_info(scenario_templates)")).fetchall()
        table_columns = {row[1] for row in table_info}

        insert_values = {
            "scenario_key": scenario_key,
            "title": title,
            "sort_order": next_order,
            "scenario_kind": kind,
            "role_scope": "all",
            "employee_scope": "all",
            "trigger_mode": "manual_only",
            "target_employee_id": None,
            "description": description,
        }
        if "created_at" in table_columns:
            insert_values["created_at"] = now
        if "updated_at" in table_columns:
            insert_values["updated_at"] = now

        required_columns_without_default = {
            row[1]
            for row in table_info
            if row[5] == 0 and row[3] == 1 and row[4] is None
        }
        missing_required_columns = sorted(required_columns_without_default - set(insert_values.keys()))
        if missing_required_columns:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Не удалось создать {item_label}: в БД есть обязательные колонки без поддержки в UI ({', '.join(missing_required_columns)}).",
            )

        columns_sql = ", ".join(insert_values.keys())
        placeholders_sql = ", ".join(f":{key}" for key in insert_values.keys())
        db.execute(
            text(f"INSERT INTO scenario_templates ({columns_sql}) VALUES ({placeholders_sql})"),
            insert_values,
        )
        db.commit()
        scenario = db.query(ScenarioTemplate).filter(ScenarioTemplate.scenario_key == scenario_key).first()
        if not scenario:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Не удалось создать {item_label}: запись не найдена после сохранения.",
            )
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Не удалось создать элемент. Попробуй ещё раз.")
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Не удалось создать элемент.")

    return {
        "message": f"{item_label.capitalize()} создан",
        "scenario_id": scenario.id,
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=kind),
    }


@router.post("/api/flows/workspace/scenarios/{scenario_id}/settings")
def update_workspace_scenario_api(
    request: Request,
    scenario_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    scenario = db.get(ScenarioTemplate, scenario_id)
    if not scenario or scenario.scenario_kind not in {"scenario", "survey"}:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий или опрос не найден")

    description = str(payload.get("description") or "").strip()
    role_scope = str(payload.get("role_scope") or "").strip()
    employee_scope = str(payload.get("employee_scope") or "").strip()
    trigger_mode = str(payload.get("trigger_mode") or "").strip()
    target_employee_id = str(payload.get("target_employee_id") or "").strip()

    scenario.description = description[:50] or None
    scenario.role_scope = role_scope if role_scope in ROLE_SCOPE_LABELS else "all"
    scenario.employee_scope = employee_scope if employee_scope in EMPLOYEE_SCOPE_LABELS else "all"
    scenario.trigger_mode = trigger_mode if trigger_mode in TRIGGER_MODE_LABELS else "manual_only"
    scenario.target_employee_id = int(target_employee_id) if target_employee_id.isdigit() else None
    db.commit()

    return {
        "message": "Настройки сохранены",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
    }


@router.post("/api/flows/workspace/scenarios/reorder")
def reorder_workspace_scenarios_api(
    request: Request,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    kind = _normalize_workspace_kind(str(payload.get("kind") or "scenario"))
    scenario_ids = [int(value) for value in (payload.get("scenario_ids") or []) if str(value).isdigit()]
    if not scenario_ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не передан порядок сценариев")

    scenarios = (
        db.query(ScenarioTemplate)
        .filter(
            ScenarioTemplate.id.in_(scenario_ids),
            ScenarioTemplate.scenario_kind == kind,
        )
        .all()
    )
    scenario_map = {scenario.id: scenario for scenario in scenarios}
    for index, scenario_id in enumerate(scenario_ids):
        scenario = scenario_map.get(scenario_id)
        if scenario:
            scenario.sort_order = (index + 1) * 10
    db.commit()

    selected_scenario_id = next((scenario_id for scenario_id in scenario_ids if scenario_id in scenario_map), None)
    return {
        "message": "Порядок обновлён",
        "payload": _build_scenario_workspace_payload(db, selected_scenario_id, kind=kind),
    }


@router.post("/api/flows/workspace/scenarios/bulk-copy")
def copy_workspace_scenarios_api(
    request: Request,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    kind = _normalize_workspace_kind(str(payload.get("kind") or "scenario"))
    scenario_ids = [int(value) for value in (payload.get("scenario_ids") or []) if str(value).isdigit()]
    if not scenario_ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не выбраны сценарии для копирования")

    scenarios = (
        db.query(ScenarioTemplate)
        .filter(
            ScenarioTemplate.id.in_(scenario_ids),
            ScenarioTemplate.scenario_kind == kind,
        )
        .order_by(ScenarioTemplate.sort_order, ScenarioTemplate.id)
        .all()
    )
    if not scenarios:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарии не найдены")

    copied_items = [_copy_template_entity(db, scenario) for scenario in scenarios]
    db.commit()

    return {
        "message": f"Скопировано: {len(copied_items)}",
        "payload": _build_scenario_workspace_payload(db, copied_items[-1].id, kind=kind),
    }


@router.post("/api/flows/workspace/scenarios/bulk-delete")
def delete_workspace_scenarios_api(
    request: Request,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    kind = _normalize_workspace_kind(str(payload.get("kind") or "scenario"))
    scenario_ids = [int(value) for value in (payload.get("scenario_ids") or []) if str(value).isdigit()]
    if not scenario_ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не выбраны сценарии для удаления")

    scenarios = (
        db.query(ScenarioTemplate)
        .filter(
            ScenarioTemplate.id.in_(scenario_ids),
            ScenarioTemplate.scenario_kind == kind,
        )
        .all()
    )
    if not scenarios:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарии не найдены")

    deleted_ids = {scenario.id for scenario in scenarios}
    for scenario in scenarios:
        _delete_template_entity(db, scenario)
    db.commit()

    remaining = (
        db.query(ScenarioTemplate)
        .filter(ScenarioTemplate.scenario_kind == kind)
        .order_by(ScenarioTemplate.sort_order, ScenarioTemplate.id)
        .all()
    )
    selected_scenario_id = next((scenario.id for scenario in remaining if scenario.id not in deleted_ids), None)

    return {
        "message": f"Удалено: {len(deleted_ids)}",
        "payload": _build_scenario_workspace_payload(db, selected_scenario_id, kind=kind),
    }


@router.post("/api/flows/workspace/steps/{step_id}")
def update_workspace_step_api(
    request: Request,
    step_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    step = db.get(FlowStepTemplate, step_id)
    if not step:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Шаг не найден")

    scenario = _get_workspace_scenario_by_flow_key(db, step.flow_key)
    if not scenario:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")

    _apply_workspace_step_update(step, payload, scenario.scenario_kind)
    _sync_workspace_button_notifications(db, step, payload, scenario.scenario_kind)
    _sync_workspace_step_send_notifications(db, step, payload, scenario.scenario_kind)
    db.commit()

    return {
        "message": "Вопрос сохранён" if scenario.scenario_kind == "survey" else "Шаг сохранён",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "step_id": step.id,
    }


@router.post("/api/flows/workspace/scenarios/{scenario_id}/steps")
def create_workspace_root_step_api(
    request: Request,
    scenario_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    scenario = db.get(ScenarioTemplate, scenario_id)
    if not scenario or scenario.scenario_kind not in {"scenario", "survey"}:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")

    last_step = (
        db.query(FlowStepTemplate)
        .filter(
            FlowStepTemplate.flow_key == scenario.scenario_key,
            FlowStepTemplate.parent_step_id.is_(None),
        )
        .order_by(FlowStepTemplate.sort_order.desc(), FlowStepTemplate.id.desc())
        .first()
    )
    next_order = (last_step.sort_order + 10) if last_step else 10
    title = str(payload.get("title") or ("Новый вопрос" if scenario.scenario_kind == "survey" else "Новый шаг")).strip()
    title = title or ("Новый вопрос" if scenario.scenario_kind == "survey" else "Новый шаг")

    step = FlowStepTemplate(
        flow_key=scenario.scenario_key,
        step_key=f"{scenario.scenario_key}_step_{int(utc_now().timestamp())}",
        step_title=title,
        sort_order=next_order,
        default_text="Новый вопрос" if scenario.scenario_kind == "survey" else "Новое сообщение сценария.",
        custom_text=None,
        response_type="text" if scenario.scenario_kind == "survey" else "none",
        button_options=None,
        send_mode="immediate",
        send_time=None,
        day_offset_workdays=0,
        target_field=None,
        send_employee_card=False,
    )
    db.add(step)
    db.commit()

    return {
        "message": "Вопрос добавлен" if scenario.scenario_kind == "survey" else "Шаг добавлен",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "step_id": step.id,
    }


@router.post("/api/flows/workspace/scenarios/{scenario_id}/steps/reorder")
def reorder_workspace_root_steps_api(
    request: Request,
    scenario_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    scenario = db.get(ScenarioTemplate, scenario_id)
    if not scenario or scenario.scenario_kind not in {"scenario", "survey"}:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")

    step_ids = [int(value) for value in (payload.get("step_ids") or []) if str(value).isdigit()]
    if not step_ids:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не передан порядок шагов")

    steps = (
        db.query(FlowStepTemplate)
        .filter(
            FlowStepTemplate.flow_key == scenario.scenario_key,
            FlowStepTemplate.parent_step_id.is_(None),
            FlowStepTemplate.id.in_(step_ids),
        )
        .all()
    )
    step_map = {step.id: step for step in steps}
    for index, step_id in enumerate(step_ids):
        step = step_map.get(step_id)
        if step:
            step.sort_order = (index + 1) * 10
    db.commit()

    return {
        "message": "Порядок вопросов обновлён" if scenario.scenario_kind == "survey" else "Порядок шагов обновлён",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
    }


@router.post("/api/flows/workspace/steps/{step_id}/branches")
def create_workspace_branch_step_api(
    request: Request,
    step_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    parent_step = db.get(FlowStepTemplate, step_id)
    if not parent_step:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Шаг не найден")

    scenario = _get_workspace_scenario_by_flow_key(db, parent_step.flow_key)
    if not scenario:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")
    if parent_step.response_type != "branching":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Ветки можно создавать только для шага с ветвлением")

    try:
        option_index = int(payload.get("option_index"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Не удалось определить кнопку для ветки")

    button_labels = [item.strip() for item in (parent_step.button_options or "").splitlines() if item.strip()]
    if option_index < 0 or option_index >= len(button_labels):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Кнопка для ветки не найдена")

    existing_branch = (
        db.query(FlowStepTemplate)
        .filter(
            FlowStepTemplate.flow_key == parent_step.flow_key,
            FlowStepTemplate.parent_step_id == parent_step.id,
            FlowStepTemplate.branch_option_index == option_index,
        )
        .first()
    )
    if existing_branch:
        return {
            "message": "Ветка уже существует",
            "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
            "step_id": existing_branch.id,
        }

    button_label = button_labels[option_index]
    branch_step = FlowStepTemplate(
        flow_key=parent_step.flow_key,
        step_key=f"{parent_step.step_key}__branch_{option_index}",
        parent_step_id=parent_step.id,
        branch_option_index=option_index,
        step_title=f"Ветка: {button_label}",
        sort_order=(parent_step.sort_order or 0) * 100 + option_index + 1,
        default_text="Новое сообщение сценария.",
        custom_text=None,
        response_type="none",
        button_options=None,
        send_mode="immediate",
        send_time=None,
        day_offset_workdays=0,
        target_field=None,
        send_employee_card=False,
    )
    db.add(branch_step)
    db.commit()

    return {
        "message": "Ветка создана",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "step_id": branch_step.id,
    }


@router.post("/api/flows/workspace/steps/{step_id}/chain")
def create_workspace_chain_step_api(
    request: Request,
    step_id: int,
    payload: dict = Body(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    parent_step = db.get(FlowStepTemplate, step_id)
    if not parent_step:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Шаг не найден")

    scenario = _get_workspace_scenario_by_flow_key(db, parent_step.flow_key)
    if not scenario:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")
    if parent_step.parent_step_id is None or parent_step.branch_option_index is None or parent_step.response_type != "chain":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Шаг цепочки можно добавить только внутри ветки с типом «Цепочка шагов»")

    last_step = (
        db.query(FlowStepTemplate)
        .filter(
            FlowStepTemplate.flow_key == parent_step.flow_key,
            FlowStepTemplate.parent_step_id == parent_step.id,
            FlowStepTemplate.branch_option_index.is_(None),
        )
        .order_by(FlowStepTemplate.sort_order.desc(), FlowStepTemplate.id.desc())
        .first()
    )
    next_order = (last_step.sort_order + 10) if last_step else 10
    title = str(payload.get("title") or "Шаг цепочки").strip() or "Шаг цепочки"

    chain_step = FlowStepTemplate(
        flow_key=parent_step.flow_key,
        step_key=f"{parent_step.step_key}__chain_{int(utc_now().timestamp())}",
        parent_step_id=parent_step.id,
        branch_option_index=None,
        step_title=title,
        sort_order=next_order,
        default_text="Новое сообщение сценария.",
        custom_text=None,
        response_type="none",
        button_options=None,
        send_mode="immediate",
        send_time=None,
        day_offset_workdays=0,
        target_field=None,
        send_employee_card=False,
    )
    db.add(chain_step)
    db.commit()

    return {
        "message": "Шаг цепочки добавлен",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "step_id": chain_step.id,
    }


@router.post("/api/flows/workspace/steps/{step_id}/delete")
def delete_workspace_step_api(
    request: Request,
    step_id: int,
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    step = db.get(FlowStepTemplate, step_id)
    if not step:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Шаг не найден")

    scenario = _get_workspace_scenario_by_flow_key(db, step.flow_key)
    if not scenario:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")

    deleted_kind = _workspace_node_kind(step)
    _delete_step_subtree(db, step)
    db.commit()

    return {
        "message": "Элемент удалён",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "deleted_kind": deleted_kind,
    }


@router.get("/app/flows/workspace")
def scenario_workspace_legacy_redirect(
    request: Request,
    scenario_id: Optional[int] = None,
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    target = f"/app/flows/workspace-v2?scenario_id={scenario_id}" if scenario_id else "/app/flows/workspace-v2"
    return RedirectResponse(url=target, status_code=status.HTTP_303_SEE_OTHER)


@router.post("/api/flows/workspace/steps/{step_id}/attachment")
async def upload_workspace_step_attachment_api(
    request: Request,
    step_id: int,
    upload: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    step = db.get(FlowStepTemplate, step_id)
    if not step:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Шаг не найден")
    scenario = _get_workspace_scenario_by_flow_key(db, step.flow_key)
    if not scenario:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")
    await _save_step_attachment(step, upload)
    db.commit()
    return {
        "message": "Вложение добавлено",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "step_id": step.id,
    }


@router.post("/api/flows/workspace/steps/{step_id}/attachment/delete")
def delete_workspace_step_attachment_api(
    request: Request,
    step_id: int,
    db: Session = Depends(get_db),
):
    require_api_auth(request)
    step = db.get(FlowStepTemplate, step_id)
    if not step:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Шаг не найден")
    scenario = _get_workspace_scenario_by_flow_key(db, step.flow_key)
    if not scenario:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Сценарий не найден")
    _delete_step_attachment_file(step)
    db.commit()
    return {
        "message": "Вложение удалено",
        "payload": _build_scenario_workspace_payload(db, scenario.id, kind=scenario.scenario_kind),
        "step_id": step.id,
    }


@router.get("/app/flows/workspace-v2")
def scenario_workspace_v2_page(
    request: Request,
    scenario_id: Optional[int] = None,
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    return render_template(
        request,
        templates,
        "react_scenario_workspace_v2.html",
        {
            "active_tab": "flows",
            "react_api_url": "/api/flows/workspace",
            "react_selected_scenario_id": scenario_id or "",
            "classic_list_url": "/flows",
            "workspace_kind": "scenario",
        },
    )


@router.get("/app/surveys/workspace")
def survey_workspace_page(
    request: Request,
    scenario_id: Optional[int] = None,
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    return render_template(
        request,
        templates,
        "react_scenario_workspace_v2.html",
        {
            "active_tab": "surveys",
            "react_api_url": "/api/flows/workspace",
            "react_selected_scenario_id": scenario_id or "",
            "classic_list_url": "/surveys",
            "workspace_kind": "survey",
        },
    )


@router.get("/surveys")
def surveys_page(request: Request):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    return RedirectResponse(url="/app/surveys/workspace", status_code=status.HTTP_303_SEE_OTHER)


def _template_entity_meta(kind: str) -> dict[str, str]:
    if kind == "survey":
        return {
            "kind": "survey",
            "active_tab": "surveys",
            "collection_title": "Опросы",
            "collection_title_single": "Опрос",
            "collection_description": "Список всех опросов проекта. Детальные вопросы редактируются на отдельной странице опроса.",
            "create_label": "Создать опрос",
            "new_title": "Новый опрос",
            "collection_path": "/surveys",
            "item_label": "опрос",
            "item_label_cap": "Опрос",
            "edit_title": "Редактировать опрос",
            "back_label": "К списку опросов",
        }
    return {
        "kind": "scenario",
        "active_tab": "flows",
        "collection_title": "Сценарии",
        "collection_title_single": "Сценарий",
        "collection_description": "Список всех сценариев проекта. Детальные шаги редактируются на отдельной странице сценария.",
        "create_label": "Создать сценарий",
        "new_title": "Новый сценарий",
        "collection_path": "/flows",
        "item_label": "сценарий",
        "item_label_cap": "Сценарий",
        "edit_title": "Редактировать сценарий",
        "back_label": "К списку сценариев",
    }


def _template_edit_redirect(
    scenario: ScenarioTemplate,
    flash_message: Optional[str] = None,
    flash_type: str = "success",
    legacy: bool = False,
) -> RedirectResponse:
    meta = _template_entity_meta(getattr(scenario, "scenario_kind", "scenario"))
    url = f"{meta['collection_path']}/{scenario.id}"
    query: dict[str, str | int] = {}
    if legacy:
        query["legacy"] = 1
    if flash_message:
        query["flash_message"] = flash_message
        query["flash_type"] = flash_type
    if query:
        from urllib.parse import urlencode

        url = f"{url}?{urlencode(query)}"
    return RedirectResponse(url=url, status_code=status.HTTP_303_SEE_OTHER)


def _template_workspace_redirect(
    kind: str,
    scenario_id: Optional[int] = None,
    flash_message: Optional[str] = None,
    flash_type: str = "success",
) -> RedirectResponse:
    workspace_path = "/app/surveys/workspace" if kind == "survey" else "/app/flows/workspace-v2"
    query = {}
    if scenario_id:
        query["scenario_id"] = scenario_id
    if flash_message:
        query["flash_message"] = flash_message
        query["flash_type"] = flash_type
    if query:
        from urllib.parse import urlencode

        workspace_path = f"{workspace_path}?{urlencode(query)}"
    return RedirectResponse(url=workspace_path, status_code=status.HTTP_303_SEE_OTHER)


def _template_workspace_redirect_from_request(
    request: Request,
    kind: str,
    scenario_id: Optional[int] = None,
) -> RedirectResponse:
    flash_message = request.query_params.get("flash_message")
    flash_type = request.query_params.get("flash_type", "success")
    return _template_workspace_redirect(kind, scenario_id=scenario_id, flash_message=flash_message, flash_type=flash_type)


def _request_prefers_legacy_editor(request: Request) -> bool:
    return request.query_params.get("legacy") in {"1", "true", "True"}


def _create_template_entity(
    request: Request,
    kind: str,
    title: str,
    role_scope: str,
    employee_scope: str,
    target_employee_id: str,
    trigger_mode: str,
    description: str,
    db: Session,
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    meta = _template_entity_meta(kind)
    last_scenario = (
        db.query(ScenarioTemplate)
        .filter(ScenarioTemplate.scenario_kind == kind)
        .order_by(ScenarioTemplate.sort_order.desc(), ScenarioTemplate.id.desc())
        .first()
    )
    scenario = ScenarioTemplate(
        scenario_key=f"custom_{kind}_{int(utc_now().timestamp())}",
        scenario_kind=kind,
        title=title.strip() or meta["new_title"],
        sort_order=(last_scenario.sort_order + 10) if last_scenario else 10,
        role_scope=role_scope if role_scope in ROLE_SCOPE_LABELS else "all",
        employee_scope=employee_scope if employee_scope in EMPLOYEE_SCOPE_LABELS else "all",
        target_employee_id=int(target_employee_id) if (target_employee_id or "").strip().isdigit() else None,
        trigger_mode=trigger_mode if trigger_mode in TRIGGER_MODE_LABELS else "manual_only",
        description=description.strip() or None,
    )
    db.add(scenario)
    db.commit()
    db.refresh(scenario)
    return _template_workspace_redirect(
        kind,
        scenario.id,
        flash_message=f"{meta['item_label_cap']} создан",
    )


def _edit_template_page(
    request: Request,
    scenario_id: int,
    kind: str,
    db: Session,
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    scenario = db.get(ScenarioTemplate, scenario_id)
    meta = _template_entity_meta(kind)
    if not scenario or scenario.scenario_kind != kind:
        return RedirectResponse(url=meta["collection_path"], status_code=status.HTTP_303_SEE_OTHER)
    editor_data = _load_scenario_editor_data(db, scenario)
    return render_template(
        request,
        templates,
        "scenario_edit.html",
        {
            "active_tab": meta["active_tab"],
            "scenario": scenario,
            "steps": editor_data["steps"],
            "role_scope_labels": ROLE_SCOPE_LABELS,
            "employee_scope_labels": EMPLOYEE_SCOPE_LABELS,
            "trigger_mode_labels": TRIGGER_MODE_LABELS,
            "response_type_labels": RESPONSE_TYPE_LABELS,
            "send_mode_labels": SEND_MODE_LABELS,
            "target_field_labels": TARGET_FIELD_LABELS,
            "notification_recipient_scope_labels": NOTIFICATION_RECIPIENT_SCOPE_LABELS,
            "branch_steps_by_parent": editor_data["branch_steps_by_parent"],
            "chain_steps_by_parent": editor_data["chain_steps_by_parent"],
            "button_notifications_by_step": editor_data["button_notifications_by_step"],
            "available_scenarios": editor_data["available_scenarios"],
            "employee_options": editor_data["employee_options"],
            "document_tag_titles": editor_data["document_tag_titles"],
            "collection_path": meta["collection_path"],
            "collection_title": meta["collection_title"],
            "collection_title_single": meta["collection_title_single"],
            "edit_title": meta["edit_title"],
            "item_label_cap": meta["item_label_cap"],
            "back_label": meta["back_label"],
            "kind": meta["kind"],
            "flash_message": request.query_params.get("flash_message"),
            "flash_type": request.query_params.get("flash_type", "success"),
        },
    )


@router.post("/flows")
def create_scenario(
    request: Request,
    title: str = Form("Новый сценарий"),
    role_scope: str = Form("all"),
    employee_scope: str = Form("all"),
    target_employee_id: str = Form(""),
    trigger_mode: str = Form("manual_only"),
    description: str = Form(""),
    db: Session = Depends(get_db),
):
    return _create_template_entity(request, "scenario", title, role_scope, employee_scope, target_employee_id, trigger_mode, description, db)


@router.post("/surveys")
def create_survey(
    request: Request,
    title: str = Form("Новый опрос"),
    role_scope: str = Form("all"),
    employee_scope: str = Form("all"),
    target_employee_id: str = Form(""),
    trigger_mode: str = Form("manual_only"),
    description: str = Form(""),
    db: Session = Depends(get_db),
):
    return _create_template_entity(request, "survey", title, role_scope, employee_scope, target_employee_id, trigger_mode, description, db)


@router.get("/flows/{scenario_id}")
def edit_scenario_page(
    request: Request,
    scenario_id: int,
    legacy: bool = False,
    db: Session = Depends(get_db),
):
    if not legacy:
        return _template_workspace_redirect_from_request(request, "scenario", scenario_id)
    return _edit_template_page(request, scenario_id, "scenario", db)


@router.get("/surveys/{scenario_id}")
def edit_survey_page(
    request: Request,
    scenario_id: int,
    legacy: bool = False,
    db: Session = Depends(get_db),
):
    if not legacy:
        return _template_workspace_redirect_from_request(request, "survey", scenario_id)
    return _edit_template_page(request, scenario_id, "survey", db)


@router.get("/flows/steps/{step_id}/attachment")
def download_step_attachment(
    request: Request,
    step_id: int,
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    step = db.get(FlowStepTemplate, step_id)
    attachment_path = (getattr(step, "attachment_path", None) or "").strip() if step else ""
    if not step or not attachment_path:
        return RedirectResponse(url="/flows", status_code=status.HTTP_303_SEE_OTHER)
    path = Path(attachment_path)
    if not path.exists():
        return RedirectResponse(url="/flows", status_code=status.HTTP_303_SEE_OTHER)
    return FileResponse(path, filename=getattr(step, "attachment_filename", None) or path.name)


@router.post("/flows/steps/{step_id}/attachment/delete")
def delete_step_attachment(
    request: Request,
    step_id: int,
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    step = db.get(FlowStepTemplate, step_id)
    if not step:
        return RedirectResponse(url="/flows", status_code=status.HTTP_303_SEE_OTHER)
    scenario = db.query(ScenarioTemplate).filter(ScenarioTemplate.scenario_key == step.flow_key).first()
    _delete_step_attachment_file(step)
    db.commit()
    if scenario:
        return _template_edit_redirect(scenario, legacy=_request_prefers_legacy_editor(request))
    return RedirectResponse(url="/flows", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/flows/{scenario_id}")
@router.post("/surveys/{scenario_id}")
async def update_scenario(
    request: Request,
    scenario_id: int,
    title: str = Form(...),
    role_scope: str = Form("all"),
    employee_scope: str = Form("all"),
    target_employee_id: str = Form(""),
    trigger_mode: str = Form("manual_only"),
    description: str = Form(""),
    action: str = Form("save"),
    target_step_id: str = Form(""),
    step_id: Optional[List[int]] = Form(None),
    step_title: Optional[List[str]] = Form(None),
    custom_text: Optional[List[str]] = Form(None),
    response_type: Optional[List[str]] = Form(None),
    button_options: Optional[List[str]] = Form(None),
    send_mode: Optional[List[str]] = Form(None),
    send_time: Optional[List[str]] = Form(None),
    day_offset_workdays: Optional[List[int]] = Form(None),
    target_field: Optional[List[str]] = Form(None),
    send_employee_card: Optional[List[str]] = Form(None),
    notify_on_send_text: Optional[List[str]] = Form(None),
    notify_on_send_recipient_ids: Optional[List[str]] = Form(None),
    notify_on_send_recipient_scope: Optional[List[str]] = Form(None),
    remove_attachment_step_id: Optional[List[int]] = Form(None),
    branch_parent_step_id: Optional[List[str]] = Form(None),
    branch_parent_step_ref: Optional[List[str]] = Form(None),
    branch_option_index: Optional[List[str]] = Form(None),
    branch_step_id: Optional[List[str]] = Form(None),
    branch_step_title: Optional[List[str]] = Form(None),
    branch_custom_text: Optional[List[str]] = Form(None),
    branch_response_type: Optional[List[str]] = Form(None),
    branch_button_options: Optional[List[str]] = Form(None),
    branch_launch_scenario_key: Optional[List[str]] = Form(None),
    branch_send_employee_card: Optional[List[str]] = Form(None),
    branch_notify_on_send_text: Optional[List[str]] = Form(None),
    branch_notify_on_send_recipient_ids: Optional[List[str]] = Form(None),
    branch_notify_on_send_recipient_scope: Optional[List[str]] = Form(None),
    branch_button_notification_text: Optional[List[str]] = Form(None),
    branch_button_notification_recipient_ids: Optional[List[str]] = Form(None),
    branch_button_notification_recipient_scope: Optional[List[str]] = Form(None),
    button_notification_parent_step_id: Optional[List[str]] = Form(None),
    button_notification_parent_step_ref: Optional[List[str]] = Form(None),
    button_notification_option_index: Optional[List[str]] = Form(None),
    button_notification_text: Optional[List[str]] = Form(None),
    button_notification_recipient_ids: Optional[List[str]] = Form(None),
    button_notification_recipient_scope: Optional[List[str]] = Form(None),
    branch_remove_attachment_key: Optional[List[str]] = Form(None),
    chain_parent_step_id: Optional[List[str]] = Form(None),
    chain_branch_option_index: Optional[List[str]] = Form(None),
    chain_step_id: Optional[List[str]] = Form(None),
    chain_row_ref: Optional[List[str]] = Form(None),
    chain_step_title: Optional[List[str]] = Form(None),
    chain_custom_text: Optional[List[str]] = Form(None),
    chain_response_type: Optional[List[str]] = Form(None),
    chain_button_options: Optional[List[str]] = Form(None),
    chain_send_mode: Optional[List[str]] = Form(None),
    chain_send_time: Optional[List[str]] = Form(None),
    chain_target_field: Optional[List[str]] = Form(None),
    chain_send_employee_card: Optional[List[str]] = Form(None),
    chain_notify_on_send_text: Optional[List[str]] = Form(None),
    chain_notify_on_send_recipient_ids: Optional[List[str]] = Form(None),
    chain_notify_on_send_recipient_scope: Optional[List[str]] = Form(None),
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    scenario = db.get(ScenarioTemplate, scenario_id)
    if scenario:
        request_form = await request.form()

        def form_list(name: str) -> list[str]:
            return [str(value) for value in request_form.getlist(name)]

        target_step_id_int = int(target_step_id) if (target_step_id or "").strip().isdigit() else None
        scenario.title = title.strip() or scenario.title
        scenario.role_scope = role_scope if role_scope in ROLE_SCOPE_LABELS else "all"
        scenario.employee_scope = employee_scope if employee_scope in EMPLOYEE_SCOPE_LABELS else "all"
        scenario.target_employee_id = int(target_employee_id) if (target_employee_id or "").strip().isdigit() else None
        scenario.trigger_mode = "manual_only" if scenario.scenario_kind == "survey" else (trigger_mode if trigger_mode in TRIGGER_MODE_LABELS else "manual_only")
        scenario.description = description.strip() or None
        step_ids = step_id or []
        step_titles = step_title or []
        custom_texts = custom_text or []
        response_types = response_type or []
        button_values = button_options or []
        send_modes = send_mode or []
        send_times = send_time or []
        day_offsets = day_offset_workdays or []
        target_fields = target_field or []
        send_employee_card_values = send_employee_card or []
        notify_on_send_text_values = notify_on_send_text or []
        notify_on_send_recipient_ids_values = notify_on_send_recipient_ids or []
        notify_on_send_recipient_scope_values = notify_on_send_recipient_scope or []
        removed_attachment_step_ids = set(remove_attachment_step_id or [])
        branch_parent_ids = form_list("branch_parent_step_id")
        branch_parent_refs = form_list("branch_parent_step_ref")
        branch_option_indexes = form_list("branch_option_index")
        branch_step_ids = form_list("branch_step_id")
        branch_step_titles = form_list("branch_step_title")
        branch_custom_texts = form_list("branch_custom_text")
        branch_response_types = form_list("branch_response_type")
        branch_button_values = form_list("branch_button_options")
        branch_launch_scenario_keys = form_list("branch_launch_scenario_key")
        branch_send_employee_card_values = form_list("branch_send_employee_card")
        branch_notify_on_send_text_values = form_list("branch_notify_on_send_text")
        branch_notify_on_send_recipient_ids_values = form_list("branch_notify_on_send_recipient_ids")
        branch_notify_on_send_recipient_scope_values = form_list("branch_notify_on_send_recipient_scope")
        branch_button_notification_texts = form_list("branch_button_notification_text")
        branch_button_notification_recipient_ids_values = form_list("branch_button_notification_recipient_ids")
        branch_button_notification_recipient_scope_values = form_list("branch_button_notification_recipient_scope")
        button_notification_parent_ids = form_list("button_notification_parent_step_id")
        button_notification_parent_refs = form_list("button_notification_parent_step_ref")
        button_notification_option_indexes = form_list("button_notification_option_index")
        button_notification_texts = form_list("button_notification_text")
        button_notification_recipient_ids_values = form_list("button_notification_recipient_ids")
        button_notification_recipient_scope_values = form_list("button_notification_recipient_scope")
        removed_branch_attachment_keys = set(form_list("branch_remove_attachment_key"))
        chain_parent_ids = form_list("chain_parent_step_id")
        chain_branch_option_indexes = form_list("chain_branch_option_index")
        chain_step_ids = form_list("chain_step_id")
        chain_row_refs = form_list("chain_row_ref")
        chain_step_titles = form_list("chain_step_title")
        chain_custom_texts = form_list("chain_custom_text")
        chain_response_types = form_list("chain_response_type")
        chain_button_values = form_list("chain_button_options")
        chain_send_modes = form_list("chain_send_mode")
        chain_send_times = form_list("chain_send_time")
        chain_target_fields = form_list("chain_target_field")
        chain_send_employee_card_values = form_list("chain_send_employee_card")
        chain_notify_on_send_text_values = form_list("chain_notify_on_send_text")
        chain_notify_on_send_recipient_ids_values = form_list("chain_notify_on_send_recipient_ids")
        chain_notify_on_send_recipient_scope_values = form_list("chain_notify_on_send_recipient_scope")

        for index, current_step_id in enumerate(step_ids):
            step = db.get(FlowStepTemplate, current_step_id)
            if not step or step.flow_key != scenario.scenario_key:
                continue
            step.sort_order = (index + 1) * 10
            if index < len(step_titles):
                step.step_title = step_titles[index].strip() or step.step_title
            if index < len(custom_texts):
                step.custom_text = custom_texts[index].strip()
            if index < len(response_types):
                current_response_type = response_types[index]
                step.response_type = current_response_type if current_response_type in {"none", "text", "file", "buttons", "branching"} else "none"
            if index < len(button_values):
                step.button_options = button_values[index].strip() or None
            if index < len(send_modes):
                current_send_mode = send_modes[index]
                step.send_mode = current_send_mode if current_send_mode in {"immediate", "specific_time"} else "immediate"
            if index < len(send_times):
                step.send_time = send_times[index].strip() or None
            if index < len(day_offsets):
                step.day_offset_workdays = int(day_offsets[index])
            if index < len(target_fields):
                step.target_field = target_fields[index].strip() or None
            if index < len(send_employee_card_values):
                step.send_employee_card = send_employee_card_values[index] == "true"
            if index < len(notify_on_send_text_values):
                step.notify_on_send_text = notify_on_send_text_values[index].strip() or None
            if index < len(notify_on_send_recipient_ids_values):
                step.notify_on_send_recipient_ids = notify_on_send_recipient_ids_values[index].strip() or None
            if index < len(notify_on_send_recipient_scope_values):
                step.notify_on_send_recipient_scope = _normalize_notification_scope(notify_on_send_recipient_scope_values[index])
            if scenario.scenario_kind == "survey":
                step.send_mode = "immediate"
                step.send_time = None
                step.day_offset_workdays = 0
                step.target_field = None
                step.send_employee_card = False
            if step.response_type == "buttons":
                options = [item.strip() for item in (step.button_options or "").splitlines() if item.strip()]
                preserved_option_indexes: set[int] = set()
                for option_idx, _ in enumerate(options):
                    payload = submitted_button_notification_rows.get((step.id, option_idx), {})
                    _sync_button_notification(
                        db,
                        step,
                        option_idx,
                        str(payload.get("text") or ""),
                        str(payload.get("recipient_ids") or ""),
                        str(payload.get("recipient_scope") or ""),
                    )
                    preserved_option_indexes.add(option_idx)
                for notification in db.query(StepButtonNotification).filter(StepButtonNotification.step_id == step.id).all():
                    if notification.option_index not in preserved_option_indexes:
                        db.delete(notification)
            else:
                db.query(StepButtonNotification).filter(StepButtonNotification.step_id == step.id).delete()
            if step.id in removed_attachment_step_ids:
                _delete_step_attachment_file(step)
            upload = request_form.get(f"step_attachment_{step.id}")
            if upload is not None and getattr(upload, "filename", ""):
                await _save_step_attachment(step, upload)

        if action == "delete_step" and target_step_id_int is not None:
            step = db.get(FlowStepTemplate, target_step_id_int)
            if step and step.flow_key == scenario.scenario_key:
                _delete_step_tree(db, step)
        elif action == "reset_step" and target_step_id_int is not None:
            step = db.get(FlowStepTemplate, target_step_id_int)
            if step and step.flow_key == scenario.scenario_key:
                step.custom_text = None
                step.button_options = None
        elif action == "add_step":
            last_step = (
                db.query(FlowStepTemplate)
                .filter(
                    FlowStepTemplate.flow_key == scenario.scenario_key,
                    FlowStepTemplate.parent_step_id.is_(None),
                )
                .order_by(FlowStepTemplate.sort_order.desc(), FlowStepTemplate.id.desc())
                .first()
            )
            next_order = (last_step.sort_order + 10) if last_step else 10
            db.add(
                FlowStepTemplate(
                    flow_key=scenario.scenario_key,
                    step_key=f"{scenario.scenario_key}_step_{int(utc_now().timestamp())}",
                    step_title="Новый вопрос" if scenario.scenario_kind == "survey" else "Новый шаг",
                    sort_order=next_order,
                    default_text="Новое сообщение опроса." if scenario.scenario_kind == "survey" else "Новое сообщение сценария.",
                    custom_text=None,
                    response_type="none",
                    button_options=None,
                    send_mode="immediate",
                    send_time=None,
                    day_offset_workdays=0,
                    target_field=None,
                    send_employee_card=False,
                )
            )

        submitted_branch_rows = {}
        submitted_branch_rows_by_ref = {}
        for index, parent_id in enumerate(branch_parent_ids):
            parent_id_value = int(parent_id) if str(parent_id).strip().isdigit() else None
            parent_ref_value = branch_parent_refs[index].strip() if index < len(branch_parent_refs) else ""
            option_idx_raw = branch_option_indexes[index] if index < len(branch_option_indexes) else None
            option_idx = int(option_idx_raw) if str(option_idx_raw).strip().isdigit() else None
            branch_step_id_raw = branch_step_ids[index] if index < len(branch_step_ids) else None
            branch_step_id_value = int(branch_step_id_raw) if str(branch_step_id_raw).strip().isdigit() else None
            if (parent_id_value is None and not parent_ref_value) or option_idx is None:
                continue
            payload = {
                "branch_step_id": branch_step_id_value,
                "title": branch_step_titles[index] if index < len(branch_step_titles) else "",
                "custom_text": branch_custom_texts[index] if index < len(branch_custom_texts) else "",
                "response_type": branch_response_types[index] if index < len(branch_response_types) else "none",
                "button_options": branch_button_values[index] if index < len(branch_button_values) else "",
                "launch_scenario_key": branch_launch_scenario_keys[index] if index < len(branch_launch_scenario_keys) else "",
                "send_employee_card": branch_send_employee_card_values[index] if index < len(branch_send_employee_card_values) else "false",
                "notify_on_send_text": branch_notify_on_send_text_values[index] if index < len(branch_notify_on_send_text_values) else "",
                "notify_on_send_recipient_ids": branch_notify_on_send_recipient_ids_values[index] if index < len(branch_notify_on_send_recipient_ids_values) else "",
                "notify_on_send_recipient_scope": branch_notify_on_send_recipient_scope_values[index] if index < len(branch_notify_on_send_recipient_scope_values) else "",
                "button_notification_text": branch_button_notification_texts[index] if index < len(branch_button_notification_texts) else "",
                "button_notification_recipient_ids": branch_button_notification_recipient_ids_values[index] if index < len(branch_button_notification_recipient_ids_values) else "",
                "button_notification_recipient_scope": branch_button_notification_recipient_scope_values[index] if index < len(branch_button_notification_recipient_scope_values) else "",
            }
            if parent_id_value is not None:
                submitted_branch_rows[(parent_id_value, option_idx)] = payload
            if parent_ref_value:
                submitted_branch_rows_by_ref[(parent_ref_value, option_idx)] = payload

        submitted_chain_rows: dict[tuple[int, int], list[dict[str, object]]] = defaultdict(list)
        for index, parent_id in enumerate(chain_parent_ids):
            parent_id_value = int(parent_id) if str(parent_id).strip().isdigit() else None
            branch_option_raw = chain_branch_option_indexes[index] if index < len(chain_branch_option_indexes) else None
            branch_option_value = int(branch_option_raw) if str(branch_option_raw).strip().isdigit() else None
            chain_step_id_raw = chain_step_ids[index] if index < len(chain_step_ids) else None
            chain_step_id_value = int(chain_step_id_raw) if str(chain_step_id_raw).strip().isdigit() else None
            if parent_id_value is None or branch_option_value is None:
                continue
            submitted_chain_rows[(parent_id_value, branch_option_value)].append(
                {
                    "chain_step_id": chain_step_id_value,
                    "row_ref": chain_row_refs[index] if index < len(chain_row_refs) else "",
                    "title": chain_step_titles[index] if index < len(chain_step_titles) else "",
                    "custom_text": chain_custom_texts[index] if index < len(chain_custom_texts) else "",
                    "response_type": chain_response_types[index] if index < len(chain_response_types) else "none",
                    "button_options": chain_button_values[index] if index < len(chain_button_values) else "",
                    "send_mode": chain_send_modes[index] if index < len(chain_send_modes) else "immediate",
                    "send_time": chain_send_times[index] if index < len(chain_send_times) else "",
                    "target_field": chain_target_fields[index] if index < len(chain_target_fields) else "",
                    "send_employee_card": chain_send_employee_card_values[index] if index < len(chain_send_employee_card_values) else "false",
                    "notify_on_send_text": chain_notify_on_send_text_values[index] if index < len(chain_notify_on_send_text_values) else "",
                    "notify_on_send_recipient_ids": chain_notify_on_send_recipient_ids_values[index] if index < len(chain_notify_on_send_recipient_ids_values) else "",
                    "notify_on_send_recipient_scope": chain_notify_on_send_recipient_scope_values[index] if index < len(chain_notify_on_send_recipient_scope_values) else "",
                    "row_index": len(submitted_chain_rows[(parent_id_value, branch_option_value)]),
                }
            )

        submitted_button_notification_rows = {}
        submitted_button_notification_rows_by_ref = {}
        for index, parent_id in enumerate(button_notification_parent_ids):
            parent_id_value = int(parent_id) if str(parent_id).strip().isdigit() else None
            parent_ref_value = button_notification_parent_refs[index].strip() if index < len(button_notification_parent_refs) else ""
            option_idx_raw = button_notification_option_indexes[index] if index < len(button_notification_option_indexes) else ""
            option_idx = int(option_idx_raw) if str(option_idx_raw).strip().isdigit() else None
            if (parent_id_value is None and not parent_ref_value) or option_idx is None:
                continue
            payload = {
                "text": button_notification_texts[index] if index < len(button_notification_texts) else "",
                "recipient_ids": button_notification_recipient_ids_values[index] if index < len(button_notification_recipient_ids_values) else "",
                "recipient_scope": button_notification_recipient_scope_values[index] if index < len(button_notification_recipient_scope_values) else "",
            }
            if parent_id_value is not None:
                submitted_button_notification_rows[(parent_id_value, option_idx)] = payload
            if parent_ref_value:
                submitted_button_notification_rows_by_ref[(parent_ref_value, option_idx)] = payload

        chain_step_id_by_ref: dict[str, int] = {}

        top_level_steps = (
            db.query(FlowStepTemplate)
            .filter(
                FlowStepTemplate.flow_key == scenario.scenario_key,
                FlowStepTemplate.parent_step_id.is_(None),
            )
            .order_by(FlowStepTemplate.sort_order, FlowStepTemplate.id)
            .all()
        )
        for step in top_level_steps:
            existing_children = {
                child.branch_option_index: child
                for child in db.query(FlowStepTemplate)
                .filter(FlowStepTemplate.parent_step_id == step.id)
                .all()
            }
            if step.response_type != "branching":
                for child in existing_children.values():
                    _delete_step_tree(db, child)
                continue

            options = [item.strip() for item in (step.button_options or "").splitlines() if item.strip()]
            for option_idx, option_label in enumerate(options):
                payload = submitted_branch_rows.get((step.id, option_idx), {})
                branch_step = None
                branch_step_id_value = payload.get("branch_step_id")
                if branch_step_id_value is not None:
                    branch_step = db.get(FlowStepTemplate, branch_step_id_value)
                    if branch_step and branch_step.parent_step_id != step.id:
                        branch_step = None
                if branch_step is None:
                    branch_step = existing_children.pop(option_idx, None)
                else:
                    existing_children.pop(option_idx, None)
                if not branch_step:
                    branch_step = FlowStepTemplate(
                        flow_key=scenario.scenario_key,
                        step_key=f"{step.step_key}__branch_{option_idx}",
                        parent_step_id=step.id,
                        branch_option_index=option_idx,
                        step_title=f"Ветка: {option_label}",
                        sort_order=step.sort_order * 100 + option_idx + 1,
                        default_text=f"Сообщение для варианта \"{option_label}\".",
                        custom_text=None,
                        response_type="none",
                        button_options=None,
                        send_mode="immediate",
                        send_time=None,
                        day_offset_workdays=0,
                        target_field=None,
                        launch_scenario_key=None,
                        send_employee_card=False,
                    )
                    db.add(branch_step)
                branch_step.flow_key = scenario.scenario_key
                branch_step.parent_step_id = step.id
                branch_step.branch_option_index = option_idx
                branch_step.sort_order = step.sort_order * 100 + option_idx + 1
                branch_step.step_title = (payload.get("title") or "").strip() or f"Ветка: {option_label}"
                branch_step.custom_text = (payload.get("custom_text") or "").strip()
                current_branch_response_type = (payload.get("response_type") or "none").strip()
                branch_step.response_type = (
                    current_branch_response_type
                    if current_branch_response_type in {"none", "text", "file", "buttons", "chain", "launch_scenario"}
                    else "none"
                )
                branch_step.button_options = (
                    (payload.get("button_options") or "").strip() or None
                    if branch_step.response_type == "buttons"
                    else None
                )
                branch_step.launch_scenario_key = (
                    (payload.get("launch_scenario_key") or "").strip() or None
                    if branch_step.response_type == "launch_scenario"
                    else None
                )
                branch_attachment_key = f"{step.id}:{option_idx}"
                if branch_attachment_key in removed_branch_attachment_keys:
                    _delete_step_attachment_file(branch_step)
                branch_upload = request_form.get(f"branch_attachment_{step.id}_{option_idx}")
                if branch_upload is not None and getattr(branch_upload, "filename", ""):
                    await _save_step_attachment(branch_step, branch_upload)
                branch_step.send_mode = "immediate"
                branch_step.send_time = None
                branch_step.day_offset_workdays = 0
                branch_step.target_field = None
                branch_step.send_employee_card = str(payload.get("send_employee_card") or "false").strip() == "true"
                branch_step.notify_on_send_text = str(payload.get("notify_on_send_text") or "").strip() or None
                branch_step.notify_on_send_recipient_ids = str(payload.get("notify_on_send_recipient_ids") or "").strip() or None
                branch_step.notify_on_send_recipient_scope = _normalize_notification_scope(str(payload.get("notify_on_send_recipient_scope") or ""))
                _sync_button_notification(
                    db,
                    step,
                    option_idx,
                    str(payload.get("button_notification_text") or ""),
                    str(payload.get("button_notification_recipient_ids") or ""),
                    str(payload.get("button_notification_recipient_scope") or ""),
                )
                db.flush()

                existing_chain_children = {
                    child.id: child
                    for child in db.query(FlowStepTemplate)
                    .filter(
                        FlowStepTemplate.parent_step_id == branch_step.id,
                        FlowStepTemplate.branch_option_index.is_(None),
                    )
                    .all()
                }
                chain_payloads = submitted_chain_rows.get((step.id, option_idx), [])
                if branch_step.response_type == "chain":
                    preserved_chain_ids: set[int] = set()
                    for chain_index, chain_payload in enumerate(chain_payloads):
                        chain_step = None
                        chain_step_id_value = chain_payload.get("chain_step_id")
                        is_new_chain_step = not isinstance(chain_step_id_value, int)
                        if isinstance(chain_step_id_value, int):
                            chain_step = existing_chain_children.get(chain_step_id_value)
                        if chain_step is None:
                            chain_step = FlowStepTemplate(
                                flow_key=scenario.scenario_key,
                                step_key=f"{branch_step.step_key}__chain_{chain_index}",
                                parent_step_id=branch_step.id,
                                branch_option_index=None,
                                step_title=f"Шаг {chain_index + 1}",
                                sort_order=(chain_index + 1) * 10,
                                default_text="Новое сообщение сценария.",
                                custom_text=None,
                                response_type="none",
                                button_options=None,
                                send_mode="immediate",
                                send_time=None,
                                day_offset_workdays=0,
                                target_field=None,
                                launch_scenario_key=None,
                                send_employee_card=False,
                            )
                            db.add(chain_step)
                            db.flush()
                        chain_step.flow_key = scenario.scenario_key
                        chain_step.parent_step_id = branch_step.id
                        chain_step.branch_option_index = None
                        chain_step.step_key = f"{branch_step.step_key}__chain_{chain_index}"
                        chain_step.sort_order = (chain_index + 1) * 10
                        chain_step.step_title = (str(chain_payload.get("title") or "").strip() or f"Шаг {chain_index + 1}")
                        chain_text_value = str(chain_payload.get("custom_text") or "").strip()
                        chain_step.custom_text = None if is_new_chain_step and not chain_text_value else chain_text_value
                        chain_response_type_value = str(chain_payload.get("response_type") or "none").strip()
                        chain_step.response_type = chain_response_type_value if chain_response_type_value in {"none", "text", "file", "buttons", "branching"} else "none"
                        chain_step.button_options = (
                            str(chain_payload.get("button_options") or "").strip() or None
                            if chain_step.response_type in {"buttons", "branching"}
                            else None
                        )
                        chain_step.launch_scenario_key = None
                        chain_send_mode_value = str(chain_payload.get("send_mode") or "immediate").strip()
                        chain_step.send_mode = chain_send_mode_value if chain_send_mode_value in {"immediate", "specific_time"} else "immediate"
                        chain_step.send_time = (
                            str(chain_payload.get("send_time") or "").strip() or None
                            if chain_step.send_mode == "specific_time"
                            else None
                        )
                        chain_step.day_offset_workdays = 0
                        chain_target_field_value = str(chain_payload.get("target_field") or "").strip()
                        chain_step.target_field = chain_target_field_value if chain_target_field_value in TARGET_FIELD_LABELS else None
                        chain_step.send_employee_card = str(chain_payload.get("send_employee_card") or "false").strip() == "true"
                        chain_step.notify_on_send_text = str(chain_payload.get("notify_on_send_text") or "").strip() or None
                        chain_step.notify_on_send_recipient_ids = str(chain_payload.get("notify_on_send_recipient_ids") or "").strip() or None
                        chain_step.notify_on_send_recipient_scope = _normalize_notification_scope(str(chain_payload.get("notify_on_send_recipient_scope") or ""))
                        chain_row_ref_value = str(chain_payload.get("row_ref") or "").strip()
                        if chain_row_ref_value:
                            chain_step_id_by_ref[chain_row_ref_value] = chain_step.id
                        preserved_chain_ids.add(chain_step.id)
                        chain_upload = request_form.get(f"chain_attachment_{step.id}_{option_idx}_{chain_index}")
                        if chain_upload is not None and getattr(chain_upload, "filename", ""):
                            await _save_step_attachment(chain_step, chain_upload)
                        if chain_step.response_type == "buttons":
                            child_options = [item.strip() for item in (chain_step.button_options or "").splitlines() if item.strip()]
                            preserved_child_option_indexes: set[int] = set()
                            for child_option_idx, _ in enumerate(child_options):
                                payload_by_id = submitted_button_notification_rows.get((chain_step.id, child_option_idx), {})
                                payload_by_ref = submitted_button_notification_rows_by_ref.get((chain_row_ref_value, child_option_idx), {}) if chain_row_ref_value else {}
                                button_payload = payload_by_id or payload_by_ref
                                _sync_button_notification(
                                    db,
                                    chain_step,
                                    child_option_idx,
                                    str(button_payload.get("text") or ""),
                                    str(button_payload.get("recipient_ids") or ""),
                                    str(button_payload.get("recipient_scope") or ""),
                                )
                                preserved_child_option_indexes.add(child_option_idx)
                            for notification in db.query(StepButtonNotification).filter(StepButtonNotification.step_id == chain_step.id).all():
                                if notification.option_index not in preserved_child_option_indexes:
                                    db.delete(notification)
                        else:
                            db.query(StepButtonNotification).filter(StepButtonNotification.step_id == chain_step.id).delete()
                    for existing_id, existing_child in existing_chain_children.items():
                        if existing_id not in preserved_chain_ids:
                            _delete_step_tree(db, existing_child)
                else:
                    db.query(StepButtonNotification).filter(StepButtonNotification.step_id == branch_step.id).delete()
                    for existing_child in existing_chain_children.values():
                        _delete_step_tree(db, existing_child)
            for child in existing_children.values():
                _delete_step_tree(db, child)

        chain_parent_steps = (
            db.query(FlowStepTemplate)
            .filter(
                FlowStepTemplate.flow_key == scenario.scenario_key,
                FlowStepTemplate.parent_step_id.is_not(None),
                FlowStepTemplate.branch_option_index.is_(None),
            )
            .all()
        )
        for chain_parent_step in chain_parent_steps:
            chain_parent_ref_candidates = [f"existing:{chain_parent_step.id}"]
            chain_parent_ref_candidates.extend(
                ref_value for ref_value, mapped_id in chain_step_id_by_ref.items() if mapped_id == chain_parent_step.id
            )
            if chain_parent_step.response_type == "buttons":
                options = [item.strip() for item in (chain_parent_step.button_options or "").splitlines() if item.strip()]
                preserved_option_indexes: set[int] = set()
                for option_idx, _ in enumerate(options):
                    payload = submitted_button_notification_rows.get((chain_parent_step.id, option_idx), {})
                    if not payload:
                        for parent_ref_candidate in chain_parent_ref_candidates:
                            payload = submitted_button_notification_rows_by_ref.get((parent_ref_candidate, option_idx), {})
                            if payload:
                                break
                    _sync_button_notification(
                        db,
                        chain_parent_step,
                        option_idx,
                        str(payload.get("text") or ""),
                        str(payload.get("recipient_ids") or ""),
                        str(payload.get("recipient_scope") or ""),
                    )
                    preserved_option_indexes.add(option_idx)
                for notification in db.query(StepButtonNotification).filter(StepButtonNotification.step_id == chain_parent_step.id).all():
                    if notification.option_index not in preserved_option_indexes:
                        db.delete(notification)
            elif chain_parent_step.response_type != "branching":
                db.query(StepButtonNotification).filter(StepButtonNotification.step_id == chain_parent_step.id).delete()
            existing_children = {
                child.branch_option_index: child
                for child in db.query(FlowStepTemplate)
                .filter(
                    FlowStepTemplate.parent_step_id == chain_parent_step.id,
                    FlowStepTemplate.branch_option_index.is_not(None),
                )
                .all()
            }
            if chain_parent_step.response_type != "branching":
                for child in existing_children.values():
                    _delete_step_tree(db, child)
                continue

            options = [item.strip() for item in (chain_parent_step.button_options or "").splitlines() if item.strip()]
            for option_idx, option_label in enumerate(options):
                payload = submitted_branch_rows.get((chain_parent_step.id, option_idx), {})
                if not payload:
                    for parent_ref_candidate in chain_parent_ref_candidates:
                        payload = submitted_branch_rows_by_ref.get((parent_ref_candidate, option_idx), {})
                        if payload:
                            break
                branch_step = None
                branch_step_id_value = payload.get("branch_step_id")
                if branch_step_id_value is not None:
                    branch_step = db.get(FlowStepTemplate, branch_step_id_value)
                    if branch_step and branch_step.parent_step_id != chain_parent_step.id:
                        branch_step = None
                if branch_step is None:
                    branch_step = existing_children.pop(option_idx, None)
                else:
                    existing_children.pop(option_idx, None)
                if not branch_step:
                    branch_step = FlowStepTemplate(
                        flow_key=scenario.scenario_key,
                        step_key=f"{chain_parent_step.step_key}__branch_{option_idx}",
                        parent_step_id=chain_parent_step.id,
                        branch_option_index=option_idx,
                        step_title=f"Ветка: {option_label}",
                        sort_order=chain_parent_step.sort_order * 100 + option_idx + 1,
                        default_text=f"Сообщение для варианта \"{option_label}\".",
                        custom_text=None,
                        response_type="none",
                        button_options=None,
                        send_mode="immediate",
                        send_time=None,
                        day_offset_workdays=0,
                        target_field=None,
                        launch_scenario_key=None,
                        send_employee_card=False,
                    )
                    db.add(branch_step)
                branch_step.flow_key = scenario.scenario_key
                branch_step.parent_step_id = chain_parent_step.id
                branch_step.branch_option_index = option_idx
                branch_step.sort_order = chain_parent_step.sort_order * 100 + option_idx + 1
                branch_step.step_title = (payload.get("title") or "").strip() or f"Ветка: {option_label}"
                branch_step.custom_text = (payload.get("custom_text") or "").strip()
                current_branch_response_type = (payload.get("response_type") or "none").strip()
                branch_step.response_type = (
                    current_branch_response_type
                    if current_branch_response_type in {"none", "text", "file", "buttons", "chain", "launch_scenario"}
                    else "none"
                )
                branch_step.button_options = (
                    (payload.get("button_options") or "").strip() or None
                    if branch_step.response_type == "buttons"
                    else None
                )
                branch_step.launch_scenario_key = (
                    (payload.get("launch_scenario_key") or "").strip() or None
                    if branch_step.response_type == "launch_scenario"
                    else None
                )
                branch_attachment_key = f"{chain_parent_step.id}:{option_idx}"
                if branch_attachment_key in removed_branch_attachment_keys:
                    _delete_step_attachment_file(branch_step)
                branch_upload = request_form.get(f"branch_attachment_{chain_parent_step.id}_{option_idx}")
                if branch_upload is not None and getattr(branch_upload, "filename", ""):
                    await _save_step_attachment(branch_step, branch_upload)
                branch_step.send_mode = "immediate"
                branch_step.send_time = None
                branch_step.day_offset_workdays = 0
                branch_step.target_field = None
                branch_step.send_employee_card = str(payload.get("send_employee_card") or "false").strip() == "true"
                branch_step.notify_on_send_text = str(payload.get("notify_on_send_text") or "").strip() or None
                branch_step.notify_on_send_recipient_ids = str(payload.get("notify_on_send_recipient_ids") or "").strip() or None
                branch_step.notify_on_send_recipient_scope = _normalize_notification_scope(str(payload.get("notify_on_send_recipient_scope") or ""))
                _sync_button_notification(
                    db,
                    chain_parent_step,
                    option_idx,
                    str(payload.get("button_notification_text") or ""),
                    str(payload.get("button_notification_recipient_ids") or ""),
                    str(payload.get("button_notification_recipient_scope") or ""),
                )

                db.flush()

                existing_chain_children = {
                    child.id: child
                    for child in db.query(FlowStepTemplate)
                    .filter(
                        FlowStepTemplate.parent_step_id == branch_step.id,
                        FlowStepTemplate.branch_option_index.is_(None),
                    )
                    .all()
                }
                chain_payloads = submitted_chain_rows.get((chain_parent_step.id, option_idx), [])
                if branch_step.response_type == "chain":
                    preserved_chain_ids: set[int] = set()
                    for chain_index, chain_payload in enumerate(chain_payloads):
                        child_chain_step = None
                        child_chain_step_id_value = chain_payload.get("chain_step_id")
                        is_new_chain_step = not isinstance(child_chain_step_id_value, int)
                        if isinstance(child_chain_step_id_value, int):
                            child_chain_step = existing_chain_children.get(child_chain_step_id_value)
                        if child_chain_step is None:
                            child_chain_step = FlowStepTemplate(
                                flow_key=scenario.scenario_key,
                                step_key=f"{branch_step.step_key}__chain_{chain_index}",
                                parent_step_id=branch_step.id,
                                branch_option_index=None,
                                step_title=f"Шаг {chain_index + 1}",
                                sort_order=(chain_index + 1) * 10,
                                default_text="Новое сообщение сценария.",
                                custom_text=None,
                                response_type="none",
                                button_options=None,
                                send_mode="immediate",
                                send_time=None,
                                day_offset_workdays=0,
                                target_field=None,
                                launch_scenario_key=None,
                                send_employee_card=False,
                            )
                            db.add(child_chain_step)
                            db.flush()
                        child_chain_step.flow_key = scenario.scenario_key
                        child_chain_step.parent_step_id = branch_step.id
                        child_chain_step.branch_option_index = None
                        child_chain_step.step_key = f"{branch_step.step_key}__chain_{chain_index}"
                        child_chain_step.sort_order = (chain_index + 1) * 10
                        child_chain_step.step_title = (str(chain_payload.get("title") or "").strip() or f"Шаг {chain_index + 1}")
                        chain_text_value = str(chain_payload.get("custom_text") or "").strip()
                        child_chain_step.custom_text = None if is_new_chain_step and not chain_text_value else chain_text_value
                        chain_response_type_value = str(chain_payload.get("response_type") or "none").strip()
                        child_chain_step.response_type = chain_response_type_value if chain_response_type_value in {"none", "text", "file", "buttons", "branching"} else "none"
                        child_chain_step.button_options = (
                            str(chain_payload.get("button_options") or "").strip() or None
                            if child_chain_step.response_type in {"buttons", "branching"}
                            else None
                        )
                        child_chain_step.launch_scenario_key = None
                        child_chain_send_mode_value = str(chain_payload.get("send_mode") or "immediate").strip()
                        child_chain_step.send_mode = child_chain_send_mode_value if child_chain_send_mode_value in {"immediate", "specific_time"} else "immediate"
                        child_chain_step.send_time = (
                            str(chain_payload.get("send_time") or "").strip() or None
                            if child_chain_step.send_mode == "specific_time"
                            else None
                        )
                        child_chain_step.day_offset_workdays = 0
                        child_chain_target_field_value = str(chain_payload.get("target_field") or "").strip()
                        child_chain_step.target_field = child_chain_target_field_value if child_chain_target_field_value in TARGET_FIELD_LABELS else None
                        child_chain_step.send_employee_card = str(chain_payload.get("send_employee_card") or "false").strip() == "true"
                        child_chain_step.notify_on_send_text = str(chain_payload.get("notify_on_send_text") or "").strip() or None
                        child_chain_step.notify_on_send_recipient_ids = str(chain_payload.get("notify_on_send_recipient_ids") or "").strip() or None
                        child_chain_step.notify_on_send_recipient_scope = _normalize_notification_scope(str(chain_payload.get("notify_on_send_recipient_scope") or ""))
                        child_chain_row_ref_value = str(chain_payload.get("row_ref") or "").strip()
                        if child_chain_row_ref_value:
                            chain_step_id_by_ref[child_chain_row_ref_value] = child_chain_step.id
                        preserved_chain_ids.add(child_chain_step.id)
                        chain_upload = request_form.get(f"chain_attachment_{chain_parent_step.id}_{option_idx}_{chain_index}")
                        if chain_upload is not None and getattr(chain_upload, "filename", ""):
                            await _save_step_attachment(child_chain_step, chain_upload)
                        if child_chain_step.response_type == "buttons":
                            child_options = [item.strip() for item in (child_chain_step.button_options or "").splitlines() if item.strip()]
                            preserved_child_option_indexes: set[int] = set()
                            for child_option_idx, _ in enumerate(child_options):
                                payload_by_id = submitted_button_notification_rows.get((child_chain_step.id, child_option_idx), {})
                                payload_by_ref = submitted_button_notification_rows_by_ref.get((child_chain_row_ref_value, child_option_idx), {}) if child_chain_row_ref_value else {}
                                button_payload = payload_by_id or payload_by_ref
                                _sync_button_notification(
                                    db,
                                    child_chain_step,
                                    child_option_idx,
                                    str(button_payload.get("text") or ""),
                                    str(button_payload.get("recipient_ids") or ""),
                                    str(button_payload.get("recipient_scope") or ""),
                                )
                                preserved_child_option_indexes.add(child_option_idx)
                            for notification in db.query(StepButtonNotification).filter(StepButtonNotification.step_id == child_chain_step.id).all():
                                if notification.option_index not in preserved_child_option_indexes:
                                    db.delete(notification)
                        else:
                            db.query(StepButtonNotification).filter(StepButtonNotification.step_id == child_chain_step.id).delete()
                    for existing_id, existing_child in existing_chain_children.items():
                        if existing_id not in preserved_chain_ids:
                            _delete_step_tree(db, existing_child)
                else:
                    db.query(StepButtonNotification).filter(StepButtonNotification.step_id == branch_step.id).delete()
                    for existing_child in existing_chain_children.values():
                        _delete_step_tree(db, existing_child)
            for child in existing_children.values():
                _delete_step_tree(db, child)
        db.commit()
        return _template_edit_redirect(scenario, legacy=_request_prefers_legacy_editor(request))
    return RedirectResponse(url="/flows", status_code=status.HTTP_303_SEE_OTHER)


@router.post("/flows/{scenario_id}/delete")
@router.post("/surveys/{scenario_id}/delete")
def delete_scenario(
    request: Request,
    scenario_id: int,
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    scenario = db.get(ScenarioTemplate, scenario_id)
    if not scenario:
        requested_kind = "survey" if request.url.path.startswith("/surveys/") else "scenario"
        return _template_workspace_redirect(requested_kind)
    kind = getattr(scenario, "scenario_kind", "scenario")
    item_label = _template_entity_meta(kind)["item_label_cap"]
    _delete_template_entity(db, scenario)
    db.commit()
    return _template_workspace_redirect(kind, flash_message=f"{item_label} удалён")


@router.post("/flows/{scenario_id}/copy")
@router.post("/surveys/{scenario_id}/copy")
def copy_scenario(
    request: Request,
    scenario_id: int,
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    scenario = db.get(ScenarioTemplate, scenario_id)
    if not scenario:
        requested_kind = "survey" if request.url.path.startswith("/surveys/") else "scenario"
        return _template_workspace_redirect(requested_kind)
    scenario_copy = _copy_template_entity(db, scenario)
    kind = getattr(scenario, "scenario_kind", "scenario")
    item_label = _template_entity_meta(kind)["item_label_cap"]
    return _template_workspace_redirect(kind, scenario_copy.id, flash_message=f"{item_label} скопирован")


@router.get("/surveys/{scenario_id}/export")
def export_survey_results(
    request: Request,
    scenario_id: int,
    db: Session = Depends(get_db),
):
    auth_redirect = require_auth(request)
    if auth_redirect:
        return auth_redirect
    scenario = db.get(ScenarioTemplate, scenario_id)
    if not scenario or scenario.scenario_kind != "survey":
        return RedirectResponse(url="/surveys", status_code=status.HTTP_303_SEE_OTHER)

    try:
        from openpyxl import Workbook
    except Exception:
        return _template_edit_redirect(
            scenario,
            "Для выгрузки Excel нужен пакет openpyxl.",
            "error",
            legacy=_request_prefers_legacy_editor(request),
        )

    steps = (
        db.query(FlowStepTemplate)
        .filter(FlowStepTemplate.flow_key == scenario.scenario_key)
        .order_by(
            FlowStepTemplate.parent_step_id.is_not(None),
            FlowStepTemplate.sort_order,
            FlowStepTemplate.id,
        )
        .all()
    )
    answers = (
        db.query(SurveyAnswer)
        .filter(SurveyAnswer.scenario_key == scenario.scenario_key)
        .order_by(SurveyAnswer.employee_id, SurveyAnswer.answered_at, SurveyAnswer.id)
        .all()
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результаты"
    step_labels: dict[str, str] = {}
    for step in steps:
        label = (step.custom_text if step.custom_text is not None else step.default_text or "").strip()
        if not label:
            label = (step.step_title or step.step_key).strip() or step.step_key
        label = " ".join(label.split())
        if len(label) > 240:
            label = f"{label[:237]}..."
        step_labels[step.step_key] = label

    sheet.append(["Пользователь ФИО", "Вопрос", "Ответ"])

    for answer in answers:
        employee = db.get(Employee, answer.employee_id)
        if not employee:
            continue
        answer_value = answer.file_name or answer.answer_value or ""
        sheet.append([
            employee.full_name or "",
            step_labels.get(answer.step_key, answer.step_key),
            answer_value,
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"survey_{scenario.id}_results.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename=\"{filename}\"'},
    )


