"""Guarded first-review import; caller owns the transaction."""
import hashlib
import json
from io import BytesIO
from uuid import UUID, uuid5

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Employee, GradeAssessment
from app.grades.workbook import parse_skills_workbook
from app.grades.history import matrix_version
from app.web import grades


def _name(value):
    return ' '.join(value.casefold().replace('ё', 'е').split())


def preview(db, content, aliases=None, zero_sheets=None):
    aliases = aliases or {}
    zero_sheets = zero_sheets or []
    if not isinstance(aliases, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in aliases.items()):
        raise ValueError('Aliases must map source names to reviewed names')
    if not isinstance(zero_sheets, list) or not all(isinstance(x, str) for x in zero_sheets):
        raise ValueError('zero_sheets must be a list of sheet names')
    if len(content) > 5_000_000:
        raise ValueError('Workbook exceeds 5 MB')
    # Reuse parser expansion limits before opening the workbook to discover sheets.
    import zipfile
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if sum(x.file_size for x in archive.infolist()) > 50_000_000:
            raise ValueError('Expanded workbook exceeds 50 MB')
    book = load_workbook(BytesIO(content), read_only=True, keep_links=False)
    try:
        names = [s.title for s in book if next(s.iter_rows(max_row=1, max_col=2, values_only=True), ()) == ('Категория', 'Навык')]
    finally:
        book.close()
    if not names:
        raise ValueError('No competency matrices found')
    parsed = [parse_skills_workbook(content, name, name in zero_sheets) for name in names]
    employees = list(db.scalars(select(Employee).order_by(Employee.id)))
    assessments = list(db.scalars(select(GradeAssessment).order_by(GradeAssessment.id)))
    active_grades = {g['slug']: g['id'] for g in grades.workspace_payload(db)['grades'] if g['active']}
    rows, seen = [], set()
    for sheet in parsed:
        for candidate in sheet['assessment_candidates']:
            source = candidate['name']
            query = _name(aliases.get(source, source)).split()
            matches = [e for e in employees if _name(e.full_name or '').split() == query or
                       (len(query) == 2 and _name(e.full_name or '').split()[:2] == query)]
            employee = matches[0] if len(matches) == 1 else None
            grade = candidate['grade']
            grade_id = active_grades.get(str(grade).strip().casefold()) if grade is not None else None
            known = sum(v['level'] is not None for v in candidate['values'])
            state = 'ready'
            if employee is None:
                state = 'unmatched' if not matches else 'ambiguous'
            elif employee.id in seen:
                state = 'duplicate_employee'
            elif any(a.employee_id == employee.id for a in assessments):
                state = 'existing_review'
            elif grade is not None and grade_id is None:
                state = 'unknown_grade'
            elif not known:
                state = 'no_scores'
            if employee:
                seen.add(employee.id)
            rows.append(dict(sheet=sheet['sheet'], name=source, employee_id=employee.id if employee else None,
                             employee_name=employee.full_name if employee else None, grade=grade,
                             grade_id=grade_id, assessed=known, missing=len(candidate['values']) - known, state=state))
    state = dict(workspace=grades.workspace_payload(db), employees=[(e.id, e.full_name) for e in employees],
                 assessments=[(a.id, a.employee_id, a.status, str(a.archived_at), str(a.updated_at)) for a in assessments],
                 source=hashlib.sha256(content).hexdigest(), aliases=aliases, zero_sheets=zero_sheets)
    return dict(rows=rows, templates=[dict(name=s['sheet'], skills=sum(len(c['skills']) for c in s['catalog']['categories'])) for s in parsed],
                fingerprint=matrix_version(state), source_digest=state['source'],
                warnings=[w for s in parsed for w in s['warnings']]), parsed


def apply(db, content, actor_id, fingerprint, aliases=None, zero_sheets=None):
    plan, parsed = preview(db, content, aliases, zero_sheets)
    if plan['fingerprint'] != fingerprint:
        raise ValueError('Import source or database changed; repeat preview')
    blocked = [r for r in plan['rows'] if r['state'] not in {'ready', 'existing_review', 'no_scores'}]
    if blocked:
        raise ValueError('Resolve unmatched, duplicate or invalid employees before import')
    ready = [r for r in plan['rows'] if r['state'] == 'ready']
    if not ready:
        return dict(imported=[], skipped=plan['rows'], backup=None)
    backup = grades._backup_before_grade_purge(db)
    imported = []
    for sheet in parsed:
        payload = dict(dryRun=True, mode='create-specialization', catalog=sheet['catalog'])
        check = grades.import_catalog(db, payload)
        result = grades.import_catalog(db, dict(payload, dryRun=False, expectedFingerprint=check['fingerprint']))
        spec_id = result['specialization']['id']
        for candidate in sheet['assessment_candidates']:
            match = next(r for r in ready if r['sheet'] == sheet['sheet'] and r['name'] == candidate['name']) if any(r['sheet'] == sheet['sheet'] and r['name'] == candidate['name'] for r in ready) else None
            if match is None:
                continue
            employee_id = match['employee_id']
            view = grades.preview_review(db, employee_id, spec_id, match['grade_id'])
            skills = {(s['categoryName'], s['skillName']): s['skillId'] for s in view['scores']}
            values = [dict(skill_id=skills[(v['category'], v['skill'])], level=v['level']) for v in candidate['values']]
            request_id = str(uuid5(UUID('5c4e9b48-1570-44c5-a084-7128f48ac806'), f"{plan['source_digest']}:{sheet['sheet']}:{candidate['column']}:{employee_id}"))
            review = grades.save_review(db, employee_id, actor_id, dict(request_id=request_id, specialization_id=spec_id,
                current_grade_id=match['grade_id'], matrix_version=view['matrix_version'], values=values,
                comment='Первичная оценка из Skills.xlsx. Дата исходного ревью неизвестна; дата записи — дата импорта.'))
            assessment = db.get(GradeAssessment, review['id'])
            review['import_source'] = dict(digest=plan['source_digest'], sheet=sheet['sheet'], column=candidate['column'], original_review_date=None)
            assessment.catalog_snapshot = json.dumps(review, ensure_ascii=False, allow_nan=False)
            db.flush()
            imported.append(dict(employee_id=employee_id, assessment_id=assessment.id))
    return dict(imported=imported, skipped=[r for r in plan['rows'] if r['state'] != 'ready'], backup=backup)
