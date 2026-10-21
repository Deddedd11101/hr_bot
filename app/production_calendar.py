"""Validated annual calendars. Does not alter scenario schedules."""

import calendar
import csv
import hashlib
import io
import json
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

MONTHS = ("ЯНВАРЬ", "ФЕВРАЛЬ", "МАРТ", "АПРЕЛЬ", "МАЙ", "ИЮНЬ", "ИЮЛЬ", "АВГУСТ", "СЕНТЯБРЬ", "ОКТЯБРЬ", "НОЯБРЬ", "ДЕКАБРЬ")
MAX_UPLOAD = 1024 * 1024


def _flag(value) -> bool:
    if value is True or value == 1 or value == "1":
        return True
    if value is False or value == 0 or value == "0":
        return False
    raise ValueError("Статусы дней должны быть 0 или 1")


def _normal_rows(rows) -> list[dict]:
    days = []
    for row in rows:
        value = row.get("date")
        day = value.date() if isinstance(value, datetime) else value if isinstance(value, date) else date.fromisoformat(str(value))
        workday = _flag(row.get("is_workday"))
        shortened = _flag(row.get("is_shortened", 0) or 0)
        if shortened and not workday:
            raise ValueError("Сокращённый день должен быть рабочим")
        description = str(row.get("description") or "").strip()
        if len(description) > 255 or description.startswith("="):
            raise ValueError("Недопустимое описание дня")
        days.append({"date": day.isoformat(), "is_workday": workday, "is_shortened": shortened, "description": description})
    return days


def _print_layout(sheet) -> list[dict]:
    year = sheet["S1"].value
    if not isinstance(year, int) or not 2000 <= year <= 2100:
        raise ValueError("Не распознана печатная раскладка календаря")
    days = []
    for month in range(1, 13):
        top, left = (2, 11, 20, 29)[(month - 1) // 3], (2, 11, 20)[(month - 1) % 3]
        if sheet.cell(top, left).value != MONTHS[month - 1]:
            raise ValueError("Не совпадают заголовки месяцев печатного календаря")
        for row in sheet.iter_rows(min_row=top + 2, max_row=top + 7, min_col=left, max_col=left + 6):
            for cell in row:
                if cell.value is None or str(cell.value).strip() == "":
                    continue
                if type(cell.value) is not int:
                    raise ValueError("В сетке календаря должны быть номера дней")
                day = date(year, month, cell.value)
                if day.weekday() != cell.column - left:
                    raise ValueError(f"Дата {day} не соответствует дню недели")
                color = cell.font.color
                rgb = color.rgb[-6:].upper() if color and color.type == "rgb" else "000000" if color is None else None
                if rgb not in {"000000", "FF0000", "0000FF"}:
                    raise ValueError(f"Неизвестный цвет дня {day}; используйте табличный формат")
                days.append({"date": day.isoformat(), "is_workday": rgb != "FF0000", "is_shortened": rgb == "0000FF", "description": ""})
    return days


def parse_calendar(contents: bytes, filename: str) -> dict:
    if not contents or len(contents) > MAX_UPLOAD:
        raise ValueError("Размер файла должен быть от 1 байта до 1 МБ")
    if filename.lower().endswith(".csv"):
        stream = io.StringIO(contents.decode("utf-8-sig"))
        sample = stream.read(2048)
        stream.seek(0)
        rows = csv.DictReader(stream, delimiter=";" if sample.splitlines()[0].count(";") > sample.splitlines()[0].count(",") else ",")
        if not {"date", "is_workday"}.issubset(rows.fieldnames or []):
            raise ValueError("Нужны колонки date и is_workday")
        days = _normal_rows(rows)
        layout = "table"
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if len(archive.infolist()) > 1000 or sum(entry.file_size for entry in archive.infolist()) > 20 * MAX_UPLOAD:
                raise ValueError("Слишком большой распакованный Excel")
            if any("externallinks/" in name.lower() or name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
                raise ValueError("Внешние связи и макросы не поддерживаются")
        workbook = load_workbook(io.BytesIO(contents), data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("Календарь должен содержать один лист")
            sheet = workbook.active
            if sheet.max_row > 500 or sheet.max_column > 40:
                raise ValueError("Слишком большая область листа")
            if any(cell.data_type == "f" for row in sheet for cell in row):
                raise ValueError("Формулы в календаре не поддерживаются")
            if sheet["A1"].value == "date":
                headers = [cell.value for cell in sheet[1]]
                if "is_workday" not in headers or len(headers) != len(set(headers)):
                    raise ValueError("Нужны уникальные заголовки и колонка is_workday")
                days = _normal_rows(dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row))
                layout = "table"
            else:
                days = _print_layout(sheet)
                layout = "print_a4"
        finally:
            workbook.close()
    else:
        raise ValueError("Поддерживаются XLSX и UTF-8 CSV")
    if not days:
        raise ValueError("В файле нет дат")
    year = date.fromisoformat(days[0]["date"]).year
    expected = {date(year, month, day).isoformat() for month in range(1, 13) for day in range(1, calendar.monthrange(year, month)[1] + 1)}
    if not 2000 <= year <= 2100 or len(days) != len(expected) or {item["date"] for item in days} != expected:
        raise ValueError("Нужны все даты одного года без пропусков и повторов")
    days.sort(key=lambda item: item["date"])
    canonical = json.dumps(days, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return {"year": year, "days": days, "digest": hashlib.sha256(canonical.encode()).hexdigest(), "layout": layout,
            "workdays": sum(item["is_workday"] for item in days), "days_off": sum(not item["is_workday"] for item in days),
            "shortened": sum(item["is_shortened"] for item in days)}
