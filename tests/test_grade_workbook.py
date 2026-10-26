from io import BytesIO
import unittest
from openpyxl import Workbook
from app.grades.workbook import parse_skills_workbook


class GradeWorkbookTests(unittest.TestCase):
    def book(self, *, blank=False, score=4):
        book = Workbook()
        sheet = book.active
        sheet.title = "Системный аналитик"
        sheet.append(["Категория", "Навык", "Junior", "Middle", "Senior", "Lead", "Важность", "Тестовый сотрудник", "Столбец 1"])
        sheet.append([None] * 7 + ["Middle"])
        sheet.append(["Hard skills", "Моделирование", None if blank else "Понимание", "Умение", "Экспертность", "Лидерство", "Критичный", score, "Умение"])
        survey = book.create_sheet("Вопросы опроса")
        survey.append(["Do not import this"])
        stream = BytesIO()
        book.save(stream)
        return stream.getvalue()

    def test_numeric_scores_and_missing_values_are_separate(self):
        parsed = parse_skills_workbook(self.book())
        self.assertEqual(len(parsed["catalog"]["categories"]), 1)
        candidate = parsed["assessment_candidates"][0]
        self.assertEqual(candidate["values"][0]["level"], 4)
        self.assertEqual(candidate["grade"], "Middle")
        self.assertIsNone(parse_skills_workbook(self.book(score="н/д"))["assessment_candidates"][0]["values"][0]["level"])
        self.assertEqual(parse_skills_workbook(self.book(score=0))["assessment_candidates"][0]["values"][0]["level"], 0)
        self.assertIsNone(parse_skills_workbook(self.book(score="—"))["assessment_candidates"][0]["values"][0]["level"])

    def test_blank_requirement_needs_explicit_policy(self):
        with self.assertRaises(ValueError):
            parse_skills_workbook(self.book(blank=True))
        parsed = parse_skills_workbook(self.book(blank=True), blank_requirements_zero=True)
        self.assertEqual(parsed["catalog"]["categories"][0]["skills"][0]["levels"]["junior"], 0)
        self.assertEqual(len(parsed["warnings"]), 1)

    def test_formula_and_survey_sheet_are_rejected(self):
        with self.assertRaises(ValueError):
            parse_skills_workbook(self.book(score="=2+2"))
        with self.assertRaises(ValueError):
            parse_skills_workbook(self.book(), "Вопросы опроса")

    def test_template_without_grade_row_keeps_first_skill(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(self.book()))
        sheet = book.active
        sheet.delete_rows(2)
        stream = BytesIO()
        book.save(stream)
        parsed = parse_skills_workbook(stream.getvalue())
        self.assertEqual(len(parsed["catalog"]["categories"][0]["skills"]), 1)
        candidate = parsed["assessment_candidates"][0]
        self.assertIsNone(candidate["grade"])
        self.assertEqual(candidate["values"][0]["row"], 2)
        self.assertEqual(candidate["values"][0]["level"], 4)

    def test_rows_outside_import_limit_are_not_silently_ignored(self):
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(self.book()))
        book.active.cell(1001, 1, "Outside import range")
        stream = BytesIO()
        book.save(stream)
        with self.assertRaisesRegex(ValueError, "exceeds"):
            parse_skills_workbook(stream.getvalue())
