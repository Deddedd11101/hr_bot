import calendar
import csv
import io
import unittest
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.auth import authenticate_account, create_admin_session_token
from app.database import SessionLocal, init_db
from app.main import AUTH_COOKIE_NAME, app
from app.models import ProductionCalendarVersion, ProductionCalendarYear
from app.production_calendar import MONTHS, parse_calendar


def csv_calendar(year=2092, override=None):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["date", "is_workday", "is_shortened"])
    for month in range(1, 13):
        for day in range(1, calendar.monthrange(year, month)[1] + 1):
            current = date(year, month, day)
            writer.writerow([current.isoformat(), override if override is not None and month == 1 and day == 1 else int(current.weekday() < 5), 0])
    return output.getvalue().encode()


class CalendarParserTests(unittest.TestCase):
    def test_xml_entities_rejected_before_loading_workbook(self):
        import zipfile
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("xl/workbook.xml", '<!DOCTYPE test [<!ENTITY x "data">]><test>&x;</test>'.encode("utf-16"))
        with self.assertRaisesRegex(ValueError, "DTD"):
            parse_calendar(output.getvalue(), "hostile.xlsx")

    def test_complete_leap_year_and_missing_duplicate_dates(self):
        data = csv_calendar()
        self.assertEqual(len(parse_calendar(data, "year.csv")["days"]), 366)
        for broken in (data[:data.rfind(b"\r\n", 0, len(data) - 2)], data + data.splitlines()[1] + b"\r\n"):
            with self.assertRaises(ValueError):
                parse_calendar(broken, "year.csv")

    def test_print_layout_color_mapping_and_alignment(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["S1"] = 2026
        from openpyxl.styles import Font
        for month in range(1, 13):
            top, left = (2, 11, 20, 29)[(month - 1) // 3], (2, 11, 20)[(month - 1) % 3]
            sheet.cell(top, left, MONTHS[month - 1])
            for day in range(1, calendar.monthrange(2026, month)[1] + 1):
                current = date(2026, month, day)
                offset = date(2026, month, 1).weekday() + day - 1
                cell = sheet.cell(top + 2 + offset // 7, left + offset % 7, day)
                cell.font = Font(color="FF0000" if current.weekday() >= 5 else "000000")
        sheet["E17"].font = Font(color="0000FF")
        output = io.BytesIO()
        workbook.save(output)
        parsed = parse_calendar(output.getvalue(), "print.xlsx")
        self.assertEqual(len(parsed["days"]), 365)
        self.assertTrue(next(d for d in parsed["days"] if d["date"] == "2026-04-30")["is_shortened"])
        sheet["E4"] = "=1+1"
        output = io.BytesIO()
        workbook.save(output)
        with self.assertRaises(ValueError):
            parse_calendar(output.getvalue(), "bad.xlsx")


class CalendarImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.client = TestClient(app)
        with SessionLocal() as db:
            actor = authenticate_account(db, "admin", "admin123")
            cls.client.cookies.set(AUTH_COOKIE_NAME, create_admin_session_token(actor.id))

    def tearDown(self):
        with SessionLocal() as db:
            db.query(ProductionCalendarYear).filter_by(year=2092).delete()
            db.query(ProductionCalendarVersion).filter_by(year=2092).delete()
            db.commit()

    def post(self, contents, **data):
        return self.client.post("/api/settings/production-calendars/import", data=data, files={"file": ("calendar.csv", contents, "text/csv")})

    def test_preview_readonly_apply_history_and_stale_request(self):
        data = csv_calendar()
        preview = self.post(data).json()
        with SessionLocal() as db:
            self.assertIsNone(db.get(ProductionCalendarYear, 2092))
        applied = self.post(data, mode="apply", expected_digest=preview["digest"], expected_version_id=0)
        self.assertEqual(applied.status_code, 200, applied.text)
        first_id = applied.json()["version"]["id"]
        stale = self.post(data, mode="apply", expected_digest=preview["digest"], expected_version_id=0)
        self.assertEqual(stale.status_code, 409)
        current_preview = self.post(data).json()
        repeat = self.post(data, mode="apply", expected_digest=current_preview["digest"], expected_version_id=first_id)
        self.assertEqual(repeat.json()["version"]["id"], first_id)
        different = csv_calendar(override=0)
        changed = self.post(different).json()
        self.assertEqual(changed["changed_days"], 1)
        self.assertEqual(self.post(different, mode="apply", expected_digest=changed["digest"], expected_version_id=first_id).status_code, 200)
        versions = self.client.get("/api/settings/production-calendars").json()["versions"]
        own_versions = [v for v in versions if v["year"] == 2092]
        self.assertEqual(len(own_versions), 2)
        self.assertEqual(sum(v["is_current"] for v in own_versions), 1)

    def test_tampered_digest_and_permission_checks(self):
        data = csv_calendar()
        self.assertEqual(self.post(data, mode="apply", expected_digest="wrong", expected_version_id=0).status_code, 409)
        client = TestClient(app)
        self.assertEqual(client.get("/api/settings/production-calendars").status_code, 401)
        with SessionLocal() as db:
            account = authenticate_account(db, "hr", "hr123")
            client.cookies.set(AUTH_COOKIE_NAME, create_admin_session_token(account.id))
        self.assertEqual(client.get("/api/settings/production-calendars").status_code, 403)
        self.assertEqual(client.post("/api/settings/production-calendars/import", files={"file": ("calendar.csv", data)}).status_code, 403)

    def test_concurrent_first_apply_has_one_winner(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        data = csv_calendar()
        preview = self.post(data).json()
        barrier = Barrier(2)
        token = self.client.cookies.get(AUTH_COOKIE_NAME)
        def claim(_):
            client = TestClient(app)
            client.cookies.set(AUTH_COOKIE_NAME, token)
            barrier.wait(timeout=10)
            return client.post("/api/settings/production-calendars/import", data={"mode": "apply", "expected_digest": preview["digest"], "expected_version_id": 0}, files={"file": ("calendar.csv", data)}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(claim, range(2))), [200, 409])
        with SessionLocal() as db:
            self.assertEqual(db.query(ProductionCalendarVersion).filter_by(year=2092).count(), 1)
