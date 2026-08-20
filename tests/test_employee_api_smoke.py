import unittest
import asyncio
from io import BytesIO
from datetime import UTC, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import authenticate_account, create_admin_session_token
from app.database import SessionLocal, init_db
from app.flow_templates import CANDIDATE_WORK_STAGE_LABELS
from app.main import AUTH_COOKIE_NAME, app
from app.messaging.identity import get_primary_chat_id, set_primary_chat_id
from app.messaging.service import (
    BLOCKED_USER_TEXT,
    MENU_BACK_BUTTON_TEXT,
    MENU_HOME_BUTTON_TEXT,
    current_menu_set,
    get_or_create_employee_by_chat,
    handle_menu_button,
    handle_start_command,
    handle_text_event,
)
from app.scenario_engine import SINGLE_STEP_REQUEST_PREFIX, handle_button_response, send_step
from app.models import (
    AdminAccount,
    BotMenuButton,
    BotMenuSet,
    DocumentLibraryItem,
    Employee,
    EmployeeAssignmentHistory,
    EmployeeDocumentLink,
    EmployeeMessengerAccount,
    EmployeeFile,
    EmployeeManualBotMessage,
    FlowLaunchRequest,
    FlowStepTemplate,
    HrSettings,
    MassMessageAction,
    MassScenarioAction,
    Position,
    ScenarioProgress,
    ScenarioTemplate,
    StepButtonNotification,
    StepSendNotification,
    SurveyAnswer,
)


class DummyMessenger:
    def __init__(self) -> None:
        self.sent_texts: list[tuple[str, str]] = []
        self.sent_menus: list[tuple[str, str, list[str]]] = []
        self.sent_documents: list[tuple[str, str | None, str | None]] = []

    async def send_text(self, chat_id: str, text: str, reply_markup=None) -> None:
        self.sent_texts.append((chat_id, text))

    async def send_menu(self, chat_id: str, text: str, buttons: list[str]) -> None:
        self.sent_menus.append((chat_id, text, buttons))

    async def send_photo_path(self, chat_id: str, path, filename=None, reply_markup=None, caption=None) -> None:
        return None

    async def send_photo_bytes(self, chat_id: str, data: bytes, filename: str, reply_markup=None, caption=None) -> None:
        return None

    async def send_document_path(self, chat_id: str, path, filename=None, reply_markup=None, caption=None) -> None:
        self.sent_documents.append((chat_id, str(path) if path is not None else None, filename))
        return None

    async def close(self) -> None:
        return None


class FailingMessenger(DummyMessenger):
    async def send_text(self, chat_id: str, text: str, reply_markup=None) -> None:
        raise RuntimeError("telegram send failed")


class EmployeeApiSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        init_db()
        cls.client = TestClient(app)
        with SessionLocal() as db:
            account = authenticate_account(db, "admin", "admin123")
            if account is None:
                raise AssertionError("Admin account is not available for API smoke tests.")
            cls.client.cookies.set(AUTH_COOKIE_NAME, create_admin_session_token(account.id))

    def setUp(self) -> None:
        self.unique_tag = uuid4().hex[:12]
        with SessionLocal() as db:
            employee = Employee(
                full_name="API Smoke Employee",
                telegram_user_id=None,
                telegram_username=None,
                first_workday=datetime.now(UTC).date(),
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                candidate_status="new",
                employee_stage="candidate",
                candidate_work_stage="testing",
            )
            db.add(employee)
            db.commit()
            db.refresh(employee)
            self.employee_id = employee.id
            hr_settings = db.query(HrSettings).first()
            self.hr_settings_snapshot = None
            if hr_settings is not None:
                self.hr_settings_snapshot = {
                    "id": hr_settings.id,
                    "hr_name": hr_settings.hr_name,
                    "telegram_user_id": hr_settings.telegram_user_id,
                    "notification_recipient_ids": hr_settings.notification_recipient_ids,
                    "notify_scenario_completed": hr_settings.notify_scenario_completed,
                    "notify_test_task_received": hr_settings.notify_test_task_received,
                    "notify_user_actions": hr_settings.notify_user_actions,
                    "default_menu_set_id": hr_settings.default_menu_set_id,
                    "default_employee_menu_set_id": hr_settings.default_employee_menu_set_id,
                    "default_candidate_menu_set_id": hr_settings.default_candidate_menu_set_id,
                }

    @staticmethod
    def _initial_candidate_stage_key() -> str:
        if "hr_interview" not in CANDIDATE_WORK_STAGE_LABELS:
            raise AssertionError("Expected hr_interview candidate stage key is missing.")
        return "hr_interview"

    def _staff_update_payload(self, **overrides) -> dict:
        payload = {
            "full_name": "API Smoke Employee",
            "chat_id": "",
            "chat_handle": "",
            "first_workday": "2026-06-10",
            "desired_position": "Аналитик",
            "birth_date": "1999-04-10",
            "work_email": "employee@example.com",
            "work_hours": "10:00-19:00",
            "is_manager": False,
            "is_mentor": False,
            "manager_employee_id": "",
            "mentor_adaptation_employee_id": "",
            "mentor_ipr_employee_id": "",
            "adaptation_tasks_url": "",
            "adaptation_feedback_url": "",
            "adaptation_midpoint": "",
            "adaptation_end": "",
            "employee_stage": "staff",
            "candidate_work_stage": "",
            "salary_expectation": "",
            "personal_data_consent": False,
            "employee_data_consent": True,
            "is_bot_blocked": False,
            "test_task_due_at": "",
            "notes": "",
        }
        payload.update(overrides)
        return payload

    def _create_staff_employee(
        self,
        *,
        full_name: str,
        telegram_user_id: str,
        is_manager: bool = False,
        is_mentor: bool = False,
    ) -> int:
        with SessionLocal() as db:
            employee = Employee(
                full_name=full_name,
                telegram_user_id=telegram_user_id,
                first_workday=None,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                employee_stage="staff",
                is_manager=is_manager,
                is_mentor=is_mentor,
            )
            db.add(employee)
            db.commit()
            db.refresh(employee)
            return employee.id

    def tearDown(self) -> None:
        with SessionLocal() as db:
            db.query(BotMenuButton).filter(BotMenuButton.label.like(f"codex-%-{self.unique_tag}%")).delete(synchronize_session=False)
            created_menu_sets = db.query(BotMenuSet).filter(BotMenuSet.title.like(f"codex-%-{self.unique_tag}%")).all()
            created_menu_set_ids = [menu_set.id for menu_set in created_menu_sets]
            if created_menu_set_ids:
                db.query(BotMenuButton).filter(BotMenuButton.menu_set_id.in_(created_menu_set_ids)).delete(synchronize_session=False)
                db.query(BotMenuButton).filter(BotMenuButton.target_menu_set_id.in_(created_menu_set_ids)).update(
                    {
                        BotMenuButton.action_type: "inactive",
                        BotMenuButton.target_menu_set_id: None,
                    },
                    synchronize_session=False,
                )
                db.query(Employee).filter(Employee.current_menu_set_id.in_(created_menu_set_ids)).update(
                    {Employee.current_menu_set_id: None, Employee.current_menu_path: None},
                    synchronize_session=False,
                )
                for menu_set in created_menu_sets:
                    db.delete(menu_set)
            created_documents = db.query(DocumentLibraryItem).filter(DocumentLibraryItem.title.like(f"codex-%-{self.unique_tag}%")).all()
            for item in created_documents:
                stored_path = (item.stored_path or "").strip()
                if stored_path:
                    path = Path(stored_path)
                    if path.exists():
                        path.unlink()
                db.delete(item)
            db.query(MassMessageAction).filter(MassMessageAction.message_text.like(f"%{self.unique_tag}%")).delete(synchronize_session=False)
            created_scenarios = db.query(ScenarioTemplate).filter(ScenarioTemplate.title.like(f"codex-%-{self.unique_tag}%")).all()
            for scenario in created_scenarios:
                scenario_step_ids = [
                    row[0]
                    for row in db.query(FlowStepTemplate.id).filter(FlowStepTemplate.flow_key == scenario.scenario_key).all()
                ]
                if scenario_step_ids:
                    db.query(StepButtonNotification).filter(StepButtonNotification.step_id.in_(scenario_step_ids)).delete(synchronize_session=False)
                    db.query(StepSendNotification).filter(StepSendNotification.step_id.in_(scenario_step_ids)).delete(synchronize_session=False)
                db.query(FlowStepTemplate).filter(FlowStepTemplate.flow_key == scenario.scenario_key).delete(synchronize_session=False)
                db.query(SurveyAnswer).filter(SurveyAnswer.scenario_key == scenario.scenario_key).delete(synchronize_session=False)
                db.delete(scenario)
            db.query(FlowLaunchRequest).filter(FlowLaunchRequest.employee_id == self.employee_id).delete(synchronize_session=False)
            db.query(EmployeeManualBotMessage).filter(EmployeeManualBotMessage.employee_id == self.employee_id).delete(
                synchronize_session=False
            )
            db.query(EmployeeFile).filter(EmployeeFile.employee_id == self.employee_id).delete(synchronize_session=False)
            db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == self.employee_id).delete()
            extra_employees = db.query(Employee).filter(Employee.full_name.like(f"%{self.unique_tag}%")).all()
            for extra_employee in extra_employees:
                if extra_employee.id == self.employee_id:
                    continue
                db.query(FlowLaunchRequest).filter(FlowLaunchRequest.employee_id == extra_employee.id).delete(synchronize_session=False)
                db.query(EmployeeManualBotMessage).filter(
                    EmployeeManualBotMessage.employee_id == extra_employee.id
                ).delete(synchronize_session=False)
                db.query(EmployeeFile).filter(EmployeeFile.employee_id == extra_employee.id).delete(synchronize_session=False)
                db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == extra_employee.id).delete()
                db.query(EmployeeAssignmentHistory).filter(
                    (EmployeeAssignmentHistory.subject_employee_id == extra_employee.id)
                    | (EmployeeAssignmentHistory.assigned_employee_id == extra_employee.id)
                ).delete(synchronize_session=False)
                db.delete(extra_employee)
            db.query(EmployeeAssignmentHistory).filter(
                EmployeeAssignmentHistory.subject_employee_id == self.employee_id
            ).delete(synchronize_session=False)
            employee = db.get(Employee, self.employee_id)
            if employee is not None:
                db.delete(employee)
            if self.hr_settings_snapshot is not None:
                hr_settings = db.get(HrSettings, self.hr_settings_snapshot["id"])
                if hr_settings is not None:
                    hr_settings.hr_name = self.hr_settings_snapshot["hr_name"]
                    hr_settings.telegram_user_id = self.hr_settings_snapshot["telegram_user_id"]
                    hr_settings.notification_recipient_ids = self.hr_settings_snapshot["notification_recipient_ids"]
                    hr_settings.notify_scenario_completed = self.hr_settings_snapshot["notify_scenario_completed"]
                    hr_settings.notify_test_task_received = self.hr_settings_snapshot["notify_test_task_received"]
                    hr_settings.notify_user_actions = self.hr_settings_snapshot["notify_user_actions"]
                    hr_settings.default_menu_set_id = self.hr_settings_snapshot["default_menu_set_id"]
                    hr_settings.default_employee_menu_set_id = self.hr_settings_snapshot["default_employee_menu_set_id"]
                    hr_settings.default_candidate_menu_set_id = self.hr_settings_snapshot["default_candidate_menu_set_id"]
            db.commit()

    def test_employee_detail_api_returns_ok(self) -> None:
        response = self.client.get(f"/api/employees/{self.employee_id}")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["employee"]["id"], self.employee_id)

    def test_employee_offer_document_file_api_persists_slot_payload(self) -> None:
        response = self.client.post(
            f"/api/employees/{self.employee_id}/document-slots/offer/file",
            files={"upload": ("offer.pdf", b"fake-offer-pdf", "application/pdf")},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["item"]["slot_key"], "offer")
        self.assertEqual(payload["item"]["item_kind"], "file")
        self.assertEqual(payload["item"]["title"], "Оффер")
        self.assertIn(f"/employees/{self.employee_id}/files/", payload["item"]["url"])
        self.assertEqual(payload["payload"]["offer_document"]["slot_key"], "offer")

        with SessionLocal() as db:
            link_row = (
                db.query(EmployeeDocumentLink)
                .filter(EmployeeDocumentLink.employee_id == self.employee_id)
                .first()
            )
            self.assertIsNotNone(link_row)
            if link_row is not None:
                self.assertEqual(link_row.slot_key, "offer")
                self.assertEqual(link_row.item_kind, "file")
                self.assertIsNotNone(link_row.employee_file_id)

    def test_workspace_payload_exposes_hr_notification_recipient(self) -> None:
        with SessionLocal() as db:
            hr_settings = db.get(HrSettings, 1)
            self.assertIsNotNone(hr_settings)
            hr_settings.hr_name = f"HR {self.unique_tag}"
            hr_settings.telegram_user_id = "770001"
            db.commit()

        scenario_key = f"codex-hr-recipient-{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-hr-recipient-{self.unique_tag}",
                role_scope="all",
                scenario_kind="scenario",
                sort_order=0,
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        response = self.client.get(f"/api/flows/workspace?scenario_id={scenario_id}")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        options = payload["workspace"]["notification_recipient_options"]
        self.assertIn(
            {
                "token": "hr",
                "label": f"HR {self.unique_tag}",
                "description": "HR из системных настроек",
                "kind": "hr",
            },
            options,
        )
        self.assertIn(
            {
                "token": "manager",
                "label": "Руководитель",
                "description": "Руководитель из карточки сотрудника",
                "kind": "role",
            },
            options,
        )
        self.assertIn(
            {
                "token": "mentor_adaptation",
                "label": "Наставник адаптации",
                "description": "Наставник адаптации из карточки сотрудника",
                "kind": "role",
            },
            options,
        )
        self.assertIn(
            {
                "token": "mentor_ipr",
                "label": "Наставник ИПР",
                "description": "Наставник ИПР из карточки сотрудника",
                "kind": "role",
            },
            options,
        )
        self.assertFalse(any(str(option.get("token") or "").startswith("employee:") for option in options))

    def test_employee_detail_update_supports_staff_selects_and_adaptation_dates(self) -> None:
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            manager = Employee(
                full_name=f"Manager {self.unique_tag}",
                telegram_user_id="700001",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
                is_manager=True,
            )
            mentor_adaptation = Employee(
                full_name=f"Mentor Adaptation {self.unique_tag}",
                telegram_user_id="700002",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
                is_mentor=True,
            )
            mentor_ipr = Employee(
                full_name=f"Mentor IPR {self.unique_tag}",
                telegram_user_id="700003",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
                is_mentor=True,
            )
            plain_staff = Employee(
                full_name=f"Plain Staff {self.unique_tag}",
                telegram_user_id="700004",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
            )
            db.add(manager)
            db.add(mentor_adaptation)
            db.add(mentor_ipr)
            db.add(plain_staff)
            db.commit()
            db.refresh(manager)
            db.refresh(mentor_adaptation)
            db.refresh(mentor_ipr)
            db.refresh(plain_staff)
            manager_id = manager.id
            mentor_adaptation_id = mentor_adaptation.id
            mentor_ipr_id = mentor_ipr.id
            plain_staff_id = plain_staff.id

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "API Smoke Employee",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "2026-06-10",
                "desired_position": "Аналитик",
                "birth_date": "1999-04-10",
                "work_email": "employee@example.com",
                "work_hours": "10:00-19:00",
                "is_manager": True,
                "is_mentor": True,
                "manager_employee_id": str(manager_id),
                "mentor_adaptation_employee_id": str(mentor_adaptation_id),
                "mentor_ipr_employee_id": str(mentor_ipr_id),
                "adaptation_tasks_url": "https://example.com/tasks",
                "adaptation_feedback_url": "https://example.com/feedback",
                "adaptation_midpoint": "2026-06-24",
                "adaptation_end": "2026-07-10",
                "employee_stage": "staff",
                "candidate_work_stage": "",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": True,
                "is_bot_blocked": False,
                "test_task_due_at": "",
                "notes": "updated",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["employee"]["manager_employee_id"], str(manager_id))
        self.assertEqual(payload["employee"]["mentor_adaptation_employee_id"], str(mentor_adaptation_id))
        self.assertEqual(payload["employee"]["mentor_ipr_employee_id"], str(mentor_ipr_id))
        self.assertTrue(payload["employee"]["is_manager"])
        self.assertTrue(payload["employee"]["is_mentor"])
        self.assertEqual(payload["employee"]["adaptation_tasks_url"], "https://example.com/tasks")
        self.assertEqual(payload["employee"]["adaptation_feedback_url"], "https://example.com/feedback")
        self.assertEqual(payload["employee"]["adaptation_midpoint"], "2026-06-24")
        self.assertEqual(payload["employee"]["adaptation_end"], "2026-07-10")
        self.assertTrue(any(option["value"] == str(manager_id) for option in payload["options"]["staff_employee_values"]))
        self.assertTrue(any(option["value"] == str(manager_id) for option in payload["options"]["manager_employee_values"]))
        self.assertFalse(any(option["value"] == str(plain_staff_id) for option in payload["options"]["manager_employee_values"]))
        self.assertTrue(any(option["value"] == str(mentor_adaptation_id) for option in payload["options"]["mentor_employee_values"]))
        self.assertFalse(any(option["value"] == str(plain_staff_id) for option in payload["options"]["mentor_employee_values"]))

        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertEqual(employee.manager_employee_id, manager_id)
            self.assertEqual(employee.mentor_adaptation_employee_id, mentor_adaptation_id)
            self.assertEqual(employee.mentor_ipr_employee_id, mentor_ipr_id)
            self.assertTrue(employee.is_manager)
            self.assertTrue(employee.is_mentor)
            self.assertEqual(employee.manager_telegram_id, "700001")
            self.assertEqual(employee.adaptation_midpoint.isoformat(), "2026-06-24")
            self.assertEqual(employee.adaptation_end.isoformat(), "2026-07-10")

            for related_employee_id in (manager_id, mentor_adaptation_id, mentor_ipr_id, plain_staff_id):
                related_employee = db.get(Employee, related_employee_id)
                if related_employee is not None:
                    db.delete(related_employee)
            db.commit()

    def test_assign_manager_creates_history_row(self) -> None:
        manager_id = self._create_staff_employee(
            full_name=f"Manager History {self.unique_tag}",
            telegram_user_id="730001",
            is_manager=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(manager_id)),
        )

        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            rows = (
                db.query(EmployeeAssignmentHistory)
                .filter(
                    EmployeeAssignmentHistory.subject_employee_id == self.employee_id,
                    EmployeeAssignmentHistory.assignment_role == "manager",
                )
                .order_by(EmployeeAssignmentHistory.id.asc())
                .all()
            )
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].assigned_employee_id, manager_id)
            self.assertIsNone(rows[0].ended_at)
            self.assertIsNotNone(rows[0].assigned_by_account_id)

    def test_replace_manager_closes_old_row_and_creates_new_row(self) -> None:
        first_manager_id = self._create_staff_employee(
            full_name=f"Manager Old {self.unique_tag}",
            telegram_user_id="730002",
            is_manager=True,
        )
        second_manager_id = self._create_staff_employee(
            full_name=f"Manager New {self.unique_tag}",
            telegram_user_id="730003",
            is_manager=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        first_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(first_manager_id)),
        )
        second_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(second_manager_id)),
        )

        self.assertEqual(first_response.status_code, 200)
        self.assertEqual(second_response.status_code, 200)
        with SessionLocal() as db:
            rows = (
                db.query(EmployeeAssignmentHistory)
                .filter(
                    EmployeeAssignmentHistory.subject_employee_id == self.employee_id,
                    EmployeeAssignmentHistory.assignment_role == "manager",
                )
                .order_by(EmployeeAssignmentHistory.id.asc())
                .all()
            )
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0].assigned_employee_id, first_manager_id)
            self.assertIsNotNone(rows[0].ended_at)
            self.assertEqual(rows[1].assigned_employee_id, second_manager_id)
            self.assertIsNone(rows[1].ended_at)

    def test_clear_mentor_closes_active_row(self) -> None:
        mentor_id = self._create_staff_employee(
            full_name=f"Mentor Clear {self.unique_tag}",
            telegram_user_id="730004",
            is_mentor=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        assign_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(mentor_adaptation_employee_id=str(mentor_id)),
        )
        clear_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(mentor_adaptation_employee_id=""),
        )

        self.assertEqual(assign_response.status_code, 200)
        self.assertEqual(clear_response.status_code, 200)
        with SessionLocal() as db:
            rows = (
                db.query(EmployeeAssignmentHistory)
                .filter(
                    EmployeeAssignmentHistory.subject_employee_id == self.employee_id,
                    EmployeeAssignmentHistory.assignment_role == "mentor_adaptation",
                )
                .order_by(EmployeeAssignmentHistory.id.asc())
                .all()
            )
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].assigned_employee_id, mentor_id)
            self.assertIsNotNone(rows[0].ended_at)

    def test_saving_unchanged_employee_does_not_duplicate_history(self) -> None:
        manager_id = self._create_staff_employee(
            full_name=f"Manager Stable {self.unique_tag}",
            telegram_user_id="730005",
            is_manager=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        first_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(manager_id)),
        )
        second_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(manager_id), notes="same assignment"),
        )

        self.assertEqual(first_response.status_code, 200)
        self.assertEqual(second_response.status_code, 200)
        with SessionLocal() as db:
            rows = (
                db.query(EmployeeAssignmentHistory)
                .filter(
                    EmployeeAssignmentHistory.subject_employee_id == self.employee_id,
                    EmployeeAssignmentHistory.assignment_role == "manager",
                )
                .order_by(EmployeeAssignmentHistory.id.asc())
                .all()
            )
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].assigned_employee_id, manager_id)
            self.assertIsNone(rows[0].ended_at)

    def test_separate_roles_do_not_conflict(self) -> None:
        manager_id = self._create_staff_employee(
            full_name=f"Manager Separate {self.unique_tag}",
            telegram_user_id="730006",
            is_manager=True,
        )
        mentor_adaptation_id = self._create_staff_employee(
            full_name=f"Mentor Adapt Separate {self.unique_tag}",
            telegram_user_id="730007",
            is_mentor=True,
        )
        mentor_ipr_id = self._create_staff_employee(
            full_name=f"Mentor IPR Separate {self.unique_tag}",
            telegram_user_id="730008",
            is_mentor=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(
                manager_employee_id=str(manager_id),
                mentor_adaptation_employee_id=str(mentor_adaptation_id),
                mentor_ipr_employee_id=str(mentor_ipr_id),
            ),
        )

        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            rows = (
                db.query(EmployeeAssignmentHistory)
                .filter(EmployeeAssignmentHistory.subject_employee_id == self.employee_id)
                .order_by(EmployeeAssignmentHistory.assignment_role.asc(), EmployeeAssignmentHistory.id.asc())
                .all()
            )
            self.assertEqual(len(rows), 3)
            by_role = {row.assignment_role: row for row in rows}
            self.assertEqual(by_role["manager"].assigned_employee_id, manager_id)
            self.assertEqual(by_role["mentor_adaptation"].assigned_employee_id, mentor_adaptation_id)
            self.assertEqual(by_role["mentor_ipr"].assigned_employee_id, mentor_ipr_id)
            self.assertTrue(all(row.ended_at is None for row in rows))

    def test_employee_detail_api_includes_assignment_history(self) -> None:
        manager_id = self._create_staff_employee(
            full_name=f"Manager Payload {self.unique_tag}",
            telegram_user_id="730009",
            is_manager=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        update_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(manager_id)),
        )
        self.assertEqual(update_response.status_code, 200)

        response = self.client.get(f"/api/employees/{self.employee_id}")

        self.assertEqual(response.status_code, 200)
        history_items = response.json()["assignment_history"]
        self.assertEqual(len(history_items), 1)
        self.assertEqual(history_items[0]["assignment_role"], "manager")
        self.assertEqual(history_items[0]["role_label"], "Руководитель")
        self.assertEqual(history_items[0]["assigned_employee_id"], manager_id)
        self.assertEqual(history_items[0]["assigned_employee_name"], f"Manager Payload {self.unique_tag}")
        self.assertTrue(history_items[0]["is_active"])
        self.assertIsNone(history_items[0]["ended_at"])

    def test_employee_detail_api_serializes_active_and_closed_assignment_history_rows(self) -> None:
        first_manager_id = self._create_staff_employee(
            full_name=f"Manager Closed {self.unique_tag}",
            telegram_user_id="730010",
            is_manager=True,
        )
        second_manager_id = self._create_staff_employee(
            full_name=f"Manager Active {self.unique_tag}",
            telegram_user_id="730011",
            is_manager=True,
        )
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        first_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(first_manager_id)),
        )
        second_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json=self._staff_update_payload(manager_employee_id=str(second_manager_id)),
        )

        self.assertEqual(first_response.status_code, 200)
        self.assertEqual(second_response.status_code, 200)

        response = self.client.get(f"/api/employees/{self.employee_id}")

        self.assertEqual(response.status_code, 200)
        history_items = response.json()["assignment_history"]
        self.assertEqual(len(history_items), 2)
        self.assertEqual(history_items[0]["assigned_employee_id"], second_manager_id)
        self.assertEqual(history_items[0]["assigned_employee_name"], f"Manager Active {self.unique_tag}")
        self.assertTrue(history_items[0]["is_active"])
        self.assertIsNone(history_items[0]["ended_at"])
        self.assertEqual(history_items[1]["assigned_employee_id"], first_manager_id)
        self.assertEqual(history_items[1]["assigned_employee_name"], f"Manager Closed {self.unique_tag}")
        self.assertFalse(history_items[1]["is_active"])
        self.assertIsNotNone(history_items[1]["ended_at"])

    def test_employee_detail_api_tolerates_missing_assignment_history_employee(self) -> None:
        missing_assigned_employee_id = 999999000 + self.employee_id
        with SessionLocal() as db:
            db.add(
                EmployeeAssignmentHistory(
                    subject_employee_id=self.employee_id,
                    assigned_employee_id=missing_assigned_employee_id,
                    assignment_role="mentor_adaptation",
                    started_at=datetime(2026, 6, 10, 9, 0),
                    ended_at=None,
                    assigned_by_account_id=None,
                    created_at=datetime(2026, 6, 10, 9, 0),
                )
            )
            db.commit()

        response = self.client.get(f"/api/employees/{self.employee_id}")

        self.assertEqual(response.status_code, 200)
        history_items = response.json()["assignment_history"]
        self.assertEqual(len(history_items), 1)
        self.assertEqual(history_items[0]["assignment_role"], "mentor_adaptation")
        self.assertEqual(history_items[0]["role_label"], "Наставник адаптации")
        self.assertEqual(history_items[0]["assigned_employee_id"], missing_assigned_employee_id)
        self.assertEqual(history_items[0]["assigned_employee_name"], "")
        self.assertTrue(history_items[0]["is_active"])

    def test_update_employee_api_enqueues_manager_trigger_for_adaptation_assignment(self) -> None:
        scenario_key = f"manager_assign_{self.unique_tag}"
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "adaptation"
            employee.first_workday = datetime(2026, 6, 10).date()
            manager = Employee(
                full_name=f"Manager Trigger {self.unique_tag}",
                telegram_user_id="710001",
                first_workday=None,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                employee_stage="staff",
                is_manager=True,
            )
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"Manager Trigger {self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="employees",
                trigger_mode="manager_assigned_adaptation",
            )
            db.add(manager)
            db.add(scenario)
            db.commit()
            db.refresh(manager)
            manager_id = manager.id

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "API Smoke Employee",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "2026-06-10",
                "desired_position": "Аналитик",
                "birth_date": "1999-04-10",
                "work_email": "employee@example.com",
                "work_hours": "10:00-19:00",
                "is_manager": False,
                "is_mentor": False,
                "manager_employee_id": str(manager_id),
                "mentor_adaptation_employee_id": "",
                "mentor_ipr_employee_id": "",
                "adaptation_tasks_url": "",
                "adaptation_feedback_url": "",
                "adaptation_midpoint": "",
                "adaptation_end": "",
                "employee_stage": "adaptation",
                "candidate_work_stage": "",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": True,
                "is_bot_blocked": False,
                "test_task_due_at": "",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200)

        with SessionLocal() as db:
            request_row = (
                db.query(FlowLaunchRequest)
                .filter(
                    FlowLaunchRequest.employee_id == self.employee_id,
                    FlowLaunchRequest.flow_key == scenario_key,
                    FlowLaunchRequest.launch_type == "trigger",
                    FlowLaunchRequest.processed_at.is_(None),
                )
                .first()
            )
            self.assertIsNotNone(request_row)
            self.assertEqual(request_row.employee_id, self.employee_id)
            manager = db.get(Employee, manager_id)
            scenario = db.query(ScenarioTemplate).filter(ScenarioTemplate.scenario_key == scenario_key).first()
            if manager is not None:
                db.delete(manager)
            if scenario is not None:
                db.delete(scenario)
            db.commit()

    def test_employee_detail_payload_keeps_selected_manager_and_mentors_visible_when_flags_missing(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            manager = Employee(
                full_name=f"Legacy Manager {self.unique_tag}",
                telegram_user_id="720001",
                first_workday=None,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                employee_stage="staff",
                is_manager=False,
            )
            mentor_adaptation = Employee(
                full_name=f"Legacy Adapt Mentor {self.unique_tag}",
                telegram_user_id="720002",
                first_workday=None,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                employee_stage="staff",
                is_mentor=False,
            )
            mentor_ipr = Employee(
                full_name=f"Legacy IPR Mentor {self.unique_tag}",
                telegram_user_id="720003",
                first_workday=None,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                employee_stage="staff",
                is_mentor=False,
            )
            db.add_all([manager, mentor_adaptation, mentor_ipr])
            db.commit()
            db.refresh(manager)
            db.refresh(mentor_adaptation)
            db.refresh(mentor_ipr)
            employee.manager_employee_id = manager.id
            employee.mentor_adaptation_employee_id = mentor_adaptation.id
            employee.mentor_ipr_employee_id = mentor_ipr.id
            db.commit()

            payload = self.client.get(f"/api/employees/{self.employee_id}").json()
            manager_values = payload["options"]["manager_employee_values"]
            mentor_values = payload["options"]["mentor_employee_values"]

            self.assertTrue(any(option["value"] == str(manager.id) for option in manager_values))
            self.assertTrue(any(option["value"] == str(mentor_adaptation.id) for option in mentor_values))
            self.assertTrue(any(option["value"] == str(mentor_ipr.id) for option in mentor_values))

            for related_employee in (manager, mentor_adaptation, mentor_ipr):
                row = db.get(Employee, related_employee.id)
                if row is not None:
                    db.delete(row)
            db.commit()

    def test_promote_candidate_to_adaptation_api_switches_employee_stage(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "offer"
            employee.first_workday = datetime(2026, 6, 10).date()
            employee.adaptation_midpoint = None
            employee.adaptation_end = None
            db.commit()

        response = self.client.post(f"/api/employees/{self.employee_id}/promote-to-adaptation")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["employee"]["employee_stage"], "adaptation")
        self.assertEqual(payload["employee"]["candidate_work_stage"], "")
        self.assertEqual(payload["employee"]["adaptation_midpoint"], "2026-07-08")
        self.assertEqual(payload["employee"]["adaptation_end"], "2026-08-05")

        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertEqual(employee.employee_stage, "adaptation")
            self.assertIsNone(employee.candidate_work_stage)

    def test_promote_candidate_to_adaptation_requires_first_workday(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "offer"
            employee.first_workday = None
            db.commit()

        response = self.client.post(f"/api/employees/{self.employee_id}/promote-to-adaptation")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.json()["detail"],
            "Для перевода в адаптацию сначала укажите первый день сотрудника.",
        )

    def test_employee_detail_api_includes_manual_launch_history(self) -> None:
        scenario_key = f"codex_manual_launch_{self.unique_tag}"
        processed_at = datetime(2026, 6, 1, 9, 45)
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                scenario_kind="scenario",
                title=f"Manual launch {self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="manual launch history smoke",
            )
            db.add(scenario)
            db.flush()
            launch_request = FlowLaunchRequest(
                employee_id=self.employee_id,
                flow_key=scenario_key,
                requested_at=datetime(2026, 6, 1, 9, 30),
                processed_at=processed_at,
                launch_type="manual",
            )
            db.add(launch_request)
            db.commit()

        response = self.client.get(f"/api/employees/{self.employee_id}")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(len(payload["manual_launch_history"]), 1)
        launch_item = payload["manual_launch_history"][0]
        self.assertEqual(launch_item["flow_key"], scenario_key)
        self.assertEqual(launch_item["scenario_title"], f"Manual launch {self.unique_tag}")
        self.assertEqual(launch_item["processed_at_label"], "01.06.2026 09:45")

    def test_manual_bot_message_send_logs_sent_history(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "700001", db=db)
            db.commit()

        messenger = DummyMessenger()
        with (
            patch("app.web.employees.settings.TELEGRAM_BOT_TOKEN", "test-token"),
            patch("app.web.employees.create_telegram_messenger", return_value=messenger),
        ):
            response = self.client.post(
                f"/api/employees/{self.employee_id}/bot-message",
                json={"text": "  Привет из HR  "},
            )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(messenger.sent_texts, [("700001", "Привет из HR")])
        self.assertEqual(payload["manual_bot_message_history"][0]["message_text"], "Привет из HR")
        self.assertEqual(payload["manual_bot_message_history"][0]["status"], "sent")
        self.assertTrue(payload["manual_bot_message_history"][0]["sent_at"])

        with SessionLocal() as db:
            history_rows = (
                db.query(EmployeeManualBotMessage)
                .filter(EmployeeManualBotMessage.employee_id == self.employee_id)
                .order_by(EmployeeManualBotMessage.id.desc())
                .all()
            )
            self.assertEqual(len(history_rows), 1)
            self.assertEqual(history_rows[0].status, "sent")
            self.assertEqual(history_rows[0].message_text, "Привет из HR")
            self.assertIsNotNone(history_rows[0].sent_at)

    def test_manual_bot_message_without_telegram_id_logs_failed(self) -> None:
        response = self.client.post(
            f"/api/employees/{self.employee_id}/bot-message",
            json={"text": "Проверка канала"},
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "У сотрудника не указан Telegram chat id.")

        with SessionLocal() as db:
            history_rows = (
                db.query(EmployeeManualBotMessage)
                .filter(EmployeeManualBotMessage.employee_id == self.employee_id)
                .order_by(EmployeeManualBotMessage.id.desc())
                .all()
            )
            self.assertEqual(len(history_rows), 1)
            self.assertEqual(history_rows[0].status, "failed")
            self.assertEqual(history_rows[0].error_text, "У сотрудника не указан Telegram chat id.")

    def test_manual_bot_message_for_blocked_employee_logs_failed_without_send(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.is_bot_blocked = True
            set_primary_chat_id(employee, "700002", db=db)
            db.commit()

        messenger = DummyMessenger()
        with (
            patch("app.web.employees.settings.TELEGRAM_BOT_TOKEN", "test-token"),
            patch("app.web.employees.create_telegram_messenger", return_value=messenger),
        ):
            response = self.client.post(
                f"/api/employees/{self.employee_id}/bot-message",
                json={"text": "Проверка блокировки"},
            )

        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["detail"], "Для этого сотрудника доступ к боту заблокирован.")
        self.assertFalse(messenger.sent_texts)

        with SessionLocal() as db:
            history_rows = (
                db.query(EmployeeManualBotMessage)
                .filter(EmployeeManualBotMessage.employee_id == self.employee_id)
                .order_by(EmployeeManualBotMessage.id.desc())
                .all()
            )
            self.assertEqual(len(history_rows), 1)
            self.assertEqual(history_rows[0].status, "failed")
            self.assertEqual(history_rows[0].error_text, "Для этого сотрудника доступ к боту заблокирован.")

    def test_manual_bot_message_empty_text_returns_400_without_log(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "700003", db=db)
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}/bot-message",
            json={"text": "   "},
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Введите текст сообщения.")

        with SessionLocal() as db:
            history_rows = db.query(EmployeeManualBotMessage).filter(EmployeeManualBotMessage.employee_id == self.employee_id).all()
            self.assertEqual(history_rows, [])

    def test_manual_bot_message_send_exception_logs_failed(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "700004", db=db)
            db.commit()

        with (
            patch("app.web.employees.settings.TELEGRAM_BOT_TOKEN", "test-token"),
            patch("app.web.employees.create_telegram_messenger", return_value=FailingMessenger()),
        ):
            response = self.client.post(
                f"/api/employees/{self.employee_id}/bot-message",
                json={"text": "Упади, пожалуйста"},
            )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "telegram send failed")

        with SessionLocal() as db:
            history_rows = (
                db.query(EmployeeManualBotMessage)
                .filter(EmployeeManualBotMessage.employee_id == self.employee_id)
                .order_by(EmployeeManualBotMessage.id.desc())
                .all()
            )
            self.assertEqual(len(history_rows), 1)
            self.assertEqual(history_rows[0].status, "failed")
            self.assertEqual(history_rows[0].error_text, "telegram send failed")
            self.assertIsNone(history_rows[0].sent_at)

    def test_employee_detail_api_includes_manual_bot_message_history_newest_first(self) -> None:
        with SessionLocal() as db:
            older = EmployeeManualBotMessage(
                employee_id=self.employee_id,
                sender_account_id=1,
                message_text="Старое сообщение",
                status="failed",
                error_text="old error",
                sent_at=None,
                created_at=datetime(2026, 6, 1, 9, 0),
            )
            newer = EmployeeManualBotMessage(
                employee_id=self.employee_id,
                sender_account_id=1,
                message_text="Новое сообщение",
                status="sent",
                error_text=None,
                sent_at=datetime(2026, 6, 1, 10, 0),
                created_at=datetime(2026, 6, 1, 10, 0),
            )
            db.add_all([older, newer])
            db.commit()

        response = self.client.get(f"/api/employees/{self.employee_id}")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(len(payload["manual_bot_message_history"]), 2)
        self.assertEqual(payload["manual_bot_message_history"][0]["message_text"], "Новое сообщение")
        self.assertEqual(payload["manual_bot_message_history"][0]["status"], "sent")
        self.assertEqual(payload["manual_bot_message_history"][1]["message_text"], "Старое сообщение")
        self.assertEqual(payload["manual_bot_message_history"][1]["status"], "failed")

    def test_react_frontend_fetch_contract_routes_exist(self) -> None:
        route_index = {
            (getattr(route, "path", ""), method)
            for route in app.router.routes
            for method in getattr(route, "methods", set())
        }
        expected_contracts = {
            ("/login", "POST"),
            ("/api/dashboard/workspace", "GET"),
            ("/api/employees", "GET"),
            ("/api/employees", "POST"),
            ("/api/employees/{employee_id}", "GET"),
            ("/api/employees/{employee_id}", "POST"),
            ("/api/employees/{employee_id}", "DELETE"),
            ("/api/employees/{employee_id}/bot-message", "POST"),
            ("/api/employees/{employee_id}/document-links", "POST"),
            ("/api/employees/{employee_id}/document-links/{link_id}", "DELETE"),
            ("/api/employees/{employee_id}/schedule", "POST"),
            ("/api/employees/{employee_id}/schedule/{launch_request_id}", "DELETE"),
            ("/api/employees/{employee_id}/launch", "POST"),
            ("/api/employees/{employee_id}/promote-to-adaptation", "POST"),
            ("/api/employees/{employee_id}/bot-link/reset", "POST"),
            ("/api/employees/{employee_id}/files", "POST"),
            ("/api/employees/{employee_id}/files/{file_id}/send", "POST"),
            ("/api/employees/{employee_id}/files/{file_id}", "DELETE"),
            ("/api/bulk-actions/workspace", "GET"),
            ("/api/bulk-actions/preview", "POST"),
            ("/api/bulk-actions/scenarios/schedule", "POST"),
            ("/api/bulk-actions/scenarios/launch", "POST"),
            ("/api/bulk-actions/scenarios/{action_id}", "DELETE"),
            ("/api/bulk-actions/surveys/schedule", "POST"),
            ("/api/bulk-actions/surveys/launch", "POST"),
            ("/api/bulk-actions/surveys/{action_id}", "DELETE"),
            ("/api/bulk-actions/messages/schedule", "POST"),
            ("/api/bulk-actions/messages/send", "POST"),
            ("/api/bulk-actions/messages/{action_id}", "DELETE"),
            ("/api/settings/workspace", "GET"),
            ("/api/settings/hr", "POST"),
            ("/api/settings/menu-sets", "POST"),
            ("/api/settings/menu-sets/{menu_set_id}", "POST"),
            ("/api/settings/menu-sets/{menu_set_id}", "DELETE"),
            ("/api/settings/menu-sets/{menu_set_id}/buttons", "POST"),
            ("/api/settings/menu-buttons/{button_id}", "POST"),
            ("/api/settings/menu-buttons/{button_id}", "DELETE"),
            ("/api/accounts", "POST"),
            ("/api/accounts/{account_id}", "POST"),
            ("/api/accounts/{account_id}", "DELETE"),
            ("/api/flows/workspace", "GET"),
            ("/api/flows/workspace/scenarios", "POST"),
            ("/api/flows/workspace/scenarios/{scenario_id}/settings", "POST"),
            ("/api/flows/workspace/scenarios/reorder", "POST"),
            ("/api/flows/workspace/scenarios/bulk-copy", "POST"),
            ("/api/flows/workspace/scenarios/bulk-delete", "POST"),
            ("/api/flows/workspace/scenarios/{scenario_id}/steps", "POST"),
            ("/api/flows/workspace/scenarios/{scenario_id}/steps/reorder", "POST"),
            ("/api/flows/workspace/steps/{step_id}", "POST"),
            ("/api/flows/workspace/steps/{step_id}/branches", "POST"),
            ("/api/flows/workspace/steps/{step_id}/chain", "POST"),
            ("/api/flows/workspace/steps/{step_id}/delete", "POST"),
            ("/api/flows/workspace/steps/{step_id}/attachment", "POST"),
            ("/api/flows/workspace/steps/{step_id}/attachment/delete", "POST"),
        }
        missing_contracts = sorted(expected_contracts - route_index)
        self.assertEqual([], missing_contracts)

    def test_workspace_api_creates_blank_message_text_for_new_scenario_steps(self) -> None:
        scenario_key = f"codex_blank_step_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                scenario_kind="scenario",
                title=f"Blank step scenario {self.unique_tag}",
                sort_order=10,
                role_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        created_step_ids: list[int] = []
        try:
            root_response = self.client.post(
                f"/api/flows/workspace/scenarios/{scenario_id}/steps",
                json={"title": "Новый шаг"},
            )
            self.assertEqual(200, root_response.status_code)
            root_step_id = int(root_response.json()["step_id"])
            created_step_ids.append(root_step_id)

            with SessionLocal() as db:
                root_step = db.get(FlowStepTemplate, root_step_id)
                self.assertIsNotNone(root_step)
                root_step.response_type = "branching"
                root_step.button_options = "Да\nНет"
                db.commit()

            branch_response = self.client.post(
                f"/api/flows/workspace/steps/{root_step_id}/branches",
                json={"option_index": 0},
            )
            self.assertEqual(200, branch_response.status_code)
            branch_step_id = int(branch_response.json()["step_id"])
            created_step_ids.append(branch_step_id)

            with SessionLocal() as db:
                branch_step = db.get(FlowStepTemplate, branch_step_id)
                self.assertIsNotNone(branch_step)
                branch_step.response_type = "chain"
                db.commit()

            chain_response = self.client.post(
                f"/api/flows/workspace/steps/{branch_step_id}/chain",
                json={"title": "Шаг цепочки"},
            )
            self.assertEqual(200, chain_response.status_code)
            created_step_ids.append(int(chain_response.json()["step_id"]))

            with SessionLocal() as db:
                default_texts = [db.get(FlowStepTemplate, step_id).default_text for step_id in created_step_ids]
            self.assertEqual(["", "", ""], default_texts)
        finally:
            with SessionLocal() as db:
                db.query(FlowStepTemplate).filter(FlowStepTemplate.flow_key == scenario_key).delete(synchronize_session=False)
                db.query(ScenarioTemplate).filter(ScenarioTemplate.scenario_key == scenario_key).delete(synchronize_session=False)
                db.commit()

    def test_dashboard_workspace_api_returns_operational_payload(self) -> None:
        scenario_key = f"codex_dashboard_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.full_name = f"codex-dashboard-candidate-{self.unique_tag}"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                scenario_kind="scenario",
                title=f"codex-dashboard-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.add(
                EmployeeMessengerAccount(
                    employee_id=self.employee_id,
                    channel="telegram",
                    external_user_id=f"900{self.unique_tag[:6]}",
                    external_username=f"codex_dashboard_{self.unique_tag}",
                    is_primary=True,
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                )
            )
            db.add(
                EmployeeFile(
                    employee_id=self.employee_id,
                    direction="inbound",
                    category="candidate_file",
                    telegram_file_id=None,
                    telegram_file_unique_id=None,
                    original_filename=f"codex-dashboard-{self.unique_tag}.pdf",
                    stored_path=str((Path("storage") / "employee_files" / f"codex-dashboard-{self.unique_tag}.pdf").resolve()),
                    mime_type="application/pdf",
                    file_size=12,
                    created_at=now,
                )
            )
            db.add(
                FlowLaunchRequest(
                    employee_id=self.employee_id,
                    flow_key=scenario_key,
                    requested_at=now + timedelta(days=1),
                    processed_at=None,
                    launch_type="scheduled",
                )
            )
            db.commit()

        response = self.client.get("/api/dashboard/workspace")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertIn("stats", payload)
        self.assertIn("upcoming_events", payload)
        self.assertIn("telegram_links", payload)
        self.assertIn("inbound_files", payload)
        self.assertIn("attention_items", payload)
        self.assertIn("module_links", payload)
        self.assertGreaterEqual(payload["stats"]["recent_telegram_links"], 1)
        self.assertTrue(any(item["employee_id"] == self.employee_id for item in payload["telegram_links"]))
        self.assertTrue(any(item["employee_id"] == self.employee_id for item in payload["inbound_files"]))
        self.assertTrue(any(item["title"] == f"codex-dashboard-{self.unique_tag}" for item in payload["upcoming_events"]))

    def test_dashboard_is_default_authenticated_entry(self) -> None:
        root_response = self.client.get("/", follow_redirects=False)
        login_response = self.client.get("/login", follow_redirects=False)

        self.assertEqual(root_response.status_code, 303)
        self.assertEqual(root_response.headers.get("location"), "/app/dashboard")
        self.assertEqual(login_response.status_code, 303)
        self.assertEqual(login_response.headers.get("location"), "/app/dashboard")

    def test_raw_account_id_cookie_does_not_authenticate(self) -> None:
        client = TestClient(app)
        client.cookies.set(AUTH_COOKIE_NAME, "1")

        response = client.get("/api/settings/workspace")

        self.assertEqual(response.status_code, 401)

    def test_login_sets_signed_session_cookie(self) -> None:
        client = TestClient(app)

        response = client.post(
            "/login",
            data={"login": "admin", "password": "admin123"},
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers.get("location"), "/app/dashboard")
        session_cookie = response.cookies.get(AUTH_COOKIE_NAME)
        self.assertIsNotNone(session_cookie)
        self.assertNotEqual(session_cookie, "1")
        self.assertEqual(len(session_cookie.split(".")), 3)

    def test_account_api_rejects_weak_passwords(self) -> None:
        response = self.client.post(
            "/api/accounts",
            json={
                "login": f"codex-weak-{self.unique_tag}",
                "password": "short",
                "role": "hr",
                "is_active": True,
            },
        )

        self.assertEqual(response.status_code, 400)

    def test_account_api_accepts_strong_passwords(self) -> None:
        login = f"codex-strong-{self.unique_tag}"

        response = self.client.post(
            "/api/accounts",
            json={
                "login": login,
                "password": f"Strong-{self.unique_tag}-2026",
                "role": "hr",
                "is_active": True,
            },
        )

        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            account = db.query(AdminAccount).filter(AdminAccount.login == login).first()
            self.assertIsNotNone(account)
            if account is not None:
                self.assertNotEqual(account.password_hash, f"Strong-{self.unique_tag}-2026")
                db.delete(account)
                db.commit()

    def test_react_dashboard_template_mounts_bundle(self) -> None:
        response = self.client.get("/app/dashboard")

        self.assertEqual(response.status_code, 200)
        self.assertIn('id="react-dashboard-root"', response.text)
        self.assertIn('data-api-url="/api/dashboard/workspace"', response.text)
        self.assertIn("/static/workspace_v2/dashboard.js", response.text)

    def test_react_bot_menu_template_mounts_bundle(self) -> None:
        response = self.client.get("/app/bot-menu")

        self.assertEqual(response.status_code, 200)
        self.assertIn('id="react-bot-menu-root"', response.text)
        self.assertIn('data-api-url="/api/settings/workspace"', response.text)
        self.assertIn("/static/workspace_v2/bot-menu.js", response.text)

    def test_classic_operator_entrypoints_redirect_to_react_surfaces(self) -> None:
        expected_redirects = {
            "/candidates": "/app/employees?list_kind=candidates",
            "/employees": "/app/employees",
            "/bulk-actions": "/app/bulk-actions",
            "/flows": "/app/flows/workspace-v2",
            "/surveys": "/app/surveys/workspace",
            "/settings": "/app/settings",
        }

        for path, expected_location in expected_redirects.items():
            with self.subTest(path=path):
                response = self.client.get(path, follow_redirects=False)
                self.assertEqual(response.status_code, 303)
                self.assertEqual(response.headers.get("location"), expected_location)

    def test_classic_scenario_editor_route_is_served_via_router(self) -> None:
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=f"codex_scenario_{self.unique_tag}",
                scenario_kind="scenario",
                title=f"codex-scenario-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="router smoke",
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        response = self.client.get(f"/flows/{scenario_id}?legacy=1")

        self.assertEqual(response.status_code, 200)
        self.assertIn(f"codex-scenario-{self.unique_tag}", response.text)

    def test_classic_scenario_get_redirects_to_react_workspace_by_default(self) -> None:
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=f"codex_scenario_redirect_{self.unique_tag}",
                scenario_kind="scenario",
                title=f"codex-scenario-redirect-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="redirect smoke",
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        response = self.client.get(f"/flows/{scenario_id}", follow_redirects=False)

        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers.get("location"), f"/app/flows/workspace-v2?scenario_id={scenario_id}")

    def test_classic_scenario_legacy_update_stays_in_legacy_editor(self) -> None:
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=f"codex_scenario_legacy_update_{self.unique_tag}",
                scenario_kind="scenario",
                title=f"codex-scenario-legacy-update-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="legacy update smoke",
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        response = self.client.post(
            f"/flows/{scenario_id}?legacy=1",
            data={
                "title": f"codex-scenario-legacy-update-{self.unique_tag}",
                "role_scope": "all",
                "employee_scope": "all",
                "target_employee_id": "",
                "trigger_mode": "manual_only",
                "description": "legacy update smoke",
            },
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers.get("location"), f"/flows/{scenario_id}?legacy=1")

    def test_classic_survey_get_redirects_to_react_workspace_by_default(self) -> None:
        with SessionLocal() as db:
            survey = ScenarioTemplate(
                scenario_key=f"codex_survey_redirect_{self.unique_tag}",
                scenario_kind="survey",
                title=f"codex-survey-redirect-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="redirect smoke",
            )
            db.add(survey)
            db.commit()
            db.refresh(survey)
            survey_id = survey.id

        response = self.client.get(f"/surveys/{survey_id}", follow_redirects=False)

        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers.get("location"), f"/app/surveys/workspace?scenario_id={survey_id}")

    def test_classic_attachment_delete_stays_in_legacy_editor_when_requested(self) -> None:
        scenario_key = f"codex_scenario_attachment_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                scenario_kind="scenario",
                title=f"codex-scenario-attachment-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="legacy attachment smoke",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с файлом",
                sort_order=10,
                default_text="Текст",
                custom_text=None,
                response_type="none",
                button_options=None,
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add(step)
            db.commit()
            db.refresh(scenario)
            db.refresh(step)
            scenario_id = scenario.id
            step_id = step.id

        response = self.client.post(
            f"/flows/steps/{step_id}/attachment/delete?legacy=1",
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers.get("location"), f"/flows/{scenario_id}?legacy=1")

    def test_classic_survey_export_route_returns_xlsx(self) -> None:
        with SessionLocal() as db:
            survey = ScenarioTemplate(
                scenario_key=f"codex_survey_{self.unique_tag}",
                scenario_kind="survey",
                title=f"codex-survey-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="export smoke",
            )
            db.add(survey)
            db.flush()
            step = FlowStepTemplate(
                flow_key=survey.scenario_key,
                step_key=f"{survey.scenario_key}_step_1",
                step_title="Как дела?",
                sort_order=10,
                default_text="Как дела?",
                custom_text="Как дела?",
                response_type="text",
                button_options=None,
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add(step)
            db.flush()
            db.add(
                SurveyAnswer(
                    employee_id=self.employee_id,
                    scenario_key=survey.scenario_key,
                    step_key=step.step_key,
                    answer_value="Хорошо",
                    file_name=None,
                    answered_at=datetime(2026, 6, 9, 12, 0),
                )
            )
            db.commit()
            db.refresh(survey)
            survey_id = survey.id

        response = self.client.get(f"/surveys/{survey_id}/export")

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.headers.get("content-type", ""),
        )
        workbook = load_workbook(filename=BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.title, "Результаты")
        self.assertEqual(
            [sheet["A1"].value, sheet["B1"].value, sheet["C1"].value],
            ["Пользователь ФИО", "Вопрос", "Ответ"],
        )
        self.assertEqual(sheet["A2"].value, "API Smoke Employee")
        self.assertEqual(sheet["B2"].value, "Как дела?")
        self.assertEqual(sheet["C2"].value, "Хорошо")

    def test_react_scenario_workspace_template_exposes_flash_attrs(self) -> None:
        response = self.client.get(
            f"/app/flows/workspace-v2?flash_message=scenario-flash-{self.unique_tag}&flash_type=error"
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(f'data-flash-message="scenario-flash-{self.unique_tag}"', response.text)
        self.assertIn('data-flash-type="error"', response.text)

    def test_classic_scenario_create_redirects_to_react_workspace(self) -> None:
        response = self.client.post(
            "/flows",
            data={"title": f"codex-scenario-{self.unique_tag}"},
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 303)
        location = response.headers.get("location", "")
        self.assertTrue(location.startswith("/app/flows/workspace-v2?"))
        self.assertIn("scenario_id=", location)
        self.assertIn("%D0%A1%D1%86%D0%B5%D0%BD%D0%B0%D1%80%D0%B8%D0%B9+%D1%81%D0%BE%D0%B7%D0%B4%D0%B0%D0%BD", location)

    def test_classic_scenario_copy_redirects_to_react_workspace(self) -> None:
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=f"codex_copy_scenario_{self.unique_tag}",
                scenario_kind="scenario",
                title=f"codex-copy-scenario-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="copy redirect smoke",
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        response = self.client.post(f"/flows/{scenario_id}/copy", follow_redirects=False)

        self.assertEqual(response.status_code, 303)
        location = response.headers.get("location", "")
        self.assertTrue(location.startswith("/app/flows/workspace-v2?"))
        self.assertIn("scenario_id=", location)
        self.assertIn("%D0%A1%D1%86%D0%B5%D0%BD%D0%B0%D1%80%D0%B8%D0%B9+%D1%81%D0%BA%D0%BE%D0%BF%D0%B8%D1%80%D0%BE%D0%B2%D0%B0%D0%BD", location)

    def test_classic_survey_delete_redirects_to_react_workspace(self) -> None:
        with SessionLocal() as db:
            survey = ScenarioTemplate(
                scenario_key=f"codex_delete_survey_{self.unique_tag}",
                scenario_kind="survey",
                title=f"codex-delete-survey-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="delete redirect smoke",
            )
            db.add(survey)
            db.commit()
            db.refresh(survey)
            survey_id = survey.id

        response = self.client.post(f"/surveys/{survey_id}/delete", follow_redirects=False)

        self.assertEqual(response.status_code, 303)
        location = response.headers.get("location", "")
        self.assertTrue(location.startswith("/app/surveys/workspace?"))
        self.assertIn("%D0%9E%D0%BF%D1%80%D0%BE%D1%81+%D1%83%D0%B4%D0%B0%D0%BB%D1%91%D0%BD", location)

    def test_react_employee_detail_template_exposes_flash_attrs(self) -> None:
        response = self.client.get(
            f"/app/employees/{self.employee_id}?flash_message=flash-smoke-{self.unique_tag}&flash_type=error"
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(f'data-flash-message="flash-smoke-{self.unique_tag}"', response.text)
        self.assertIn('data-flash-type="error"', response.text)

    def test_classic_employee_edit_route_redirects_to_react_detail(self) -> None:
        response = self.client.get(f"/employees/{self.employee_id}/edit", follow_redirects=False)

        self.assertEqual(response.status_code, 303)
        self.assertEqual(response.headers.get("location"), f"/app/employees/{self.employee_id}")

    def test_classic_employee_action_redirects_back_to_react_detail(self) -> None:
        response = self.client.post(
            f"/employees/{self.employee_id}/profile-photo/delete",
            follow_redirects=False,
        )

        self.assertEqual(response.status_code, 303)
        self.assertTrue(response.headers.get("location", "").startswith(f"/app/employees/{self.employee_id}?"))

    def test_update_employee_api_accepts_public_chat_handle_without_chat_id(self) -> None:
        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "Updated API Smoke Employee",
                "chat_id": "",
                "chat_handle": "hr_team",
                "first_workday": "",
                "desired_position": "",
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_chat_id": "",
                "mentor_adaptation_chat_id": "",
                "mentor_ipr_chat_id": "",
                "employee_stage": "candidate",
                "candidate_work_stage": "testing",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": False,
                "test_task_due_at": "",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["employee"]["chat_id"], "")
        self.assertEqual(payload["employee"]["chat_handle"], "hr_team")

        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertEqual(employee.telegram_user_id, None)
            self.assertEqual(employee.telegram_username, "hr_team")

    def test_update_employee_api_preserves_chat_id_when_payload_omits_it(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "777000111", db=db)
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "Updated API Smoke Employee",
                "chat_handle": "hr_team",
                "first_workday": "",
                "desired_position": "",
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_chat_id": "",
                "mentor_adaptation_chat_id": "",
                "mentor_ipr_chat_id": "",
                "employee_stage": "candidate",
                "candidate_work_stage": "testing",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": False,
                "test_task_due_at": "",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["employee"]["chat_id"], "777000111")
        self.assertEqual(payload["employee"]["chat_handle"], "hr_team")

        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertEqual(employee.telegram_user_id, "777000111")
            self.assertEqual(employee.telegram_username, "hr_team")

    def test_bot_start_links_existing_employee_by_public_username(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            unique_suffix = uuid4().hex[:12]
            username = f"codex_link_{unique_suffix}"
            chat_id = str(900000000000 + (int(unique_suffix, 16) % 100000000000))
            employee.telegram_username = f"@{username}"
            db.commit()

            employee, created = get_or_create_employee_by_chat(db, chat_id, username.upper())

            self.assertFalse(created)
            self.assertEqual(employee.id, self.employee_id)
            self.assertEqual(get_primary_chat_id(employee, db=db), chat_id)
            self.assertEqual(employee.telegram_username, username.upper())
            self.assertEqual(db.query(Employee).filter(Employee.id == self.employee_id).count(), 1)

    def test_bot_start_creates_candidate_when_card_is_missing(self) -> None:
        with SessionLocal() as db:
            unique_suffix = uuid4().hex[:12]
            chat_id = str(910000000000 + (int(unique_suffix, 16) % 100000000000))
            before_count = db.query(Employee).count()

            employee, created = get_or_create_employee_by_chat(db, chat_id, f"new_candidate_{unique_suffix}")

            self.assertTrue(created)
            self.assertIsNotNone(employee)
            self.assertEqual(db.query(Employee).count(), before_count + 1)
            self.assertEqual(employee.employee_stage, "candidate")
            self.assertEqual(employee.candidate_work_stage, self._initial_candidate_stage_key())
            self.assertEqual(get_primary_chat_id(employee, db=db), chat_id)
            self.assertEqual(employee.telegram_username, f"new_candidate_{unique_suffix}")

            db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == employee.id).delete(
                synchronize_session=False
            )
            db.delete(employee)
            db.commit()

    def test_update_employee_api_persists_shared_fields_for_candidate(self) -> None:
        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "Кандидат Тест",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "2026-05-20",
                "desired_position": "Дизайнер",
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_chat_id": "",
                "mentor_adaptation_chat_id": "",
                "mentor_ipr_chat_id": "",
                "employee_stage": "candidate",
                "candidate_work_stage": "offer",
                "salary_expectation": "300000",
                "personal_data_consent": True,
                "employee_data_consent": False,
                "is_bot_blocked": False,
                "test_task_due_at": "",
                "notes": "candidate-notes",
            },
        )

        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertEqual(employee.first_workday.isoformat() if employee.first_workday else None, "2026-05-20")
            self.assertEqual(employee.desired_position, "Дизайнер")
            self.assertEqual(employee.salary_expectation, "300000")
            self.assertEqual(employee.notes, "candidate-notes")
            self.assertFalse(employee.is_bot_blocked)

    def test_update_employee_api_clears_first_workday_with_empty_value(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.first_workday = datetime(2026, 5, 20).date()
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "Кандидат без даты",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "",
                "desired_position": "Аналитик",
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_chat_id": "",
                "mentor_adaptation_chat_id": "",
                "mentor_ipr_chat_id": "",
                "employee_stage": "candidate",
                "candidate_work_stage": "testing",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": False,
                "is_bot_blocked": False,
                "test_task_due_at": "",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["employee"]["first_workday"], "")
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertIsNone(employee.first_workday)

    def test_update_employee_api_persists_shared_fields_for_employee(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "staff"
            employee.candidate_work_stage = None
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "Сотрудник Тест",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "2026-05-21",
                "desired_position": "Аналитик",
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_chat_id": "",
                "mentor_adaptation_chat_id": "",
                "mentor_ipr_chat_id": "",
                "employee_stage": "staff",
                "candidate_work_stage": "",
                "salary_expectation": "350000",
                "personal_data_consent": False,
                "employee_data_consent": True,
                "is_bot_blocked": True,
                "test_task_due_at": "",
                "notes": "employee-notes",
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["employee"]["is_bot_blocked"])
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertEqual(employee.first_workday.isoformat() if employee.first_workday else None, "2026-05-21")
            self.assertEqual(employee.desired_position, "Аналитик")
            self.assertEqual(employee.salary_expectation, "350000")
            self.assertEqual(employee.notes, "employee-notes")
            self.assertTrue(employee.is_bot_blocked)

    def test_settings_positions_api_supports_crud(self) -> None:
        position_slug = f"qa_engineer_{self.unique_tag}"
        position_title = f"QA engineer {self.unique_tag}"
        create_response = self.client.post(
            "/api/settings/positions",
            json={
                "title": position_title,
                "slug": position_slug,
            },
        )

        self.assertEqual(create_response.status_code, 200)
        workspace = create_response.json()
        created_position = next((item for item in workspace["positions"] if item["slug"] == position_slug), None)
        self.assertIsNotNone(created_position)
        self.assertEqual(created_position["title"], position_title)

        list_response = self.client.get("/api/settings/positions")
        self.assertEqual(list_response.status_code, 200)
        self.assertTrue(any(item["slug"] == position_slug for item in list_response.json()["positions"]))

        update_response = self.client.patch(
            f"/api/settings/positions/{created_position['id']}",
            json={"title": "Senior QA engineer", "sort_order": 90},
        )
        self.assertEqual(update_response.status_code, 200)
        updated_position = next(
            (item for item in update_response.json()["positions"] if item["id"] == created_position["id"]),
            None,
        )
        self.assertIsNotNone(updated_position)
        self.assertEqual(updated_position["title"], "Senior QA engineer")
        self.assertEqual(updated_position["sort_order"], 90)

        delete_response = self.client.delete(f"/api/settings/positions/{created_position['id']}")
        self.assertEqual(delete_response.status_code, 200)
        deleted_position = next(
            (item for item in delete_response.json()["positions"] if item["id"] == created_position["id"]),
            None,
        )
        self.assertIsNotNone(deleted_position)
        self.assertFalse(deleted_position["is_active"])

    def test_employee_detail_payload_uses_catalog_positions_and_preserves_legacy_value(self) -> None:
        position_slug = f"qa_engineer_{self.unique_tag}"
        position_title = f"QA engineer {self.unique_tag}"
        with SessionLocal() as db:
            ensure_position = Position(
                title=position_title,
                slug=position_slug,
                is_active=True,
                sort_order=90,
                created_at=datetime.now(UTC).replace(tzinfo=None),
            )
            db.add(ensure_position)
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.desired_position = "Редактор"
            db.commit()

        response = self.client.get(f"/api/employees/{self.employee_id}")
        self.assertEqual(response.status_code, 200)
        options = response.json()["options"]["employee_role_values"]
        self.assertIn(position_title, options)
        self.assertIn("Редактор", options)

    def test_update_employee_api_persists_position_from_catalog(self) -> None:
        position_slug = f"qa_engineer_{self.unique_tag}"
        position_title = f"QA engineer {self.unique_tag}"
        self.client.post(
            "/api/settings/positions",
            json={
                "title": position_title,
                "slug": position_slug,
            },
        )

        response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "Кандидат QA",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "",
                "desired_position": position_slug,
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_chat_id": "",
                "mentor_adaptation_chat_id": "",
                "mentor_ipr_chat_id": "",
                "employee_stage": "candidate",
                "candidate_work_stage": "testing",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": False,
                "is_bot_blocked": False,
                "test_task_due_at": "",
                "notes": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["employee"]["desired_position"], position_title)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertEqual(employee.desired_position, position_title)

    def test_workspace_scenario_settings_api_updates_scope_and_description(self) -> None:
        scenario_key = f"codex_settings_{uuid4().hex[:12]}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title="Workspace Settings Smoke",
                sort_order=999999,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                target_employee_id=None,
                description=None,
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        try:
            response = self.client.post(
                f"/api/flows/workspace/scenarios/{scenario_id}/settings",
                json={
                    "title": "Новый заголовок сценария",
                    "description": "x" * 60,
                    "role_scope": "analyst",
                    "employee_scope": "employees",
                    "trigger_mode": "candidate_hr_stage",
                    "candidate_work_stage_trigger": "offer",
                    "target_employee_id": str(self.employee_id),
                },
            )

            self.assertEqual(response.status_code, 200)
            scenario_payload = response.json()["payload"]["workspace"]["scenario"]
            self.assertEqual(scenario_payload["title"], "Новый заголовок сценария")
            self.assertEqual(scenario_payload["description"], "x" * 50)
            self.assertEqual(scenario_payload["role_scope"], "analyst")
            self.assertEqual(scenario_payload["role_scopes"], ["analyst"])
            self.assertEqual(scenario_payload["employee_scope"], "employees")
            self.assertEqual(scenario_payload["trigger_mode"], "candidate_hr_stage")
            self.assertEqual(scenario_payload["candidate_work_stage_trigger"], "offer")
            self.assertEqual(scenario_payload["target_employee_id"], self.employee_id)
            scenario_summary = next(
                item for item in response.json()["payload"]["scenarios"] if item["id"] == scenario_id
            )
            self.assertIn("created_at", scenario_summary)
            self.assertIn("updated_at", scenario_summary)
            self.assertEqual(scenario_summary["candidate_work_stage_trigger"], "offer")
            self.assertIn("candidate_hr_stage", response.json()["payload"]["workspace"]["trigger_mode_labels"])
            self.assertEqual(
                response.json()["payload"]["workspace"]["candidate_work_stage_labels"]["manager_interview"],
                "Собеседование с руководителем",
            )
        finally:
            with SessionLocal() as db:
                scenario = db.get(ScenarioTemplate, scenario_id)
                if scenario is not None:
                    db.delete(scenario)
                db.commit()

    def test_workspace_scenario_settings_api_supports_multiple_role_scopes(self) -> None:
        scenario_key = f"codex_multi_scope_{uuid4().hex[:12]}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title="Workspace Multi Scope Smoke",
                sort_order=999999,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                target_employee_id=None,
                description=None,
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        try:
            response = self.client.post(
                f"/api/flows/workspace/scenarios/{scenario_id}/settings",
                json={
                    "title": "Multi scope scenario",
                    "description": "multi positions",
                    "role_scopes": ["designer", "analyst"],
                    "employee_scope": "employees",
                    "trigger_mode": "manual_only",
                    "candidate_work_stage_trigger": "offer",
                    "target_employee_id": "",
                },
            )

            self.assertEqual(response.status_code, 200)
            scenario_payload = response.json()["payload"]["workspace"]["scenario"]
            self.assertEqual(scenario_payload["role_scope"], "designer,analyst")
            self.assertEqual(scenario_payload["role_scopes"], ["designer", "analyst"])
            self.assertEqual(scenario_payload["role_scope_label"], "Дизайнер, Аналитик")
            self.assertEqual(scenario_payload["candidate_work_stage_trigger"], "")
            scenario_summary = next(
                item for item in response.json()["payload"]["scenarios"] if item["id"] == scenario_id
            )
            self.assertEqual(scenario_summary["role_scope"], "designer,analyst")
            self.assertEqual(scenario_summary["role_scopes"], ["designer", "analyst"])
            self.assertEqual(scenario_summary["role_scope_label"], "Дизайнер, Аналитик")

            with SessionLocal() as db:
                scenario = db.get(ScenarioTemplate, scenario_id)
                self.assertIsNotNone(scenario)
                self.assertEqual(scenario.role_scope, "designer,analyst")
        finally:
            with SessionLocal() as db:
                scenario = db.get(ScenarioTemplate, scenario_id)
                if scenario is not None:
                    db.delete(scenario)
                db.commit()

    def test_candidate_stage_update_queues_status_transition_launch_once(self) -> None:
        scenario_key = f"codex_hr_stage_{self.unique_tag}"
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "hr_interview"
            employee.desired_position = "Аналитик"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"Offer trigger {self.unique_tag}",
                sort_order=15,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="candidates",
                trigger_mode="candidate_hr_stage",
                candidate_work_stage_trigger="offer",
                target_employee_id=None,
                description="status trigger smoke",
            )
            db.add(scenario)
            db.commit()

        payload = {
            "full_name": "API Smoke Employee",
            "chat_id": "",
            "chat_handle": "",
            "first_workday": "",
            "desired_position": "Аналитик",
            "birth_date": "",
            "work_email": "",
            "work_hours": "",
            "manager_employee_id": "",
            "mentor_adaptation_employee_id": "",
            "mentor_ipr_employee_id": "",
            "adaptation_tasks_url": "",
            "adaptation_feedback_url": "",
            "adaptation_midpoint": "",
            "adaptation_end": "",
            "employee_stage": "candidate",
            "candidate_work_stage": "offer",
            "salary_expectation": "",
            "personal_data_consent": False,
            "employee_data_consent": False,
            "is_bot_blocked": False,
            "test_task_due_at": "",
            "notes": "",
        }

        first_response = self.client.post(f"/api/employees/{self.employee_id}", json=payload)
        self.assertEqual(first_response.status_code, 200)
        candidate_stage_values = {
            item["value"]: item["label"]
            for item in first_response.json()["options"]["candidate_work_stage_values"]
        }
        self.assertEqual(candidate_stage_values["offer"], "Оффер")
        self.assertNotIn("contract", candidate_stage_values)

        second_response = self.client.post(f"/api/employees/{self.employee_id}", json=payload)
        self.assertEqual(second_response.status_code, 200)

        with SessionLocal() as db:
            queued_requests = (
                db.query(FlowLaunchRequest)
                .filter(
                    FlowLaunchRequest.employee_id == self.employee_id,
                    FlowLaunchRequest.flow_key == scenario_key,
                    FlowLaunchRequest.launch_type == "status_transition",
                )
                .all()
            )
            self.assertEqual(len(queued_requests), 1)
            self.assertIsNone(queued_requests[0].processed_at)

    def test_candidate_stage_update_queues_expected_hr_status_triggers(self) -> None:
        expected_stage_triggers = [
            ("company_decline", "Наш отказ кандидату"),
            ("hr_interview", "Собеседование с HR"),
            ("testing", "Тестирование"),
            ("offer", "Оффер"),
            ("preonboarding", "Преонбординг"),
        ]
        scenario_keys = []
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = None
            employee.desired_position = "Аналитик"
            for stage_key, title in expected_stage_triggers:
                scenario_key = f"codex_{stage_key}_{self.unique_tag}"
                scenario_keys.append(scenario_key)
                db.add(
                    ScenarioTemplate(
                        scenario_key=scenario_key,
                        title=f"codex-{title}-{self.unique_tag}",
                        sort_order=15,
                        scenario_kind="scenario",
                        role_scope="all",
                        employee_scope="candidates",
                        trigger_mode="candidate_hr_stage",
                        candidate_work_stage_trigger=stage_key,
                        target_employee_id=None,
                    )
                )
            db.commit()

        for stage_key, _title in expected_stage_triggers:
            response = self.client.post(
                f"/api/employees/{self.employee_id}",
                json={
                    "full_name": "API Smoke Employee",
                    "chat_id": "",
                    "chat_handle": "",
                    "first_workday": "",
                    "desired_position": "Аналитик",
                    "birth_date": "",
                    "work_email": "",
                    "work_hours": "",
                    "manager_employee_id": "",
                    "mentor_adaptation_employee_id": "",
                    "mentor_ipr_employee_id": "",
                    "adaptation_tasks_url": "",
                    "adaptation_feedback_url": "",
                    "adaptation_midpoint": "",
                    "adaptation_end": "",
                    "employee_stage": "candidate",
                    "candidate_work_stage": stage_key,
                    "salary_expectation": "",
                    "personal_data_consent": False,
                    "employee_data_consent": False,
                    "is_bot_blocked": False,
                    "test_task_due_at": "",
                    "notes": "",
                },
            )
            self.assertEqual(response.status_code, 200)
            with SessionLocal() as db:
                employee = db.get(Employee, self.employee_id)
                self.assertIsNotNone(employee)
                employee.candidate_work_stage = None
                db.commit()

        with SessionLocal() as db:
            queued_flow_keys = {
                row.flow_key
                for row in db.query(FlowLaunchRequest)
                .filter(
                    FlowLaunchRequest.employee_id == self.employee_id,
                    FlowLaunchRequest.flow_key.in_(scenario_keys),
                    FlowLaunchRequest.launch_type == "status_transition",
                    FlowLaunchRequest.processed_at.is_(None),
                )
                .all()
            }

        self.assertEqual(queued_flow_keys, set(scenario_keys))

    def test_candidate_stage_trigger_settings_accept_stage_label_and_match_key(self) -> None:
        scenario_key = f"codex_label_stage_{self.unique_tag}"
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "hr_interview"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-label-stage-{self.unique_tag}",
                sort_order=15,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="candidates",
                trigger_mode="candidate_hr_stage",
                candidate_work_stage_trigger="manual",
                target_employee_id=None,
            )
            db.add(scenario)
            db.commit()
            db.refresh(scenario)
            scenario_id = scenario.id

        settings_response = self.client.post(
            f"/api/flows/workspace/scenarios/{scenario_id}/settings",
            json={
                "title": f"codex-label-stage-{self.unique_tag}",
                "description": "",
                "role_scope": "all",
                "employee_scope": "candidates",
                "trigger_mode": "candidate_hr_stage",
                "candidate_work_stage_trigger": "Оффер",
                "target_employee_id": "",
            },
        )

        self.assertEqual(settings_response.status_code, 200)
        self.assertEqual(settings_response.json()["payload"]["workspace"]["scenario"]["candidate_work_stage_trigger"], "offer")

        update_response = self.client.post(
            f"/api/employees/{self.employee_id}",
            json={
                "full_name": "API Smoke Employee",
                "chat_id": "",
                "chat_handle": "",
                "first_workday": "",
                "desired_position": "Аналитик",
                "birth_date": "",
                "work_email": "",
                "work_hours": "",
                "manager_employee_id": "",
                "mentor_adaptation_employee_id": "",
                "mentor_ipr_employee_id": "",
                "adaptation_tasks_url": "",
                "adaptation_feedback_url": "",
                "adaptation_midpoint": "",
                "adaptation_end": "",
                "employee_stage": "candidate",
                "candidate_work_stage": "offer",
                "salary_expectation": "",
                "personal_data_consent": False,
                "employee_data_consent": False,
                "is_bot_blocked": False,
                "test_task_due_at": "",
                "notes": "",
            },
        )

        self.assertEqual(update_response.status_code, 200)
        with SessionLocal() as db:
            queued = (
                db.query(FlowLaunchRequest)
                .filter(
                    FlowLaunchRequest.employee_id == self.employee_id,
                    FlowLaunchRequest.flow_key == scenario_key,
                    FlowLaunchRequest.launch_type == "status_transition",
                    FlowLaunchRequest.processed_at.is_(None),
                )
                .one_or_none()
            )
            self.assertIsNotNone(queued)

    def test_employee_apis_hide_internal_followup_launch_requests(self) -> None:
        visible_flow_key = f"codex_visible_launch_{self.unique_tag}"
        internal_flow_key = f"codex_internal_launch_{self.unique_tag}"
        with SessionLocal() as db:
            visible_scenario = ScenarioTemplate(
                scenario_key=visible_flow_key,
                title=f"Visible launch {self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            internal_scenario = ScenarioTemplate(
                scenario_key=internal_flow_key,
                title=f"Internal launch {self.unique_tag}",
                sort_order=20,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add_all([visible_scenario, internal_scenario])
            db.flush()
            db.add_all(
                [
                    FlowLaunchRequest(
                        employee_id=self.employee_id,
                        flow_key=visible_flow_key,
                        requested_at=datetime(2026, 6, 1, 9, 30),
                        processed_at=None,
                        launch_type="scheduled",
                        skip_step_key=None,
                    ),
                    FlowLaunchRequest(
                        employee_id=self.employee_id,
                        flow_key=internal_flow_key,
                        requested_at=datetime(2026, 6, 1, 9, 45),
                        processed_at=None,
                        launch_type="scheduled",
                        skip_step_key=f"{SINGLE_STEP_REQUEST_PREFIX}step_two",
                    ),
                ]
            )
            db.commit()

        detail_response = self.client.get(f"/api/employees/{self.employee_id}")
        list_response = self.client.get("/api/employees?list_kind=candidates")

        self.assertEqual(detail_response.status_code, 200)
        detail_payload = detail_response.json()
        self.assertEqual(len(detail_payload["scheduled_launches"]), 1)
        self.assertEqual(detail_payload["scheduled_launches"][0]["flow_key"], visible_flow_key)

        self.assertEqual(list_response.status_code, 200)
        list_payload = list_response.json()
        employee_item = next(item for item in list_payload["items"] if item["id"] == self.employee_id)
        self.assertEqual(employee_item["planned_scenario_title"], f"Visible launch {self.unique_tag}")

    def test_workspace_step_api_updates_button_notifications(self) -> None:
        scenario_key = f"codex_button_notify_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-button-notify-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с кнопками",
                sort_order=10,
                default_text="Нажми кнопку",
                custom_text=None,
                response_type="buttons",
                button_options="Да\nНет",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add(step)
            db.commit()
            db.refresh(step)
            step_id = step.id

        response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Шаг с кнопками",
                "text": "Нажми кнопку",
                "response_type": "buttons",
                "button_options": "Да\nНет",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "",
                "launch_scenario_key": "",
                "send_employee_card": False,
                "notify_on_send_text": "",
                "notify_on_send_recipient_ids": "",
                "notify_on_send_recipient_scope": "",
                "button_notifications": [
                    {
                        "option_index": 0,
                        "rules": [
                            {
                                "rule_index": 0,
                                "message_text": "Нажали Да",
                                "recipient_ids": "mentor_ipr",
                                "recipient_scope": "",
                            },
                            {
                                "rule_index": 1,
                                "message_text": "Дублирующее уведомление для HR",
                                "recipient_ids": "manager",
                                "recipient_scope": "",
                            },
                        ],
                    },
                    {
                        "option_index": 1,
                        "rules": [],
                    },
                ],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(payload["button_notifications"][0]["message_text"], "Нажали Да")
        self.assertEqual(payload["button_notifications"][0]["recipient_ids"], "mentor_ipr")
        self.assertEqual(len(payload["button_notifications"][0]["rules"]), 2)
        self.assertEqual(payload["button_notifications"][0]["rules"][1]["message_text"], "Дублирующее уведомление для HR")

        with SessionLocal() as db:
            notifications = (
                db.query(StepButtonNotification)
                .filter(StepButtonNotification.step_id == step_id)
                .order_by(StepButtonNotification.option_index.asc(), StepButtonNotification.rule_index.asc())
                .all()
            )
            self.assertEqual(len(notifications), 2)
            self.assertEqual(notifications[0].option_index, 0)
            self.assertEqual(notifications[0].rule_index, 0)
            self.assertEqual(notifications[0].message_text, "Нажали Да")
            self.assertEqual(notifications[0].recipient_ids, "mentor_ipr")
            self.assertIsNone(notifications[0].recipient_scope)
            self.assertEqual(notifications[1].rule_index, 1)
            self.assertEqual(notifications[1].message_text, "Дублирующее уведомление для HR")
            self.assertEqual(notifications[1].recipient_ids, "manager")

    def test_workspace_step_api_preserves_buttons_target_field(self) -> None:
        scenario_key = f"codex_button_target_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-button-target-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с выбором дохода",
                sort_order=10,
                default_text="Выбери ожидания по доходу",
                custom_text=None,
                response_type="buttons",
                button_options="100 000\n150 000\n200 000",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field="salary_expectation",
                send_employee_card=False,
            )
            db.add(step)
            db.commit()
            db.refresh(step)
            step_id = step.id

        response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Шаг с выбором дохода",
                "text": "Выбери ожидания по доходу",
                "response_type": "buttons",
                "button_options": "100 000\n150 000\n200 000",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "salary_expectation",
                "launch_scenario_key": "",
                "send_employee_card": False,
                "notify_on_send_text": "",
                "notify_on_send_recipient_ids": "",
                "notify_on_send_recipient_scope": "",
                "button_notifications": [],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(payload["response_type"], "buttons")
        self.assertEqual(payload["response_label"], "Выбор кнопками")
        self.assertEqual(payload["target_field"], "salary_expectation")
        self.assertEqual(payload["target_field_label"], "Ожидания по доходу")

        with SessionLocal() as db:
            step = db.get(FlowStepTemplate, step_id)
            self.assertIsNotNone(step)
            self.assertEqual(step.response_type, "buttons")
            self.assertEqual(step.target_field, "salary_expectation")

    def test_workspace_step_api_persists_date_response_for_first_workday(self) -> None:
        scenario_key = f"codex-date-step-{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-date-step-{self.unique_tag}",
                role_scope="all",
                scenario_kind="scenario",
                sort_order=0,
                trigger_mode="manual_only",
            )
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key="offer_date",
                step_title="Дата выхода",
                sort_order=10,
                default_text="Когда выходишь?",
                response_type="text",
                send_mode="immediate",
                day_offset_workdays=0,
                target_field=None,
            )
            db.add_all([scenario, step])
            db.commit()
            db.refresh(step)
            step_id = step.id

        response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Дата выхода",
                "text": "{name}, с какого числа ты планируешь присоединиться к команде Зефира?",
                "response_type": "date",
                "button_options": "",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "first_workday",
                "launch_scenario_key": "",
                "send_employee_card": False,
                "notify_on_send_text": "",
                "notify_on_send_recipient_ids": "",
                "notify_on_send_recipient_scope": "",
                "step_send_notifications": [],
                "button_notifications": [],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(payload["response_type"], "date")
        self.assertEqual(payload["target_field"], "first_workday")
        self.assertEqual(payload["target_field_label"], "Первый день выхода")

        with SessionLocal() as db:
            step = db.get(FlowStepTemplate, step_id)
            self.assertIsNotNone(step)
            self.assertEqual(step.response_type, "date")
            self.assertEqual(step.target_field, "first_workday")

    def test_workspace_step_api_preserves_branching_target_field(self) -> None:
        scenario_key = f"codex_branch_target_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-branch-target-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Ожидаемый доход через ветвление",
                sort_order=10,
                default_text="Выбери ожидаемый доход",
                custom_text=None,
                response_type="branching",
                button_options="200000\n300000",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add(step)
            db.commit()
            db.refresh(step)
            step_id = step.id

        response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Ожидаемый доход через ветвление",
                "text": "Выбери ожидаемый доход",
                "response_type": "branching",
                "button_options": "200000\n300000",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "salary_expectation",
                "launch_scenario_key": "",
                "send_employee_card": False,
                "notify_on_send_text": "",
                "notify_on_send_recipient_ids": "",
                "notify_on_send_recipient_scope": "",
                "button_notifications": [],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(payload["response_type"], "branching")
        self.assertEqual(payload["target_field"], "salary_expectation")
        self.assertEqual(payload["target_field_label"], "Ожидания по доходу")

        with SessionLocal() as db:
            step = db.get(FlowStepTemplate, step_id)
            self.assertIsNotNone(step)
            self.assertEqual(step.response_type, "branching")
            self.assertEqual(step.target_field, "salary_expectation")

    def test_workspace_step_api_normalizes_survey_question_flow(self) -> None:
        scenario_key = f"codex_survey_flow_{self.unique_tag}"
        with SessionLocal() as db:
            survey = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-survey-flow-{self.unique_tag}",
                sort_order=10,
                scenario_kind="survey",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(survey)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Старый вопрос",
                sort_order=10,
                default_text="Старый вопрос",
                custom_text="Старый вопрос",
                response_type="text",
                button_options=None,
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add(step)
            db.commit()
            db.refresh(step)
            step_id = step.id

        response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Какой у тебя график?",
                "text": "Какой у тебя график?",
                "response_type": "launch_scenario",
                "button_options": "5/2\n2/2",
                "send_mode": "specific_time",
                "send_time": "10:30",
                "target_field": "full_name",
                "launch_scenario_key": "another_flow",
                "send_employee_card": True,
                "notify_on_send_text": "Не должно сохраниться",
                "notify_on_send_recipient_ids": "manager",
                "notify_on_send_recipient_scope": "manager",
                "step_send_notifications": [
                    {
                        "rule_index": 0,
                        "message_text": "Тоже не должно сохраниться",
                        "recipient_ids": "mentor_adaptation",
                        "recipient_scope": "",
                    }
                ],
                "button_notifications": [
                    {
                        "option_index": 0,
                        "rules": [
                            {
                                "rule_index": 0,
                                "message_text": "Лишнее кнопочное уведомление",
                                "recipient_ids": "mentor_ipr",
                                "recipient_scope": "",
                            }
                        ],
                    }
                ],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(payload["title"], "Какой у тебя график?")
        self.assertEqual(payload["text"], "Какой у тебя график?")
        self.assertEqual(payload["response_type"], "text")
        self.assertEqual(payload["button_options"], ["5/2", "2/2"])
        self.assertEqual(payload["send_mode"], "immediate")
        self.assertEqual(payload["send_time"], "")
        self.assertEqual(payload["target_field"], "")
        self.assertEqual(payload["launch_scenario_key"], "")
        self.assertEqual(payload["step_send_notifications"], [])
        self.assertEqual(payload["button_notifications"][0]["rules"], [])

        with SessionLocal() as db:
            step = db.get(FlowStepTemplate, step_id)
            self.assertIsNotNone(step)
            self.assertEqual(step.step_title, "Какой у тебя график?")
            self.assertEqual(step.custom_text, "Какой у тебя график?")
            self.assertEqual(step.default_text, "Какой у тебя график?")
            self.assertEqual(step.response_type, "text")
            self.assertEqual(step.button_options, "5/2\n2/2")
            self.assertEqual(step.send_mode, "immediate")
            self.assertIsNone(step.send_time)
            self.assertIsNone(step.target_field)
            self.assertIsNone(step.launch_scenario_key)
            self.assertFalse(bool(step.send_employee_card))
            self.assertIsNone(step.notify_on_send_text)
            self.assertEqual(
                db.query(StepSendNotification).filter(StepSendNotification.step_id == step_id).count(),
                0,
            )
            self.assertEqual(
                db.query(StepButtonNotification).filter(StepButtonNotification.step_id == step_id).count(),
                0,
            )

    def test_workspace_step_api_updates_step_send_notifications(self) -> None:
        scenario_key = f"codex_step_notify_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-step-notify-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с уведомлением",
                sort_order=10,
                default_text="Привет",
                custom_text=None,
                response_type="none",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add(step)
            db.commit()
            db.refresh(step)
            step_id = step.id

        response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Шаг с уведомлением",
                "text": "Привет",
                "response_type": "none",
                "button_options": "",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "",
                "launch_scenario_key": "",
                "send_employee_card": False,
                "notify_on_send_text": "",
                "notify_on_send_recipient_ids": "",
                "notify_on_send_recipient_scope": "",
                "step_send_notifications": [
                    {
                        "rule_index": 0,
                        "message_text": "Шаг отправлен HR",
                        "recipient_ids": "manager",
                        "recipient_scope": "",
                    },
                    {
                        "rule_index": 1,
                        "message_text": "Дубликат уведомления",
                        "recipient_ids": "hr",
                        "recipient_scope": "",
                    },
                ],
                "button_notifications": [],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(payload["step_send_notifications"][0]["message_text"], "Шаг отправлен HR")
        self.assertEqual(payload["step_send_notifications"][0]["recipient_ids"], "manager")
        self.assertEqual(len(payload["step_send_notifications"]), 2)
        self.assertEqual(payload["step_send_notifications"][1]["message_text"], "Дубликат уведомления")

        with SessionLocal() as db:
            notifications = (
                db.query(StepSendNotification)
                .filter(StepSendNotification.step_id == step_id)
                .order_by(StepSendNotification.rule_index.asc())
                .all()
            )
            self.assertEqual(len(notifications), 2)
            self.assertEqual(notifications[0].rule_index, 0)
            self.assertEqual(notifications[0].message_text, "Шаг отправлен HR")
            self.assertEqual(notifications[0].recipient_ids, "manager")
            self.assertIsNone(notifications[0].recipient_scope)
            self.assertEqual(notifications[1].rule_index, 1)
            self.assertEqual(notifications[1].message_text, "Дубликат уведомления")
            self.assertEqual(notifications[1].recipient_ids, "hr")

        update_response = self.client.post(
            f"/api/flows/workspace/steps/{step_id}",
            json={
                "title": "Шаг с уведомлением",
                "text": "Привет",
                "response_type": "none",
                "button_options": "",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "",
                "launch_scenario_key": "",
                "send_employee_card": False,
                "notify_on_send_text": "Обновлённое уведомление",
                "notify_on_send_recipient_ids": "mentor_adaptation",
                "notify_on_send_recipient_scope": "",
                "step_send_notifications": [
                    {
                        "rule_index": 0,
                        "message_text": "Обновлённое уведомление",
                        "recipient_ids": "mentor_adaptation",
                        "recipient_scope": "",
                    },
                ],
                "button_notifications": [],
            },
        )

        self.assertEqual(update_response.status_code, 200)
        updated_payload = update_response.json()["payload"]["workspace"]["root_steps"][0]
        self.assertEqual(len(updated_payload["step_send_notifications"]), 1)
        self.assertEqual(updated_payload["step_send_notifications"][0]["message_text"], "Обновлённое уведомление")

        with SessionLocal() as db:
            notifications = db.query(StepSendNotification).filter(StepSendNotification.step_id == step_id).all()
            self.assertEqual(len(notifications), 1)
            self.assertEqual(notifications[0].message_text, "Обновлённое уведомление")
            self.assertEqual(notifications[0].recipient_ids, "mentor_adaptation")
            self.assertIsNone(notifications[0].recipient_scope)

    def test_workspace_branch_step_api_persists_return_to_root_step(self) -> None:
        scenario_key = f"codex_branch_return_{self.unique_tag}"
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-branch-return-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            root_step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Root branching",
                sort_order=10,
                default_text="Выбери вариант",
                custom_text=None,
                response_type="branching",
                button_options="Да\nНет",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            followup_step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_2",
                step_title="Common flow",
                sort_order=20,
                default_text="Общий поток",
                custom_text=None,
                response_type="none",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            branch_step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_branch_yes",
                parent_step_id=None,
                branch_option_index=0,
                step_title="Ветка Да",
                sort_order=1001,
                default_text="Локальная ветка",
                custom_text=None,
                response_type="none",
                send_mode="immediate",
                send_time=None,
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add_all([root_step, followup_step])
            db.flush()
            branch_step.parent_step_id = root_step.id
            db.add(branch_step)
            db.commit()
            db.refresh(branch_step)
            branch_step_id = branch_step.id
            followup_step_key = followup_step.step_key

        response = self.client.post(
            f"/api/flows/workspace/steps/{branch_step_id}",
            json={
                "title": "Ветка Да",
                "text": "Локальная ветка",
                "response_type": "none",
                "button_options": "",
                "send_mode": "immediate",
                "send_time": "",
                "target_field": "",
                "launch_scenario_key": "",
                "return_to_step_key": followup_step_key,
                "send_employee_card": False,
                "notify_on_send_text": "",
                "notify_on_send_recipient_ids": "",
                "notify_on_send_recipient_scope": "",
                "step_send_notifications": [],
                "button_notifications": [],
            },
        )

        self.assertEqual(response.status_code, 200)
        root_payload = response.json()["payload"]["workspace"]["root_steps"][0]
        branch_payload = root_payload["branch_items"][0]["step"]
        self.assertEqual(branch_payload["return_to_step_key"], followup_step_key)
        self.assertEqual(branch_payload["step_key"], f"{scenario_key}_branch_yes")

        with SessionLocal() as db:
            branch_step = db.get(FlowStepTemplate, branch_step_id)
            self.assertIsNotNone(branch_step)
            self.assertEqual(branch_step.return_to_step_key, followup_step_key)

    def test_workspace_payload_includes_read_only_graph_contract(self) -> None:
        scenario_key = f"codex_graph_payload_{self.unique_tag}"
        target_scenario_key = f"{scenario_key}_target"
        with SessionLocal() as db:
            target_scenario = ScenarioTemplate(
                scenario_key=target_scenario_key,
                title="Target scenario",
                sort_order=5,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-graph-payload-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add_all([target_scenario, scenario])
            db.flush()
            branching_root = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Root branching",
                sort_order=10,
                default_text="Выбери вариант",
                response_type="branching",
                button_options="Да\nНет",
                send_mode="immediate",
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            followup_root = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_2",
                step_title="Launch step",
                sort_order=20,
                default_text="Переходим дальше",
                response_type="launch_scenario",
                launch_scenario_key=target_scenario_key,
                send_mode="immediate",
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
            )
            db.add_all([branching_root, followup_root])
            db.flush()
            branch_step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_branch_yes",
                parent_step_id=branching_root.id,
                branch_option_index=0,
                step_title="Ветка Да",
                sort_order=1001,
                default_text="Локальная ветка",
                response_type="none",
                send_mode="immediate",
                day_offset_workdays=0,
                target_field=None,
                send_employee_card=False,
                return_to_step_key=followup_root.step_key,
            )
            db.add(branch_step)
            db.commit()
            scenario_id = scenario.id

        response = self.client.get(f"/api/flows/workspace?scenario_id={scenario_id}")

        self.assertEqual(response.status_code, 200)
        workspace = response.json()["workspace"]
        graph = workspace["graph"]
        self.assertGreaterEqual(graph["meta"]["node_count"], 5)
        self.assertTrue(graph["meta"]["has_branching"])
        self.assertTrue(graph["meta"]["has_return_edges"])
        self.assertTrue(graph["meta"]["has_launch_edges"])
        self.assertTrue(graph["meta"]["has_placeholders"])

        node_kinds = {node["kind"] for node in graph["nodes"]}
        self.assertIn("root_step", node_kinds)
        self.assertIn("branch_step", node_kinds)
        self.assertIn("branch_slot", node_kinds)
        self.assertIn("launch_target", node_kinds)

        branch_edges = [edge for edge in graph["edges"] if edge["kind"] == "branch_option"]
        self.assertEqual({edge["label"] for edge in branch_edges}, {"Да", "Нет"})
        self.assertTrue(any(edge["kind"] == "return_to_root" for edge in graph["edges"]))
        self.assertTrue(any(edge["kind"] == "launch_scenario" for edge in graph["edges"]))

    def test_button_notification_to_mentor_ipr_resolves_correctly(self) -> None:
        scenario_key = f"codex_button_runtime_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "200001", db=db)
            mentor_ipr = Employee(
                full_name=f"Mentor IPR Runtime {self.unique_tag}",
                telegram_user_id="780001",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
                is_mentor=True,
            )
            db.add(mentor_ipr)
            db.flush()
            set_primary_chat_id(mentor_ipr, "210003", db=db)
            employee.mentor_ipr_employee_id = mentor_ipr.id
            employee.employee_stage = "candidate"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-button-runtime-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с кнопками",
                sort_order=10,
                default_text="Нажми кнопку",
                response_type="buttons",
                button_options="Да\nНет",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                ScenarioProgress(
                    employee_id=employee.id,
                    scenario_key=scenario_key,
                    current_step_key=step.step_key,
                    waiting_for_response=True,
                    is_completed=False,
                    started_at=now,
                    updated_at=now,
                )
            )
            db.add(
                StepButtonNotification(
                    flow_key=scenario_key,
                    step_id=step.id,
                    option_index=0,
                    rule_index=0,
                    message_text="Правило mentor_ipr",
                    recipient_ids="mentor_ipr",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            handled = asyncio.run(handle_button_response(messenger, db, employee, scenario_key, step.step_key, 0))

            self.assertTrue(handled)
            self.assertIn(("210003", "Правило mentor_ipr"), messenger.sent_texts)

    def test_send_step_notification_to_manager_resolves_correctly(self) -> None:
        scenario_key = f"codex_step_manager_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "200001", db=db)
            manager = Employee(
                full_name=f"Manager Runtime {self.unique_tag}",
                telegram_user_id="780002",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
                is_manager=True,
            )
            db.add(manager)
            db.flush()
            set_primary_chat_id(manager, "210001", db=db)
            employee.manager_employee_id = manager.id
            employee.employee_stage = "candidate"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-step-manager-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с уведомлением",
                sort_order=10,
                default_text="Текст шага",
                response_type="none",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                StepSendNotification(
                    flow_key=scenario_key,
                    step_id=step.id,
                    rule_index=0,
                    message_text="Шаг правило manager",
                    recipient_ids="manager",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(send_step(messenger, db, employee, scenario, step))

            self.assertEqual(
                messenger.sent_texts[:2],
                [("200001", "Текст шага"), ("210001", "Шаг правило manager")],
            )

    def test_send_step_notification_to_mentor_adaptation_resolves_correctly(self) -> None:
        scenario_key = f"codex_step_mentor_adapt_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "200001", db=db)
            mentor = Employee(
                full_name=f"Mentor Adapt Runtime {self.unique_tag}",
                telegram_user_id="780003",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
                is_mentor=True,
            )
            db.add(mentor)
            db.flush()
            set_primary_chat_id(mentor, "210002", db=db)
            employee.mentor_adaptation_employee_id = mentor.id
            employee.employee_stage = "candidate"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-step-mentor-adapt-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с уведомлением",
                sort_order=10,
                default_text="Текст шага",
                response_type="none",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                StepSendNotification(
                    flow_key=scenario_key,
                    step_id=step.id,
                    rule_index=0,
                    message_text="Шаг правило mentor_adaptation",
                    recipient_ids="mentor_adaptation",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(send_step(messenger, db, employee, scenario, step))

            self.assertIn(("210002", "Шаг правило mentor_adaptation"), messenger.sent_texts)

    def test_send_step_notification_to_hr_resolves_via_hr_settings(self) -> None:
        scenario_key = f"codex_step_hr_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "200001", db=db)
            hr_settings = db.get(HrSettings, 1)
            self.assertIsNotNone(hr_settings)
            hr_settings.telegram_user_id = "210004"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-step-hr-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с уведомлением",
                sort_order=10,
                default_text="Текст шага",
                response_type="none",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                StepSendNotification(
                    flow_key=scenario_key,
                    step_id=step.id,
                    rule_index=0,
                    message_text="Шаг правило hr",
                    recipient_ids="hr",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(send_step(messenger, db, employee, scenario, step))

            self.assertIn(("210004", "Шаг правило hr"), messenger.sent_texts)

    def test_send_step_notification_skips_missing_manager_without_failing_scenario(self) -> None:
        scenario_key = f"codex_step_missing_manager_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "200001", db=db)
            employee.manager_employee_id = None
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-step-missing-manager-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с уведомлением",
                sort_order=10,
                default_text="Текст шага",
                response_type="none",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                StepSendNotification(
                    flow_key=scenario_key,
                    step_id=step.id,
                    rule_index=0,
                    message_text="Не должно упасть",
                    recipient_ids="manager",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(send_step(messenger, db, employee, scenario, step))

            self.assertEqual(messenger.sent_texts, [("200001", "Текст шага")])

    def test_send_step_notification_keeps_legacy_employee_token_runtime_support(self) -> None:
        scenario_key = f"codex_step_legacy_recipient_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            set_primary_chat_id(employee, "200001", db=db)
            observer = Employee(
                full_name=f"Legacy Observer {self.unique_tag}",
                telegram_user_id="780004",
                first_workday=None,
                created_at=now,
                is_flow_scheduled=False,
                employee_stage="staff",
            )
            db.add(observer)
            db.flush()
            set_primary_chat_id(observer, "210005", db=db)
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-step-legacy-recipient-{self.unique_tag}",
                sort_order=10,
                scenario_kind="scenario",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(scenario)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Шаг с уведомлением",
                sort_order=10,
                default_text="Текст шага",
                response_type="none",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                StepSendNotification(
                    flow_key=scenario_key,
                    step_id=step.id,
                    rule_index=0,
                    message_text="Legacy recipient works",
                    recipient_ids=f"employee:{observer.id}",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(send_step(messenger, db, employee, scenario, step))

            self.assertIn(("210005", "Legacy recipient works"), messenger.sent_texts)

    def test_survey_option_buttons_work_without_branching_response_type(self) -> None:
        scenario_key = f"codex_survey_buttons_{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            survey = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-survey-buttons-{self.unique_tag}",
                sort_order=10,
                scenario_kind="survey",
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
            )
            db.add(survey)
            db.flush()
            step = FlowStepTemplate(
                flow_key=scenario_key,
                step_key=f"{scenario_key}_step_1",
                step_title="Какой график?",
                sort_order=10,
                default_text="Какой график?",
                custom_text="Какой график?",
                response_type="text",
                button_options="5/2\n2/2",
                send_mode="immediate",
            )
            db.add(step)
            db.flush()
            db.add(
                ScenarioProgress(
                    employee_id=employee.id,
                    scenario_key=scenario_key,
                    current_step_key=step.step_key,
                    waiting_for_response=True,
                    is_completed=False,
                    started_at=now,
                    updated_at=now,
                )
            )
            db.commit()
            messenger = DummyMessenger()

            handled = asyncio.run(handle_button_response(messenger, db, employee, scenario_key, step.step_key, 1))

            self.assertTrue(handled)
            stored_answer = (
                db.query(SurveyAnswer)
                .filter(
                    SurveyAnswer.employee_id == employee.id,
                    SurveyAnswer.scenario_key == scenario_key,
                    SurveyAnswer.step_key == step.step_key,
                )
                .first()
            )
            self.assertIsNotNone(stored_answer)
            self.assertEqual(stored_answer.answer_value, "2/2")

    def test_employee_file_delete_api_removes_uploaded_file(self) -> None:
        with SessionLocal() as db:
            db_file = EmployeeFile(
                employee_id=self.employee_id,
                direction="outbound",
                category="hr_file",
                telegram_file_id=None,
                telegram_file_unique_id=None,
                original_filename=f"codex-file-{self.unique_tag}.txt",
                stored_path=str((Path("storage") / "employee_files" / f"codex-file-{self.unique_tag}.txt").resolve()),
                mime_type="text/plain",
                file_size=5,
                created_at=datetime.now(UTC).replace(tzinfo=None),
            )
            file_path = Path(db_file.stored_path)
            file_path.parent.mkdir(parents=True, exist_ok=True)
            file_path.write_text("hello", encoding="utf-8")
            db.add(db_file)
            db.commit()
            db.refresh(db_file)
            file_id = db_file.id

        response = self.client.delete(f"/api/employees/{self.employee_id}/files/{file_id}")

        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            self.assertIsNone(db.get(EmployeeFile, file_id))
        self.assertFalse(file_path.exists())

    def test_update_employee_api_returns_conflict_for_duplicate_chat_id(self) -> None:
        with SessionLocal() as db:
            other_employee = Employee(
                full_name="Existing Chat Owner",
                telegram_user_id="777888999",
                telegram_username="existing_owner",
                first_workday=datetime.now(UTC).date(),
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                candidate_status="new",
                employee_stage="candidate",
                candidate_work_stage="testing",
            )
            db.add(other_employee)
            db.commit()
            db.refresh(other_employee)
            conflict_employee_id = other_employee.id

        try:
            response = self.client.post(
                f"/api/employees/{self.employee_id}",
                json={
                    "full_name": "Updated API Smoke Employee",
                    "chat_id": "777888999",
                    "chat_handle": "hr_team",
                    "first_workday": "",
                    "desired_position": "",
                    "birth_date": "",
                    "work_email": "",
                    "work_hours": "",
                    "manager_chat_id": "",
                    "mentor_adaptation_chat_id": "",
                    "mentor_ipr_chat_id": "",
                    "employee_stage": "candidate",
                    "candidate_work_stage": "testing",
                    "salary_expectation": "",
                    "personal_data_consent": False,
                    "employee_data_consent": False,
                    "test_task_due_at": "",
                    "notes": "",
                },
            )

            self.assertEqual(response.status_code, 409)
            self.assertIn("777888999", response.json()["detail"])
        finally:
            with SessionLocal() as db:
                db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == conflict_employee_id).delete()
                other_employee = db.get(Employee, conflict_employee_id)
                if other_employee is not None:
                    db.delete(other_employee)
                    db.commit()

    def test_launch_flow_api_rejects_blocked_employee(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.is_bot_blocked = True
            employee.employee_stage = "staff"
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}/launch",
            json={"flow_key": "recruitment_hiring"},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("доступ к боту заблокирован", response.json()["detail"].lower())

    def test_reset_employee_bot_linkage_api_clears_identity_and_pending_runtime_state(self) -> None:
        scenario_key = f"codex-reset-linkage-{self.unique_tag}"
        now = datetime.now(UTC).replace(tzinfo=None)
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.telegram_user_id = "777123456"
            employee.telegram_username = "reset_me"
            employee.current_menu_set_id = 123
            employee.is_flow_scheduled = True
            db.add(
                EmployeeMessengerAccount(
                    employee_id=employee.id,
                    channel="telegram",
                    external_user_id="777123456",
                    external_username="reset_me",
                    is_primary=True,
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                )
            )
            db.add(
                ScenarioProgress(
                    employee_id=employee.id,
                    scenario_key=scenario_key,
                    current_step_key="step_1",
                    waiting_for_response=True,
                    is_completed=False,
                    started_at=now,
                    updated_at=now,
                    response_undo_history='[{"step_key":"step_1"}]',
                )
            )
            db.add(
                FlowLaunchRequest(
                    employee_id=employee.id,
                    flow_key=scenario_key,
                    requested_at=now,
                    processed_at=None,
                    launch_type="scheduled",
                    skip_step_key=None,
                )
            )
            db.add(
                FlowLaunchRequest(
                    employee_id=employee.id,
                    flow_key=scenario_key,
                    requested_at=now,
                    processed_at=now,
                    launch_type="manual",
                    skip_step_key=None,
                )
            )
            db.commit()

        response = self.client.post(
            f"/api/employees/{self.employee_id}/bot-link/reset",
            headers={"Accept": "application/json"},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["employee"]["chat_id"], "")
        self.assertEqual(payload["employee"]["chat_handle"], "")
        self.assertEqual(payload["scheduled_launches"], [])

        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            self.assertIsNone(employee.telegram_user_id)
            self.assertIsNone(employee.telegram_username)
            self.assertIsNone(employee.current_menu_set_id)
            self.assertFalse(employee.is_flow_scheduled)
            self.assertEqual(
                db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == self.employee_id).count(),
                0,
            )
            self.assertEqual(
                db.query(ScenarioProgress).filter(ScenarioProgress.employee_id == self.employee_id).count(),
                0,
            )
            self.assertEqual(
                db.query(FlowLaunchRequest)
                .filter(
                    FlowLaunchRequest.employee_id == self.employee_id,
                    FlowLaunchRequest.processed_at.is_(None),
                )
                .count(),
                0,
            )
            self.assertEqual(
                db.query(FlowLaunchRequest)
                .filter(
                    FlowLaunchRequest.employee_id == self.employee_id,
                    FlowLaunchRequest.processed_at.is_not(None),
                )
                .count(),
                1,
            )

    def test_bulk_actions_preview_api_uses_candidate_stage_split(self) -> None:
        response = self.client.post(
            "/api/bulk-actions/preview",
            json={
                "target_role_scope": "",
                "target_employee_id": None,
                "target_employee_stages": [],
                "target_candidate_stages": ["testing"],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertGreaterEqual(payload["recipient_count"], 1)
        self.assertIn("Кандидаты", payload["recipient_scope"])

    def test_bulk_actions_schedule_message_api_persists_react_workspace_payload(self) -> None:
        message_text = f"codex-message-{self.unique_tag}"
        response = self.client.post(
            "/api/bulk-actions/messages/schedule",
            json={
                "message_text": message_text,
                "requested_at": "2026-06-01T10:30",
                "target_role_scope": "",
                "target_employee_id": None,
                "target_employee_stages": [],
                "target_candidate_stages": ["testing"],
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()["payload"]
        self.assertIn("scheduled_message_actions", payload)
        with SessionLocal() as db:
            action = db.query(MassMessageAction).filter(MassMessageAction.message_text == message_text).one_or_none()
            self.assertIsNotNone(action)
            self.assertEqual(action.launch_type, "scheduled")
            self.assertEqual(action.target_candidate_stages, "testing")
            self.assertEqual(action.target_employee_stages, None)
            self.assertGreaterEqual(action.recipient_count, 1)

    def test_bulk_actions_delete_scheduled_survey_uses_survey_route(self) -> None:
        scenario_key = f"codex_scheduled_survey_{self.unique_tag}"
        requested_at = datetime(2026, 6, 1, 10, 30)
        with SessionLocal() as db:
            survey = ScenarioTemplate(
                scenario_key=scenario_key,
                scenario_kind="survey",
                title=f"codex-scheduled-survey-{self.unique_tag}",
                sort_order=10,
                role_scope="all",
                employee_scope="all",
                trigger_mode="manual_only",
                description="scheduled survey delete smoke",
            )
            db.add(survey)
            db.flush()
            action = MassScenarioAction(
                flow_key=scenario_key,
                scenario_kind="survey",
                requested_at=requested_at,
                processed_at=None,
                launch_type="scheduled",
                target_all=True,
                target_statuses=None,
                target_employee_stages=None,
                target_candidate_stages=None,
                target_role_scope=None,
                target_employee_id=None,
                recipient_count=1,
                created_at=requested_at,
            )
            db.add(action)
            db.commit()
            db.refresh(action)
            action_id = action.id

        response = self.client.delete(f"/api/bulk-actions/surveys/{action_id}")

        self.assertEqual(response.status_code, 200)
        with SessionLocal() as db:
            self.assertIsNone(db.get(MassScenarioAction, action_id))

    def test_settings_hr_update_api_persists_workspace_fields(self) -> None:
        candidate_menu_title = f"codex-candidate-root-{self.unique_tag}"
        employee_menu_title = f"codex-employee-root-{self.unique_tag}"
        with SessionLocal() as db:
            candidate_menu = BotMenuSet(
                title=candidate_menu_title,
                description="candidate root",
                sort_order=10,
                employee_scope="candidates",
            )
            employee_menu = BotMenuSet(
                title=employee_menu_title,
                description="employee root",
                sort_order=20,
                employee_scope="employees",
            )
            db.add_all([candidate_menu, employee_menu])
            db.commit()
            db.refresh(candidate_menu)
            db.refresh(employee_menu)
            candidate_menu_id = candidate_menu.id
            employee_menu_id = employee_menu.id

        response = self.client.post(
            "/api/settings/hr",
            json={
                "hr_name": f"codex-hr-{self.unique_tag}",
                "telegram_user_id": f"tg-{self.unique_tag}",
                "notification_recipient_ids": f"tg-a-{self.unique_tag},tg-b-{self.unique_tag}",
                "default_menu_set_id": None,
                "default_employee_menu_set_id": employee_menu_id,
                "default_candidate_menu_set_id": candidate_menu_id,
                "notify_scenario_completed": False,
                "notify_test_task_received": True,
                "notify_user_actions": False,
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["hr_settings"]["hr_name"], f"codex-hr-{self.unique_tag}")
        with SessionLocal() as db:
            hr_settings = db.query(HrSettings).first()
            self.assertIsNotNone(hr_settings)
            self.assertEqual(hr_settings.hr_name, f"codex-hr-{self.unique_tag}")
            self.assertEqual(hr_settings.telegram_user_id, f"tg-{self.unique_tag}")
            self.assertEqual(hr_settings.default_employee_menu_set_id, employee_menu_id)
            self.assertEqual(hr_settings.default_candidate_menu_set_id, candidate_menu_id)
            self.assertFalse(hr_settings.notify_scenario_completed)
            self.assertFalse(hr_settings.notify_user_actions)

    def test_bot_root_menu_uses_explicit_audience_defaults(self) -> None:
        with SessionLocal() as db:
            candidate = db.get(Employee, self.employee_id)
            self.assertIsNotNone(candidate)
            candidate.employee_stage = "candidate"
            candidate.candidate_work_stage = "testing"

            employee = Employee(
                full_name=f"Staff Root {self.unique_tag}",
                telegram_user_id=None,
                first_workday=None,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                is_flow_scheduled=False,
                employee_stage="staff",
                candidate_work_stage=None,
            )
            candidate_root = BotMenuSet(
                title=f"codex-candidate-root-{self.unique_tag}",
                description="candidate root",
                sort_order=10,
                employee_scope="candidates",
            )
            employee_root = BotMenuSet(
                title=f"codex-employee-root-{self.unique_tag}",
                description="employee root",
                sort_order=20,
                employee_scope="employees",
            )
            db.add_all([employee, candidate_root, employee_root])
            db.commit()
            db.refresh(employee)
            db.refresh(candidate_root)
            db.refresh(employee_root)

            hr_settings = db.query(HrSettings).first()
            self.assertIsNotNone(hr_settings)
            hr_settings.default_menu_set_id = None
            hr_settings.default_candidate_menu_set_id = candidate_root.id
            hr_settings.default_employee_menu_set_id = employee_root.id
            db.commit()

            candidate_resolved = current_menu_set(db, candidate)
            employee_resolved = current_menu_set(db, employee)

            self.assertIsNotNone(candidate_resolved)
            self.assertIsNotNone(employee_resolved)
            self.assertEqual(candidate_resolved.id, candidate_root.id)
            self.assertEqual(employee_resolved.id, employee_root.id)

    def test_settings_menu_set_api_supports_react_workspace_create_and_delete(self) -> None:
        title = f"codex-menu-{self.unique_tag}"
        create_response = self.client.post(
            "/api/settings/menu-sets",
            json={"title": title, "description": "workspace smoke"},
        )

        self.assertEqual(create_response.status_code, 200)
        created_workspace = create_response.json()
        created_menu_set = next((item for item in created_workspace["menu_sets"] if item["title"] == title), None)
        self.assertIsNotNone(created_menu_set)

        delete_response = self.client.delete(f"/api/settings/menu-sets/{created_menu_set['id']}")
        self.assertEqual(delete_response.status_code, 200)
        deleted_workspace = delete_response.json()
        self.assertFalse(any(item["id"] == created_menu_set["id"] for item in deleted_workspace["menu_sets"]))

    def test_settings_menu_set_api_persists_audience_rules(self) -> None:
        title = f"codex-audience-{self.unique_tag}"
        create_response = self.client.post(
            "/api/settings/menu-sets",
            json={
                "title": title,
                "description": "candidate menu",
                "role_scope": "designer",
                "employee_scope": "candidates",
                "target_employee_ids": [self.employee_id],
            },
        )

        self.assertEqual(create_response.status_code, 200)
        created_menu_set = next((item for item in create_response.json()["menu_sets"] if item["title"] == title), None)
        self.assertIsNotNone(created_menu_set)
        self.assertEqual(created_menu_set["role_scope"], "designer")
        self.assertEqual(created_menu_set["employee_scope"], "candidates")
        self.assertEqual(created_menu_set["target_employee_ids"], [self.employee_id])

        with SessionLocal() as db:
            menu_set = db.get(BotMenuSet, created_menu_set["id"])
            self.assertIsNotNone(menu_set)
            self.assertEqual(menu_set.role_scope, "designer")
            self.assertEqual(menu_set.employee_scope, "candidates")
            self.assertEqual(menu_set.target_employee_ids, str(self.employee_id))
            self.assertIsNone(menu_set.target_employee_stages)
            self.assertIsNone(menu_set.target_candidate_stages)

    def test_settings_menu_set_api_rejects_duplicate_target_employee_assignment(self) -> None:
        first_title = f"codex-audience-first-{self.unique_tag}"
        second_title = f"codex-audience-second-{self.unique_tag}"

        first_response = self.client.post(
            "/api/settings/menu-sets",
            json={
                "title": first_title,
                "employee_scope": "candidates",
                "target_employee_ids": [self.employee_id],
            },
        )
        self.assertEqual(first_response.status_code, 200)

        second_response = self.client.post(
            "/api/settings/menu-sets",
            json={
                "title": second_title,
                "employee_scope": "candidates",
                "target_employee_ids": [self.employee_id],
            },
        )
        self.assertEqual(second_response.status_code, 409)
        self.assertIn("уже привязаны", second_response.json()["detail"])

    def test_bot_current_menu_set_prefers_matching_candidate_audience_set(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "testing"
            employee.desired_position = "Дизайнер"
            employee.current_menu_set_id = None
            hr_settings = db.query(HrSettings).first()
            if hr_settings is not None:
                hr_settings.default_menu_set_id = None
            broad_menu = BotMenuSet(
                title=f"codex-broad-{self.unique_tag}",
                description="all",
                sort_order=10,
            )
            targeted_menu = BotMenuSet(
                title=f"codex-targeted-{self.unique_tag}",
                description="candidate designer",
                sort_order=20,
                role_scope="designer",
                employee_scope="candidates",
            )
            db.add_all([broad_menu, targeted_menu])
            db.commit()
            db.refresh(targeted_menu)

            resolved = current_menu_set(db, employee)

            self.assertIsNotNone(resolved)
            self.assertEqual(resolved.id, targeted_menu.id)
            self.assertEqual(employee.current_menu_set_id, targeted_menu.id)

    def test_bot_menu_open_set_blocks_incompatible_target_set(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "testing"
            employee.telegram_user_id = "123456789"
            root_menu = BotMenuSet(
                title=f"codex-root-{self.unique_tag}",
                description="root",
                sort_order=10,
                employee_scope="candidates",
            )
            restricted_menu = BotMenuSet(
                title=f"codex-staff-{self.unique_tag}",
                description="staff only",
                sort_order=20,
                employee_scope="employees",
            )
            db.add_all([root_menu, restricted_menu])
            db.commit()
            db.refresh(root_menu)
            db.refresh(restricted_menu)
            button = BotMenuButton(
                menu_set_id=root_menu.id,
                label=f"codex-open-{self.unique_tag}",
                sort_order=10,
                action_type="open_set",
                target_menu_set_id=restricted_menu.id,
            )
            db.add(button)
            employee.current_menu_set_id = root_menu.id
            db.commit()
            messenger = DummyMessenger()

            handled = asyncio.run(handle_menu_button(messenger, db, employee, button.label))

            self.assertTrue(handled)
            self.assertEqual(employee.current_menu_set_id, root_menu.id)
            self.assertTrue(messenger.sent_menus)
            self.assertIn("недоступен", messenger.sent_menus[-1][1].lower())

    def test_bot_start_resets_user_to_root_menu(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "testing"
            chat_id = str(930000000000 + (uuid4().int % 100000000000))
            set_primary_chat_id(employee, chat_id, db=db)
            root_menu = BotMenuSet(
                title=f"codex-root-{self.unique_tag}",
                description="root",
                sort_order=10,
                employee_scope="candidates",
            )
            child_menu = BotMenuSet(
                title=f"codex-child-{self.unique_tag}",
                description="child",
                sort_order=20,
                employee_scope="candidates",
            )
            db.add_all([root_menu, child_menu])
            db.commit()
            db.refresh(root_menu)
            db.refresh(child_menu)
            hr_settings = db.query(HrSettings).first()
            self.assertIsNotNone(hr_settings)
            hr_settings.default_menu_set_id = root_menu.id
            employee.current_menu_set_id = child_menu.id
            employee.current_menu_path = f"{root_menu.id},{child_menu.id}"
            child_button = BotMenuButton(
                menu_set_id=child_menu.id,
                label=f"codex-child-btn-{self.unique_tag}",
                sort_order=10,
                action_type="inactive",
            )
            root_button = BotMenuButton(
                menu_set_id=root_menu.id,
                label=f"codex-root-btn-{self.unique_tag}",
                sort_order=10,
                action_type="inactive",
            )
            db.add_all([child_button, root_button])
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, employee.telegram_username))

            db.refresh(employee)
            self.assertEqual(employee.current_menu_set_id, root_menu.id)
            self.assertEqual(employee.current_menu_path, str(root_menu.id))
            self.assertEqual(messenger.sent_texts, [])
            self.assertTrue(messenger.sent_menus)
            self.assertEqual(messenger.sent_menus[-1][2], [root_button.label])

    def test_bot_start_without_menu_does_not_send_empty_menu_warning(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            chat_id = str(950000000000 + (uuid4().int % 100000000000))
            set_primary_chat_id(employee, chat_id, db=db)
            employee.current_menu_set_id = None
            employee.current_menu_path = None
            hr_settings = db.query(HrSettings).first()
            self.assertIsNotNone(hr_settings)
            hr_settings.default_menu_set_id = None
            hr_settings.default_employee_menu_set_id = None
            hr_settings.default_candidate_menu_set_id = None
            db.query(BotMenuButton).delete(synchronize_session=False)
            db.query(BotMenuSet).delete(synchronize_session=False)
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, employee.telegram_username))

            self.assertEqual(messenger.sent_texts, [])
            self.assertEqual(messenger.sent_menus, [])

    def test_bot_start_launches_registration_scenario_on_new_telegram_link(self) -> None:
        scenario_key = f"codex-registration-{self.unique_tag}"
        username = f"codex_user_{self.unique_tag}"
        chat_id = str(960000000000 + (uuid4().int % 100000000000))
        registration_text = f"codex registration step {self.unique_tag}"
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.telegram_username = username
            employee.telegram_user_id = None
            employee.employee_stage = "candidate"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-registration-{self.unique_tag}",
                role_scope="all",
                employee_scope="candidates",
                scenario_kind="scenario",
                sort_order=-100,
                trigger_mode="bot_registration",
            )
            db.add(scenario)
            db.add(
                FlowStepTemplate(
                    flow_key=scenario_key,
                    step_key="registration_start",
                    step_title="Registration start",
                    sort_order=10,
                    default_text=registration_text,
                    response_type="none",
                    send_mode="immediate",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, username))

            db.refresh(employee)
            self.assertEqual(get_primary_chat_id(employee, db=db), chat_id)
            self.assertEqual(messenger.sent_texts, [(chat_id, registration_text)])
            self.assertEqual(messenger.sent_menus, [])

            messenger.sent_texts.clear()
            asyncio.run(handle_start_command(messenger, db, chat_id, username))

            self.assertNotIn((chat_id, registration_text), messenger.sent_texts)

    def test_bot_start_creates_candidate_for_unknown_username_and_runs_registration(self) -> None:
        scenario_key = f"codex-registration-created-{self.unique_tag}"
        username = f"fresh_candidate_{self.unique_tag}"
        chat_id = str(961000000000 + (uuid4().int % 100000000000))
        registration_text = f"created candidate registration {self.unique_tag}"
        created_employee_id: int | None = None
        with SessionLocal() as db:
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-registration-created-{self.unique_tag}",
                role_scope="all",
                employee_scope="candidates",
                scenario_kind="scenario",
                sort_order=-100,
                trigger_mode="bot_registration",
            )
            db.add(scenario)
            db.add(
                FlowStepTemplate(
                    flow_key=scenario_key,
                    step_key="registration_start",
                    step_title="Registration start",
                    sort_order=10,
                    default_text=registration_text,
                    response_type="none",
                    send_mode="immediate",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, username))

            created_employee = db.query(Employee).filter(Employee.telegram_user_id == chat_id).order_by(Employee.id.desc()).first()
            self.assertIsNotNone(created_employee)
            created_employee_id = created_employee.id
            self.assertEqual(created_employee.employee_stage, "candidate")
            self.assertEqual(created_employee.candidate_work_stage, self._initial_candidate_stage_key())
            self.assertEqual(created_employee.telegram_username, username)
            self.assertEqual(messenger.sent_texts, [(chat_id, registration_text)])

            before_repeat_count = db.query(Employee).filter(Employee.telegram_user_id == chat_id).count()
            messenger.sent_texts.clear()
            asyncio.run(handle_start_command(messenger, db, chat_id, username))
            after_repeat_count = db.query(Employee).filter(Employee.telegram_user_id == chat_id).count()

            self.assertEqual(before_repeat_count, 1)
            self.assertEqual(after_repeat_count, 1)
            self.assertNotIn((chat_id, registration_text), messenger.sent_texts)

            if created_employee_id is not None:
                db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == created_employee_id).delete(
                    synchronize_session=False
                )
                created_employee = db.get(Employee, created_employee_id)
                if created_employee is not None:
                    db.delete(created_employee)
                    db.commit()

    def test_bot_start_creates_candidate_without_username(self) -> None:
        chat_id = str(962000000000 + (uuid4().int % 100000000000))
        created_employee_id: int | None = None
        with SessionLocal() as db:
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, None))

            created_employee = db.query(Employee).filter(Employee.telegram_user_id == chat_id).order_by(Employee.id.desc()).first()
            self.assertIsNotNone(created_employee)
            created_employee_id = created_employee.id
            self.assertEqual(created_employee.employee_stage, "candidate")
            self.assertEqual(created_employee.candidate_work_stage, self._initial_candidate_stage_key())
            self.assertIsNone(created_employee.telegram_username)
            self.assertNotIn((chat_id, "Привет! Я HR-бот."), messenger.sent_texts)
            self.assertTrue(messenger.sent_texts)

            if created_employee_id is not None:
                db.query(EmployeeMessengerAccount).filter(EmployeeMessengerAccount.employee_id == created_employee_id).delete(
                    synchronize_session=False
                )
                created_employee = db.get(Employee, created_employee_id)
                if created_employee is not None:
                    db.delete(created_employee)
                    db.commit()

    def test_bot_start_links_staff_by_username_without_registration(self) -> None:
        scenario_key = f"codex-registration-staff-{self.unique_tag}"
        username = f"staff_link_{self.unique_tag}"
        chat_id = str(963000000000 + (uuid4().int % 100000000000))
        registration_text = f"staff should not get this {self.unique_tag}"
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.telegram_username = f"@{username}"
            employee.telegram_user_id = None
            employee.employee_stage = "staff"
            scenario = ScenarioTemplate(
                scenario_key=scenario_key,
                title=f"codex-registration-staff-{self.unique_tag}",
                role_scope="all",
                employee_scope="candidates",
                scenario_kind="scenario",
                sort_order=-100,
                trigger_mode="bot_registration",
            )
            db.add(scenario)
            db.add(
                FlowStepTemplate(
                    flow_key=scenario_key,
                    step_key="registration_start",
                    step_title="Registration start",
                    sort_order=10,
                    default_text=registration_text,
                    response_type="none",
                    send_mode="immediate",
                )
            )
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, username.upper()))

            db.refresh(employee)
            self.assertEqual(get_primary_chat_id(employee, db=db), chat_id)
            self.assertEqual(messenger.sent_texts, [])
            self.assertNotIn((chat_id, registration_text), messenger.sent_texts)

    def test_bot_start_keeps_blocked_matched_user_blocked(self) -> None:
        username = f"blocked_link_{self.unique_tag}"
        chat_id = str(964000000000 + (uuid4().int % 100000000000))
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.telegram_username = f"@{username}"
            employee.telegram_user_id = None
            employee.is_bot_blocked = True
            db.commit()
            messenger = DummyMessenger()

            asyncio.run(handle_start_command(messenger, db, chat_id, username.upper()))

            db.refresh(employee)
            self.assertIsNone(get_primary_chat_id(employee, db=db))
            self.assertEqual(messenger.sent_texts, [(chat_id, BLOCKED_USER_TEXT)])

    def test_bot_menu_back_returns_to_previous_menu(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "testing"
            chat_id = str(940000000000 + (uuid4().int % 100000000000))
            set_primary_chat_id(employee, chat_id, db=db)
            root_menu = BotMenuSet(
                title=f"codex-root-{self.unique_tag}",
                description="root",
                sort_order=10,
                employee_scope="candidates",
            )
            child_menu = BotMenuSet(
                title=f"codex-child-{self.unique_tag}",
                description="child",
                sort_order=20,
                employee_scope="candidates",
            )
            db.add_all([root_menu, child_menu])
            db.commit()
            db.refresh(root_menu)
            db.refresh(child_menu)
            root_button = BotMenuButton(
                menu_set_id=root_menu.id,
                label=f"codex-root-btn-{self.unique_tag}",
                sort_order=10,
                action_type="open_set",
                target_menu_set_id=child_menu.id,
            )
            child_button = BotMenuButton(
                menu_set_id=child_menu.id,
                label=f"codex-child-btn-{self.unique_tag}",
                sort_order=10,
                action_type="inactive",
            )
            db.add_all([root_button, child_button])
            employee.current_menu_set_id = child_menu.id
            employee.current_menu_path = f"{root_menu.id},{child_menu.id}"
            db.commit()
            messenger = DummyMessenger()

            result = asyncio.run(handle_text_event(messenger, db, chat_id, employee.telegram_username, MENU_BACK_BUTTON_TEXT))

            db.refresh(employee)
            self.assertEqual(result, "handled")
            self.assertEqual(employee.current_menu_set_id, root_menu.id)
            self.assertEqual(employee.current_menu_path, str(root_menu.id))
            self.assertTrue(messenger.sent_menus)
            self.assertEqual(messenger.sent_menus[-1][2], [root_button.label])

    def test_bot_menu_api_rejects_reserved_navigation_button_labels(self) -> None:
        create_set_response = self.client.post(
            "/api/settings/menu-sets",
            json={"title": f"codex-root-{self.unique_tag}"},
        )
        self.assertEqual(create_set_response.status_code, 200)
        created_menu_set = next(
            (item for item in create_set_response.json()["menu_sets"] if item["title"] == f"codex-root-{self.unique_tag}"),
            None,
        )
        self.assertIsNotNone(created_menu_set)

        response = self.client.post(
            f"/api/settings/menu-sets/{created_menu_set['id']}/buttons",
            json={
                "label": MENU_HOME_BUTTON_TEXT,
                "action_type": "inactive",
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("зарезервированы", response.json()["detail"].lower())

    def test_bot_menu_broadcast_api_pushes_main_menu_to_linked_users(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            employee.candidate_work_stage = "testing"
            chat_id = str(950000000000 + (uuid4().int % 100000000000))
            set_primary_chat_id(employee, chat_id, db=db)
            root_menu = BotMenuSet(
                title=f"codex-root-{self.unique_tag}",
                description="root",
                sort_order=10,
                employee_scope="candidates",
            )
            child_menu = BotMenuSet(
                title=f"codex-child-{self.unique_tag}",
                description="child",
                sort_order=20,
                employee_scope="candidates",
            )
            db.add_all([root_menu, child_menu])
            db.commit()
            db.refresh(root_menu)
            db.refresh(child_menu)
            hr_settings = db.query(HrSettings).first()
            self.assertIsNotNone(hr_settings)
            hr_settings.default_menu_set_id = root_menu.id
            root_button = BotMenuButton(
                menu_set_id=root_menu.id,
                label=f"codex-root-btn-{self.unique_tag}",
                sort_order=10,
                action_type="open_set",
                target_menu_set_id=child_menu.id,
            )
            child_button = BotMenuButton(
                menu_set_id=child_menu.id,
                label=f"codex-child-btn-{self.unique_tag}",
                sort_order=10,
                action_type="inactive",
            )
            db.add_all([root_button, child_button])
            employee.current_menu_set_id = child_menu.id
            employee.current_menu_path = f"{root_menu.id},{child_menu.id}"
            db.commit()

        messenger = DummyMessenger()
        with (
            patch("app.web.settings_routes.settings.TELEGRAM_BOT_TOKEN", "test-token"),
            patch("app.web.settings_routes.create_telegram_messenger", return_value=messenger),
        ):
            response = self.client.post("/api/settings/bot-menu/broadcast")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertGreaterEqual(payload["refreshed_count"], 1)
        self.assertTrue(
            any(buttons == [f"codex-root-btn-{self.unique_tag}"] for _, _, buttons in messenger.sent_menus)
        )

        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            root_menu_id = db.query(BotMenuSet.id).filter(BotMenuSet.title == f"codex-root-{self.unique_tag}").scalar()
            self.assertIsNotNone(root_menu_id)
            self.assertEqual(employee.current_menu_set_id, root_menu_id)
            self.assertEqual(employee.current_menu_path, str(root_menu_id))

    def test_documents_workspace_api_returns_created_link_item(self) -> None:
        response = self.client.post(
            "/api/documents/links",
            json={
                "title": f"codex-doc-{self.unique_tag}",
                "description": "shared handbook",
                "category": "Регламенты",
                "external_url": "https://example.com/handbook",
                "is_active": True,
            },
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        titles = [item["title"] for item in payload["items"]]
        self.assertIn(f"codex-doc-{self.unique_tag}", titles)

    def test_documents_menu_scaffold_creates_root_and_category_sets(self) -> None:
        self.client.post(
            "/api/documents/links",
            json={
                "title": f"codex-doc-a-{self.unique_tag}",
                "description": "",
                "category": "HR",
                "external_url": "https://example.com/a",
                "is_active": True,
            },
        )
        self.client.post(
            "/api/documents/links",
            json={
                "title": f"codex-doc-b-{self.unique_tag}",
                "description": "",
                "category": "HR",
                "external_url": "https://example.com/b",
                "is_active": True,
            },
        )

        response = self.client.post(
            "/api/documents/menu-scaffold",
            json={"root_title": f"codex-doc-root-{self.unique_tag}"},
        )

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(str(payload["created_root_menu_set_id"]).isdigit())
        self.assertEqual(payload["created_root_menu_title"], f"codex-doc-root-{self.unique_tag}")

        with SessionLocal() as db:
            root_menu = db.query(BotMenuSet).filter(BotMenuSet.title == f"codex-doc-root-{self.unique_tag}").first()
            self.assertIsNotNone(root_menu)
            buttons = db.query(BotMenuButton).filter(BotMenuButton.menu_set_id == root_menu.id).all()
            self.assertTrue(any(button.action_type == "open_set" for button in buttons))

    def test_documents_menu_scaffold_rebuild_replaces_generated_sets_without_duplicates(self) -> None:
        self.client.post(
            "/api/documents/links",
            json={
                "title": f"codex-doc-a-{self.unique_tag}",
                "description": "",
                "category": "HR",
                "external_url": "https://example.com/a",
                "is_active": True,
            },
        )
        root_title = f"codex-doc-root-{self.unique_tag}"
        first_response = self.client.post(
            "/api/documents/menu-scaffold",
            json={"root_title": root_title, "mode": "create"},
        )
        self.assertEqual(first_response.status_code, 200)

        duplicate_response = self.client.post(
            "/api/documents/menu-scaffold",
            json={"root_title": root_title, "mode": "create"},
        )
        self.assertEqual(duplicate_response.status_code, 400)
        self.assertIn("пересборку", duplicate_response.json()["detail"].lower())

        rebuild_response = self.client.post(
            "/api/documents/menu-scaffold",
            json={"root_title": root_title, "mode": "rebuild"},
        )
        self.assertEqual(rebuild_response.status_code, 200)

        with SessionLocal() as db:
            system_tag = f"documents_scaffold:codex-doc-root-{self.unique_tag}"
            generated_sets = db.query(BotMenuSet).filter(BotMenuSet.system_tag == system_tag).all()
            root_menus = [menu_set for menu_set in generated_sets if menu_set.title == root_title]
            self.assertEqual(len(root_menus), 1)
            count_after_rebuild = len(generated_sets)

        rebuild_response_second = self.client.post(
            "/api/documents/menu-scaffold",
            json={"root_title": root_title, "mode": "rebuild"},
        )
        self.assertEqual(rebuild_response_second.status_code, 200)

        with SessionLocal() as db:
            system_tag = f"documents_scaffold:codex-doc-root-{self.unique_tag}"
            generated_sets_second = db.query(BotMenuSet).filter(BotMenuSet.system_tag == system_tag).all()
            root_menus_second = [menu_set for menu_set in generated_sets_second if menu_set.title == root_title]
            self.assertEqual(len(root_menus_second), 1)
            self.assertEqual(len(generated_sets_second), count_after_rebuild)

    def test_bot_menu_send_document_link_sends_text_message(self) -> None:
        with SessionLocal() as db:
            employee = db.get(Employee, self.employee_id)
            self.assertIsNotNone(employee)
            employee.employee_stage = "candidate"
            set_primary_chat_id(employee, f"77{self.employee_id}{int(self.unique_tag[:6], 16) % 1000000:06d}", db=db)
            menu_set = BotMenuSet(
                title=f"codex-docs-{self.unique_tag}",
                description="documents",
                sort_order=10,
                employee_scope="candidates",
            )
            document = DocumentLibraryItem(
                title=f"codex-link-{self.unique_tag}",
                description="policy",
                category="Docs",
                item_kind="link",
                external_url="https://example.com/policy",
                original_filename=None,
                stored_path=None,
                mime_type=None,
                file_size=None,
                is_active=True,
                sort_order=10,
                created_at=datetime.now(UTC).replace(tzinfo=None),
                updated_at=datetime.now(UTC).replace(tzinfo=None),
            )
            db.add_all([menu_set, document])
            db.commit()
            db.refresh(menu_set)
            db.refresh(document)
            button = BotMenuButton(
                menu_set_id=menu_set.id,
                label=f"codex-doc-button-{self.unique_tag}",
                sort_order=10,
                action_type="send_document",
                document_item_id=document.id,
            )
            db.add(button)
            employee.current_menu_set_id = menu_set.id
            db.commit()
            messenger = DummyMessenger()

            handled = asyncio.run(handle_menu_button(messenger, db, employee, button.label))

            self.assertTrue(handled)
            self.assertTrue(messenger.sent_texts)
            self.assertIn("https://example.com/policy", messenger.sent_texts[-1][1])
