import asyncio
from datetime import UTC, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import authenticate_account, create_admin_session_token
from app.database import SessionLocal, init_db
from app.main import AUTH_COOKIE_NAME, app
from app.models import Employee, EmployeeFeedbackRecipient, EmployeeFeedbackRun, FlowStepTemplate, ScenarioProgress, ScenarioTemplate, SurveyAnswer
from app.scenario_engine import get_or_create_progress, handle_text_response, start_scenario
from app.web.employees import _delete_employee_active_scenario_runtime_state, _delete_employee_record
from app.web.scenarios import _delete_template_entity


class FakeMessenger:
    def __init__(self):
        self.messages = []

    async def send_text(self, chat_id, text, reply_markup=None):
        self.messages.append((chat_id, text))

    async def send_photo_path(self, chat_id, path, filename=None, reply_markup=None, caption=None):
        pass

    async def close(self):
        pass


class FailingMessenger(FakeMessenger):
    async def send_text(self, chat_id, text, reply_markup=None):
        raise RuntimeError("simulated Telegram failure")


class EmployeeFeedbackSurveyTests(TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.client = TestClient(app)
        with SessionLocal() as db:
            account = authenticate_account(db, "admin", "admin123")
            cls.client.cookies.set(AUTH_COOKIE_NAME, create_admin_session_token(account.id))

    def setUp(self):
        suffix = uuid4().hex[:12]
        self.scenario_key = f"feedback_{suffix}"
        with SessionLocal() as db:
            subject = Employee(
                full_name="Анна Объект", desired_position="Дизайнер", employee_stage="staff",
                created_at=datetime.now(UTC).replace(tzinfo=None), is_flow_scheduled=False,
            )
            first = Employee(
                full_name="Борис Отвечающий", telegram_user_id=f"{uuid4().int % 10**12}",
                employee_stage="staff", created_at=datetime.now(UTC).replace(tzinfo=None), is_flow_scheduled=False,
            )
            second = Employee(
                full_name="Вера Отвечающая", telegram_user_id=f"{uuid4().int % 10**12}",
                employee_stage="staff", created_at=datetime.now(UTC).replace(tzinfo=None), is_flow_scheduled=False,
            )
            db.add_all([subject, first, second])
            db.flush()
            self.subject_id, self.first_id, self.second_id = subject.id, first.id, second.id
            db.add(ScenarioTemplate(
                scenario_key=self.scenario_key, title="Обратная связь", scenario_kind="survey",
                role_scope="all", employee_scope="employees", recipient_mode="self",
                trigger_mode="manual_only", sort_order=0,
            ))
            db.add(FlowStepTemplate(
                flow_key=self.scenario_key, step_key="question", step_title="Вопрос",
                default_text="Как работает {employee_full_name}, {position}?", response_type="text", sort_order=0,
            ))
            db.commit()

    def _launch(self, recipient_ids, messenger):
        with patch("app.web.feedback_surveys.create_telegram_messenger", return_value=messenger):
            return self.client.post(
                f"/api/employees/{self.subject_id}/feedback-surveys",
                json={"scenario_key": self.scenario_key, "recipient_employee_ids": recipient_ids},
            )

    def _answer(self, employee_id, value, messenger):
        with SessionLocal() as db:
            employee = db.get(Employee, employee_id)
            self.assertTrue(asyncio.run(handle_text_response(messenger, db, employee, SimpleNamespace(text=value))))

    def _sheet_rows(self):
        response = self.client.get(f"/api/employees/{self.subject_id}/feedback-surveys/export")
        self.assertEqual(response.status_code, 200, response.text)
        return list(load_workbook(BytesIO(response.content), read_only=True).active.values)

    def test_two_respondents_append_to_one_download_and_repeat_keeps_history(self):
        messenger = FakeMessenger()
        response = self._launch([self.first_id, self.second_id], messenger)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(messenger.messages), 2)
        self.assertTrue(all("Анна Объект" in text and "Дизайнер" in text for _, text in messenger.messages))
        self._answer(self.first_id, "Хорошо", messenger)
        self.assertEqual(len(self._sheet_rows()), 2)
        with SessionLocal() as db:
            db.get(Employee, self.first_id).full_name = "Борис Новое Имя"
            db.commit()
        self._answer(self.second_id, "Отлично", messenger)
        rows = self._sheet_rows()
        self.assertEqual(len(rows), 3)
        self.assertEqual({row[2]: row[4] for row in rows[1:]}, {
            "Борис Отвечающий": "Хорошо", "Вера Отвечающая": "Отлично",
        })
        with SessionLocal() as db:
            self.assertEqual(db.get(Employee, self.first_id).candidate_status, None)
            self.assertEqual(db.get(Employee, self.second_id).candidate_status, None)
        response = self._launch([self.first_id], messenger)
        self.assertEqual(response.status_code, 200, response.text)
        self._answer(self.first_id, "Ещё лучше", messenger)
        self.assertEqual(len(self._sheet_rows()), 4)
        with SessionLocal() as db:
            run_ids = [run.id for run in db.query(EmployeeFeedbackRun).filter_by(subject_employee_id=self.subject_id).all()]
            self.assertEqual(len(run_ids), 2)
            self.assertEqual(db.query(SurveyAnswer).filter(SurveyAnswer.feedback_run_id.in_(run_ids)).count(), 3)
            scenario_id = db.query(ScenarioTemplate).filter_by(scenario_key=self.scenario_key).one().id
        global_export = self.client.get(f"/surveys/{scenario_id}/export")
        self.assertEqual(global_export.status_code, 200)
        self.assertEqual(len(list(load_workbook(BytesIO(global_export.content), read_only=True).active.values)), 1)
        with SessionLocal() as db:
            _delete_template_entity(db, db.get(ScenarioTemplate, scenario_id))
            db.commit()
        saved_rows = self._sheet_rows()
        self.assertEqual(len(saved_rows), 4)
        self.assertEqual({row[3] for row in saved_rows[1:]}, {"Как работает Анна Объект, Дизайнер?"})

    def test_active_progress_rejects_launch_without_replacing_answers(self):
        messenger = FakeMessenger()
        self.assertEqual(self._launch([self.first_id], messenger).status_code, 200)
        response = self._launch([self.first_id], messenger)
        self.assertEqual(response.status_code, 409)
        with SessionLocal() as db:
            self.assertEqual(db.query(EmployeeFeedbackRun).filter_by(subject_employee_id=self.subject_id).count(), 1)
            progress = db.query(ScenarioProgress).filter_by(employee_id=self.first_id, scenario_key=self.scenario_key).one()
            self.assertTrue(progress.waiting_for_response)
            run_id = progress.feedback_run_id
            self.assertFalse(asyncio.run(start_scenario(messenger, db, db.get(Employee, self.first_id), self.scenario_key)))
            db.refresh(progress)
            self.assertEqual(progress.feedback_run_id, run_id)
            progress.current_step_key = None
            db.commit()
            self.assertEqual(self._launch([self.first_id], messenger).status_code, 409)
            self.assertFalse(asyncio.run(start_scenario(messenger, db, db.get(Employee, self.first_id), self.scenario_key)))
            self.assertFalse(asyncio.run(start_scenario(
                messenger, db, db.get(Employee, self.first_id), self.scenario_key, feedback_run_id=run_id,
            )))
            db.refresh(progress)
            self.assertEqual(progress.feedback_run_id, run_id)
            self.assertEqual(len(messenger.messages), 1)

    def test_empty_progress_and_progress_addressed_to_someone_else_do_not_block(self):
        with SessionLocal() as db:
            get_or_create_progress(db, self.first_id, f"empty_{self.scenario_key}")
            other = get_or_create_progress(db, self.first_id, f"other_{self.scenario_key}")
            other.recipient_employee_id = self.second_id
            other.current_step_key = "question"
            other.waiting_for_response = True
            db.commit()
        messenger = FakeMessenger()
        response = self._launch([self.first_id], messenger)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(messenger.messages), 1)

    def test_active_progress_addressed_to_respondent_explains_conflict(self):
        other_key = f"other_{self.scenario_key}"
        with SessionLocal() as db:
            progress = get_or_create_progress(db, self.second_id, other_key)
            progress.recipient_employee_id = self.first_id
            progress.current_step_key = "question"
            progress.waiting_for_response = True
            db.commit()
        response = self._launch([self.first_id], FakeMessenger())
        self.assertEqual(response.status_code, 409)
        self.assertIn("Борис Отвечающий", response.json()["detail"])
        self.assertIn(other_key, response.json()["detail"])
        with SessionLocal() as db:
            self.assertEqual(db.query(EmployeeFeedbackRun).filter_by(subject_employee_id=self.subject_id).count(), 0)

    def test_same_scenario_context_progress_for_other_recipient_is_not_reset(self):
        with SessionLocal() as db:
            progress = get_or_create_progress(db, self.first_id, self.scenario_key)
            progress.recipient_employee_id = self.second_id
            progress.current_step_key = "question"
            progress.waiting_for_response = True
            db.commit()
            progress_id = progress.id
        response = self._launch([self.first_id], FakeMessenger())
        self.assertEqual(response.status_code, 409)
        with SessionLocal() as db:
            run = EmployeeFeedbackRun(
                subject_employee_id=self.subject_id, scenario_key=self.scenario_key,
                scenario_title="Обратная связь", created_at=datetime.now(UTC).replace(tzinfo=None),
            )
            db.add(run)
            db.flush()
            self.assertFalse(asyncio.run(start_scenario(
                FakeMessenger(), db, db.get(Employee, self.first_id), self.scenario_key, feedback_run_id=run.id,
            )))
            progress = db.get(ScenarioProgress, progress_id)
            self.assertEqual(progress.recipient_employee_id, self.second_id)
            self.assertEqual(progress.current_step_key, "question")
            self.assertTrue(progress.waiting_for_response)

    def test_excel_escapes_formula_answer(self):
        messenger = FakeMessenger()
        self.assertEqual(self._launch([self.first_id], messenger).status_code, 200)
        self._answer(self.first_id, "=2+2", messenger)
        self.assertEqual(self._sheet_rows()[1][4], "'=2+2")

    def test_question_snapshot_precedes_template_edit(self):
        messenger = FakeMessenger()
        self.assertEqual(self._launch([self.first_id], messenger).status_code, 200)
        with SessionLocal() as db:
            db.query(FlowStepTemplate).filter_by(flow_key=self.scenario_key, step_key="question").one().custom_text = "Другой вопрос"
            db.commit()
        self._answer(self.first_id, "Ответ на старый вопрос", messenger)
        self.assertEqual(self._sheet_rows()[1][3], "Как работает Анна Объект, Дизайнер?")

    def test_failed_first_send_does_not_leave_active_progress(self):
        response = self._launch([self.first_id], FailingMessenger())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["runs"][0]["failed_count"], 1)
        with SessionLocal() as db:
            progress = db.query(ScenarioProgress).filter_by(employee_id=self.first_id, scenario_key=self.scenario_key).one()
            self.assertTrue(progress.is_completed)
            self.assertFalse(progress.waiting_for_response)

    def test_survey_audience_is_visible_and_enforced(self):
        with SessionLocal() as db:
            scenario = db.query(ScenarioTemplate).filter_by(scenario_key=self.scenario_key).one()
            scenario.target_employee_id = self.first_id
            db.commit()
        payload = self.client.get(f"/api/employees/{self.subject_id}/feedback-surveys").json()
        survey = next(item for item in payload["surveys"] if item["key"] == self.scenario_key)
        self.assertEqual(survey["eligible_recipient_ids"], [self.first_id])
        response = self._launch([self.second_id], FakeMessenger())
        self.assertEqual(response.status_code, 400)

    def test_deleting_subject_removes_scoped_history(self):
        messenger = FakeMessenger()
        self.assertEqual(self._launch([self.first_id], messenger).status_code, 200)
        self._answer(self.first_id, "Ответ", messenger)
        with SessionLocal() as db:
            _delete_employee_record(db, db.get(Employee, self.subject_id))
        with SessionLocal() as db:
            self.assertEqual(db.query(EmployeeFeedbackRun).filter_by(subject_employee_id=self.subject_id).count(), 0)
            self.assertEqual(db.query(EmployeeFeedbackRecipient).filter_by(respondent_employee_id=self.first_id).count(), 0)
            self.assertEqual(db.query(SurveyAnswer).filter_by(employee_id=self.first_id, scenario_key=self.scenario_key).count(), 0)
            self.assertEqual(db.query(ScenarioProgress).filter_by(employee_id=self.first_id, scenario_key=self.scenario_key).count(), 0)

    def test_resetting_respondent_marks_unfinished_delivery_unavailable(self):
        self.assertEqual(self._launch([self.first_id], FakeMessenger()).status_code, 200)
        with SessionLocal() as db:
            _delete_employee_active_scenario_runtime_state(db, self.first_id)
            db.commit()
            row = db.query(EmployeeFeedbackRecipient).filter_by(respondent_employee_id=self.first_id).one()
            self.assertEqual(row.delivery_status, "unavailable")
            self.assertEqual(db.query(ScenarioProgress).filter_by(employee_id=self.first_id, scenario_key=self.scenario_key).count(), 0)

    def test_deleting_respondent_removes_personal_feedback(self):
        messenger = FakeMessenger()
        self.assertEqual(self._launch([self.first_id], messenger).status_code, 200)
        self._answer(self.first_id, "Личный ответ", messenger)
        with SessionLocal() as db:
            _delete_employee_record(db, db.get(Employee, self.first_id))
        with SessionLocal() as db:
            self.assertEqual(db.query(SurveyAnswer).filter_by(employee_id=self.first_id, scenario_key=self.scenario_key).count(), 0)
            self.assertEqual(db.query(EmployeeFeedbackRecipient).filter_by(respondent_employee_id=self.first_id).count(), 0)
        response = self.client.get(f"/api/employees/{self.subject_id}/feedback-surveys/export")
        self.assertEqual(response.status_code, 404)
