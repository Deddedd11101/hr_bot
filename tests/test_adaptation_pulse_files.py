import io
import sqlite3
from contextlib import closing
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook

from app.database import SessionLocal, init_db
from app.messaging.identity import set_primary_chat_id
from app.messaging.service import handle_saved_document, save_incoming_file
from app.models import Employee, EmployeeDocumentLink, EmployeeFile, EmployeeMessengerAccount, FlowStepTemplate, OnboardingEvent, ScenarioProgress, ScenarioTemplate, SurveyAnswer
from app.scenario_engine import handle_text_response, render_telegram_message_html, resolve_tagged_employee_documents, start_scenario
from app.time_utils import utc_now
from app.web.scenario_routes import export_survey_results
from tests.test_scenario_engine_smoke import FakeMessenger
from tools.repair_adaptation_pulse_steps import FIRST_DAY, ORGANIZATION, PULSE, repair


class AdaptationPulseFilesTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        init_db()

    def setUp(self):
        self.key = f"adaptation_files_{uuid4().hex[:12]}"
        self.ids = []
        with SessionLocal() as db:
            subject = Employee(full_name="Subject", employee_stage="staff", created_at=utc_now())
            manager = Employee(full_name="Manager", employee_stage="staff", created_at=utc_now())
            db.add_all([subject, manager])
            db.flush()
            self.ids = [subject.id, manager.id]
            subject.manager_employee_id = manager.id
            set_primary_chat_id(subject, "700001001", db=db)
            set_primary_chat_id(manager, "700001002", db=db)
            db.commit()

    def tearDown(self):
        with SessionLocal() as db:
            for model in (SurveyAnswer, ScenarioProgress, FlowStepTemplate, ScenarioTemplate):
                column = model.flow_key if model is FlowStepTemplate else model.scenario_key
                db.query(model).filter(column == self.key).delete(synchronize_session=False)
            for model in (EmployeeDocumentLink, EmployeeFile, EmployeeMessengerAccount, OnboardingEvent):
                db.query(model).filter(model.employee_id.in_(self.ids)).delete(synchronize_session=False)
            db.query(Employee).filter(Employee.id.in_(self.ids)).delete(synchronize_session=False)
            db.commit()

    async def test_manager_plan_upload_advances_and_attaches_to_subject(self):
        messenger = FakeMessenger()
        with SessionLocal() as db, TemporaryDirectory() as temp:
            subject, manager = [db.get(Employee, i) for i in self.ids]
            db.add(ScenarioTemplate(scenario_key=self.key, title="Plan", role_scope="all", employee_scope="employees", recipient_mode="manager", trigger_mode="manual_only"))
            db.add_all([
                FlowStepTemplate(flow_key=self.key, step_key="upload", step_title="Plan", sort_order=0, default_text="Upload", response_type="file", target_field="adaptation_tasks_url", send_mode="immediate"),
                FlowStepTemplate(flow_key=self.key, step_key="done", step_title="Done", sort_order=1, default_text="Plan saved", response_type="none", send_mode="immediate", is_terminal=True),
            ])
            db.commit()
            self.assertTrue(await start_scenario(messenger, db, subject, self.key))
            path = Path(temp) / "plan.pdf"
            path.write_bytes(b"plan")
            with patch("app.messaging.service.resolve_inbound_access", return_value=SimpleNamespace(state="ok", employee=manager)):
                sender, file_row, state = await save_incoming_file(db, "700001002", None, original_name="plan.pdf", stored_path=str(path), category="candidate_file", mime_type="application/pdf", file_size=4)
            self.assertEqual(state, "saved")
            self.assertEqual(sender.id, manager.id)
            self.assertEqual(file_row.employee_id, subject.id)
            self.assertTrue(await handle_saved_document(messenger, db, sender, file_row))
            self.assertEqual(messenger.texts[-1]["text"], "Plan saved")
            slot = db.query(EmployeeDocumentLink).filter_by(employee_id=subject.id, slot_key="probation_plan").one()
            self.assertEqual(slot.employee_file_id, file_row.id)
            rendered = render_telegram_message_html(db, "{probation_plan}", subject, utc_now().date(), None)
            self.assertEqual(rendered, "plan.pdf")
            self.assertEqual([f.id for f in resolve_tagged_employee_documents(db, "{probation_plan}", subject)], [file_row.id])
            progress = db.query(ScenarioProgress).filter_by(employee_id=subject.id, scenario_key=self.key).one()
            self.assertTrue(progress.is_completed)

    async def test_scenario_answers_survive_relaunch_and_export(self):
        messenger = FakeMessenger()
        with SessionLocal() as db:
            subject = db.get(Employee, self.ids[0])
            scenario = ScenarioTemplate(scenario_key=self.key, title="Pulse", role_scope="all", employee_scope="employees", recipient_mode="self", trigger_mode="manual_only", scenario_kind="scenario")
            db.add(scenario)
            db.add(FlowStepTemplate(flow_key=self.key, step_key="question", step_title="Question", sort_order=0, default_text="How are you?", response_type="text", target_field="survey_answer", send_mode="immediate", is_terminal=True))
            db.commit()
            for answer in ("First", "=1+1"):
                self.assertTrue(await start_scenario(messenger, db, subject, self.key))
                self.assertTrue(await handle_text_response(messenger, db, subject, SimpleNamespace(text=answer)))
            rows = db.query(SurveyAnswer).filter_by(scenario_key=self.key).order_by(SurveyAnswer.id).all()
            self.assertEqual([row.answer_value for row in rows], ["First", "=1+1"])
            with patch("app.web.scenario_routes.require_auth", return_value=None):
                response = export_survey_results(SimpleNamespace(), scenario.id, db)
            contents = b"".join([chunk async for chunk in response.body_iterator])
            sheet = load_workbook(io.BytesIO(contents)).active
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet.cell(3, 3).value, "=1+1")
            self.assertEqual(sheet.cell(3, 3).data_type, "s")


class AdaptationRepairTests(unittest.TestCase):
    def test_dry_run_apply_backup_and_idempotence(self):
        with TemporaryDirectory() as temp:
            path = Path(temp) / "fixture.db"
            backup = Path(temp) / "before.db"
            with closing(sqlite3.connect(path)) as db:
                db.executescript("CREATE TABLE scenario_templates (scenario_key TEXT, scenario_kind TEXT); CREATE TABLE flow_step_templates (id INTEGER PRIMARY KEY, flow_key TEXT, step_key TEXT, response_type TEXT, target_field TEXT);")
                db.executemany("INSERT INTO scenario_templates VALUES (?, 'scenario')", [(key,) for key in (FIRST_DAY, ORGANIZATION, PULSE)])
                old = f"{FIRST_DAY}_step_1783414191"
                db.executemany("INSERT INTO flow_step_templates VALUES (?,?,?,?,?)", [(256, FIRST_DAY, old, "buttons", None), (257, FIRST_DAY, old, "buttons", None), (254, ORGANIZATION, "upload", "file", "")])
                db.executemany("INSERT INTO flow_step_templates VALUES (?,?,?,?,?)", [(500+i, PULSE, f"q{i}", "text", "") for i in range(12)])
                db.commit()
            before = path.read_bytes()
            self.assertEqual(len(repair(path)), 14)
            self.assertEqual(path.read_bytes(), before)
            with self.assertRaises(ValueError):
                repair(path, apply=True, backup=backup, expected_plan_sha256="0" * 64)
            self.assertFalse(backup.exists())
            repair(path, apply=True, backup=backup)
            self.assertEqual(repair(path), [])
            with closing(sqlite3.connect(backup)) as db:
                self.assertEqual(db.execute("SELECT step_key FROM flow_step_templates WHERE id=257").fetchone()[0], old)
                self.assertEqual(db.execute("SELECT target_field FROM flow_step_templates WHERE id=254").fetchone()[0], "")
            with closing(sqlite3.connect(path)) as db:
                db.execute("UPDATE flow_step_templates SET response_type='buttons' WHERE id=254")
                db.commit()
            before_invalid_apply = path.read_bytes()
            invalid_backup = Path(temp) / "should-not-exist.db"
            with self.assertRaises(ValueError):
                repair(path, apply=True, backup=invalid_backup)
            self.assertFalse(invalid_backup.exists())
            self.assertEqual(path.read_bytes(), before_invalid_apply)
