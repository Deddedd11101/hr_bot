import io
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import Employee, FlowStepTemplate, ScenarioAnswerRun, ScenarioProgress, ScenarioTemplate, SurveyAnswer
from app.scenario_engine import handle_text_response, start_scenario
from app.web.scenario_results import export_scenario_result, scenario_results
from tests import test_adaptation_pulse_files as fixtures
from tests.test_scenario_engine_smoke import FakeMessenger


class ScenarioAnswerRunTests(fixtures.AdaptationPulseFilesTests):
    def tearDown(self):
        with SessionLocal() as db:
            db.query(ScenarioAnswerRun).filter_by(scenario_key=self.key).delete(synchronize_session=False)
            db.commit()
        super().tearDown()

    async def test_employee_history_and_exports_are_isolated_by_launch(self):
        with SessionLocal() as db:
            subject = db.get(Employee, self.ids[0])
            db.add(ScenarioTemplate(scenario_key=self.key, title="Pulse snapshot", scenario_kind="scenario", role_scope="all", employee_scope="employees", recipient_mode="self"))
            db.add(FlowStepTemplate(flow_key=self.key, step_key="q", step_title="Question", default_text="How?", response_type="text", target_field="survey_answer", is_terminal=True, send_mode="immediate"))
            db.commit()
            for value in ("first", "=2+2"):
                self.assertTrue(await start_scenario(FakeMessenger(), db, subject, self.key))
                self.assertTrue(await handle_text_response(FakeMessenger(), db, subject, SimpleNamespace(text=value)))
            runs = db.query(ScenarioAnswerRun).filter_by(scenario_key=self.key).order_by(ScenarioAnswerRun.id).all()
            self.assertEqual(len(runs), 2)
            self.assertTrue(all(run.status == "completed" for run in runs))
            with patch("app.web.scenario_results.require_api_auth"):
                history = scenario_results(SimpleNamespace(), subject.id, db)
                self.assertEqual([item["answer_count"] for item in history["runs"]], [1, 1])
                for run, value in zip(runs, ("first", "=2+2")):
                    response = export_scenario_result(SimpleNamespace(), subject.id, run.id, db)
                    content = b"".join([chunk async for chunk in response.body_iterator])
                    sheet = load_workbook(io.BytesIO(content)).active
                    self.assertEqual(sheet.max_row, 2)
                    self.assertEqual(sheet.cell(2, 3).value, value)
                    self.assertEqual(sheet.cell(2, 3).data_type, "s")
                with self.assertRaises(HTTPException) as raised:
                    export_scenario_result(SimpleNamespace(), self.ids[1], runs[0].id, db)
                self.assertEqual(raised.exception.status_code, 404)

    async def test_restart_preserves_interrupted_result_and_reset_linkage(self):
        from app.web.employees import _reset_employee_bot_linkage
        with SessionLocal() as db:
            subject = db.get(Employee, self.ids[0])
            db.add(ScenarioTemplate(scenario_key=self.key, title="Pulse", scenario_kind="scenario", role_scope="all", employee_scope="employees", recipient_mode="self"))
            db.add(FlowStepTemplate(flow_key=self.key, step_key="q", step_title="Question", default_text="How?", response_type="text", target_field="survey_answer", send_mode="immediate"))
            db.commit()
            self.assertTrue(await start_scenario(FakeMessenger(), db, subject, self.key))
            first_id = db.query(ScenarioProgress).filter_by(employee_id=subject.id, scenario_key=self.key).one().answer_run_id
            self.assertTrue(await start_scenario(FakeMessenger(), db, subject, self.key))
            self.assertEqual(db.get(ScenarioAnswerRun, first_id).status, "interrupted")
            _reset_employee_bot_linkage(db, subject)
            self.assertTrue(all(run.status == "interrupted" for run in db.query(ScenarioAnswerRun).filter_by(scenario_key=self.key)))


if __name__ == "__main__":
    unittest.main()
